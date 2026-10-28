import io
import struct
import zipfile

import pytest
from app.archive_limits import load_workbook, validate_workbook_archive
from app.security_scan import validate_evidence_type
from fastapi import HTTPException
from openpyxl import Workbook


def archive(members):
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as source:
        for name, content in members:
            source.writestr(name, content)
    return output.getvalue()


def test_normal_workbook_remains_readable_with_safe_archive_gate():
    workbook = Workbook()
    workbook.active.append(["Work package", "Value"])
    workbook.active.append(["WP-01", 42])
    output = io.BytesIO()
    workbook.save(output)
    validate_evidence_type("normal.xlsx", output.getvalue())
    loaded = load_workbook(output, read_only=True, data_only=True)
    assert loaded.active["B2"].value == 42
    loaded.close()


@pytest.mark.parametrize(
    "name", ["../escape.xml", "/absolute.xml", "xl\\worksheet.xml", "C:/file.xml"]
)
def test_archive_names_cannot_escape_or_alias(name):
    content = archive([(name.replace("\\", "/"), "<root/>")])
    if "\\" in name:
        content = content.replace(name.replace("\\", "/").encode(), name.encode())
    with pytest.raises(ValueError, match="unsafe"):
        validate_workbook_archive(content)


@pytest.mark.parametrize(
    "xml",
    [
        '<worksheet><dimension ref="A1:XFD1048576"/></worksheet>',
        '<worksheet><sheetData><row r="999999999"/></sheetData></worksheet>',
        '<worksheet><sheetData><row><c r="ZZZ1"/></row></sheetData></worksheet>',
        '<!DOCTYPE worksheet [<!ENTITY x "untrusted">]><worksheet>&x;</worksheet>',
        "<worksheet><broken></worksheet>",
    ],
)
def test_unsafe_xml_and_iteration_dimensions_fail_before_openpyxl(xml):
    content = archive([("xl/worksheets/sheet1.xml", xml)])
    with pytest.raises(ValueError):
        validate_workbook_archive(content)
    with pytest.raises(HTTPException) as upload:
        validate_evidence_type("untrusted.xlsx", content)
    assert upload.value.status_code == 422


def test_archive_expansion_and_directory_allocation_are_bounded():
    with pytest.raises(ValueError, match="expansion"):
        validate_workbook_archive(archive([("data.bin", b"0" * 1000000)]))
    small = bytearray(archive([("normal.xml", "<root/>")]))
    end = small.rfind(b"PK\x05\x06")
    struct.pack_into("<HH", small, end + 8, 60000, 60000)
    with pytest.raises(ValueError, match="directory exceeds"):
        validate_workbook_archive(bytes(small))


def test_duplicate_names_and_excess_sheets_fail_closed():
    with pytest.raises(ValueError, match="duplicate"):
        validate_workbook_archive(archive([("a.xml", "<x/>"), ("A.xml", "<x/>")]))
    with pytest.raises(ValueError, match="too many worksheets"):
        validate_workbook_archive(
            archive([(f"xl/worksheets/sheet{i}.xml", "<worksheet/>") for i in range(33)])
        )


def test_alternative_relationship_targets_do_not_bypass_xml_dimension_gate():
    for name in ("custom/SHEET.XML", "sheet.bin"):
        with pytest.raises(ValueError, match="dimensions"):
            validate_workbook_archive(
                archive([(name, '<worksheet><dimension ref="A1:A999999"/></worksheet>')])
            )
