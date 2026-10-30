from __future__ import annotations

import io
from datetime import date

import pytest
from app.services.openai_schedule import build_openai_schedule_workbook
from app.services.progress_workbook import analyze_progress_workbook
from openpyxl import load_workbook


def _workbook(percentages: tuple[float | None, float | None], evidence: str | None = None) -> bytes:
    tasks = [
        {
            "id": 1,
            "name": "Excavation",
            "duration_days": 10,
            "start": "2026-10-01",
            "finish": "2026-10-10",
            "dependency_ids": [],
            "boq_references": ["p1:1.1"],
        },
        {
            "id": 2,
            "name": "Concrete",
            "duration_days": 10,
            "start": "2026-10-11",
            "finish": "2026-10-20",
            "dependency_ids": [1],
            "boq_references": ["p1:1.2"],
        },
    ]
    content = build_openai_schedule_workbook(tasks)
    workbook = load_workbook(io.BytesIO(content))
    sheet = workbook["Progress Report"]
    sheet["G5"], sheet["G6"], sheet["J5"] = percentages[0], percentages[1], evidence
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_progress_analysis_separates_reported_evidence_and_flags_delay() -> None:
    result = analyze_progress_workbook(_workbook((50, None)), date(2026, 10, 10))
    assert result["counts"] == {"EARLY": 0, "ON TIME": 0, "LATE": 1, "NOT REPORTED": 1}
    assert result["rows"][0]["evidence_state"] == "REPORTED ONLY"
    assert result["rows"][0]["schedule_variance_days"] == -5.0
    assert result["potential_deadline_exposure_days"] == 5.0


def test_progress_analysis_recognizes_early_and_evidence_reference() -> None:
    result = analyze_progress_workbook(_workbook((90, 0), "Inspection 14"), date(2026, 10, 5))
    assert result["rows"][0]["status"] == "EARLY"
    assert result["rows"][0]["evidence_state"] == "EVIDENCE REFERENCED"


def test_progress_analysis_rejects_invalid_percentage() -> None:
    with pytest.raises(ValueError, match="from 0 to 100"):
        analyze_progress_workbook(_workbook((120, None)), date(2026, 10, 10))


@pytest.mark.parametrize("cutoff,expected", [(date(2026, 10, 10), -6.3), (date(2026, 10, 5), 6.3)])
def test_variance_rounding_matches_excel_for_positive_and_negative_half_days(cutoff, expected):
    workbook = load_workbook(io.BytesIO(_workbook((75, None))))
    workbook["Progress Report"]["D5"] = "25 days"
    output = io.BytesIO()
    workbook.save(output)
    result = analyze_progress_workbook(output.getvalue(), cutoff)
    assert result["rows"][0]["schedule_variance_days"] == expected
