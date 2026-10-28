from __future__ import annotations

import asyncio
import io
from datetime import date

import httpx
import pytest
from app.services.ms_project_export import (
    HEADERS,
    BoqItem,
    build_ms_project_workbook,
    parse_boq_items,
)
from app.services.openai_schedule import (
    build_openai_schedule_workbook,
    create_openai_schedule,
    validate_schedule_proposal,
)
from openpyxl import Workbook, load_workbook


def _template() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Tasks"
    sheet.append(HEADERS)
    sheet.append([1, "Dismantle Steel Doors", "4 days", "August 16, 2026", "August 19, 2026", None])
    sheet.append([2, "Other task", "2 days", "August 20, 2026", "August 21, 2026", "1"])
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_boq_scope_is_not_duplicated_as_a_second_task_schedule() -> None:
    text = "1.1 Dismantle Steel Doors  No. 2\n1.2 Lay asphalt surface  m2  30\n"
    assert len(parse_boq_items(text)) == 2
    content, report = build_ms_project_workbook(text, _template())
    workbook = load_workbook(io.BytesIO(content), read_only=True)
    rows = list(workbook["Tasks"].values)
    assert rows[0] == HEADERS
    assert [row[0] for row in rows[1:]] == list(range(1, len(rows)))
    assert rows[2][5] == "1"
    assert not any(row[1] == "Lay asphalt surface" for row in rows[1:])
    assert report["boq_items"] == 2
    assert report["proposed_durations"] == 0
    progress_rows = list(workbook["Progress Report"].values)
    assert progress_rows[3] == (
        "Subdivision",
        "Task ID",
        "Task Name",
        "Proposed Duration",
        "Planned Start",
        "Planned Finish",
        "Work Complete %",
        "Reporting Date",
        "Remarks",
        "Evidence Reference",
        "Responsible Party",
        "Planning Basis",
    )
    assert not any(row[2] == "Lay asphalt surface" for row in progress_rows[4:])
    review_rows = list(workbook["BOQ Review"].values)
    assert any(row[0] == "p1:1.2" and row[1] == "Lay asphalt surface" for row in review_rows)


def test_non_boq_pdf_text_has_an_actionable_error() -> None:
    text = "Daily Project Control Assistant\nFrom reporting to decision intelligence\n"
    try:
        parse_boq_items(text)
    except ValueError as exc:
        message = str(exc)
        assert "actual BOQ" in message
        assert "presentation" in message
    else:
        raise AssertionError("Non-BOQ content was incorrectly accepted")


def test_boq_parser_preserves_part_chapter_and_section_context() -> None:
    text = "PART II\nCHAPTER MECHANICAL\nSECTION HEATING\n1.1 Supply boilers  No. 2\n"
    item = parse_boq_items(text)[0]
    assert item.reference == "p1:1.1"
    assert item.hierarchy == ("PART II", "CHAPTER MECHANICAL", "SECTION HEATING")


def test_layout_only_leaves_all_planning_fields_blank() -> None:
    content, _ = build_ms_project_workbook(
        "1.1 Lay asphalt surface  m2  30",
        _template(),
        reuse_template_plan=False,
        propose_missing_durations=False,
    )
    rows = list(load_workbook(io.BytesIO(content), read_only=True)["Tasks"].values)
    assert rows[1] == (1, "Lay asphalt surface", None, None, None, None)


def test_generated_output_cannot_be_reused_as_a_template() -> None:
    content, _ = build_ms_project_workbook("1.1 Dismantle Steel Doors  No. 2", _template())
    try:
        build_ms_project_workbook("1.1 Dismantle Steel Doors  No. 2", content)
    except ValueError as exc:
        assert "generated output workbook" in str(exc)
    else:
        raise AssertionError("Generated output was incorrectly accepted as a workflow template")


def test_incompatible_workflow_fails_instead_of_appending_boq_lines() -> None:
    text = "\n".join(
        f"{index}.1 Unrelated measured scope {index}  m2  {index * 10}" for index in range(1, 21)
    )
    try:
        build_ms_project_workbook(text, _template())
    except ValueError as exc:
        assert "not compatible with this BOQ" in str(exc)
    else:
        raise AssertionError("Incompatible workflow was incorrectly presented as a schedule")


def test_openai_schedule_export_uses_only_six_column_tasks_sheet() -> None:
    content = build_openai_schedule_workbook(
        [
            {
                "id": 1,
                "name": "Mobilization",
                "duration_days": 10,
                "start": "2026-10-01",
                "finish": "2026-10-10",
                "dependency_ids": [],
            },
            {
                "id": 2,
                "name": "Substructure",
                "duration_days": 30,
                "start": "2026-10-11",
                "finish": "2026-11-09",
                "dependency_ids": [1],
            },
        ],
    )
    workbook = load_workbook(io.BytesIO(content), read_only=True)
    assert workbook.sheetnames == ["Progress Report", "Tasks"]
    assert list(workbook["Tasks"].values) == [
        HEADERS,
        (1, "Mobilization", "10 days", "2026-10-01", "2026-10-10", None),
        (2, "Substructure", "30 days", "2026-10-11", "2026-11-09", "1"),
    ]


def test_openai_schedule_timeout_returns_an_actionable_error(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    request_body: dict[str, object] = {}

    async def raise_timeout(*_args: object, **kwargs: object) -> None:
        request_body.update(kwargs["json"])  # type: ignore[arg-type]
        request = httpx.Request("POST", "https://api.openai.com/v1/responses")
        raise httpx.ReadTimeout("response exceeded timeout", request=request)

    monkeypatch.setattr(httpx.AsyncClient, "post", raise_timeout)
    with pytest.raises(ValueError, match="longer than five minutes.*BOQ remains saved"):
        asyncio.run(
            create_openai_schedule(
                [],
                api_key="test-key",
                model="test-model",
                project_start=date(2026, 10, 13),
                target_finish=None,
            )
        )
    assert request_body["reasoning"] == {"effort": "minimal"}
    assert request_body["max_output_tokens"] == 16000
    assert request_body["text"] == {
        "verbosity": "low",
        "format": request_body["text"]["format"],  # type: ignore[index]
    }


def _governed_task(
    task_id: int,
    name: str,
    references: list[str],
    predecessors: list[int],
) -> dict[str, object]:
    return {
        "id": task_id,
        "name": name,
        "duration_days": 2,
        "start": "2026-10-12",
        "finish": "2026-10-13",
        "relationships": [
            {"predecessor_id": value, "type": "FS", "lag_days": 0, "basis": "Physical sequence"}
            for value in predecessors
        ],
        "boq_references": references,
        "scope_type": "BOQ_REQUIREMENT" if references else "EXECUTION_SUPPORT",
        "duration_basis": "DOCUMENTED_PROJECT_INPUT",
        "duration_note": "Planner supplied duration",
    }


def test_schedule_qa_blocks_missing_and_duplicate_boq_coverage() -> None:
    items = [
        BoqItem("1.1", "Excavation", "m3", None, 1),
        BoqItem("1.2", "Backfill", "m3", None, 1),
    ]
    tasks = [
        _governed_task(1, "Excavation", ["p1:1.1"], []),
        _governed_task(2, "Duplicate excavation scope", ["p1:1.1"], [1]),
    ]
    status, findings, coverage = validate_schedule_proposal(tasks, items)
    assert status == "BLOCKED"
    assert any("not allocated" in finding.issue for finding in findings)
    assert any("more than once" in finding.issue for finding in findings)
    assert [item["coverage_status"] for item in coverage] == ["Requires Clarification", "Missing"]


def test_schedule_qa_blocks_backfill_without_below_grade_waterproofing_gate() -> None:
    tasks = [
        _governed_task(1, "Below-grade waterproofing", [], []),
        _governed_task(2, "Unrelated concrete work", [], [1]),
        _governed_task(3, "Backfilling", [], [2]),
    ]
    tasks[2]["relationships"] = []
    status, findings, _ = validate_schedule_proposal(tasks, [])
    assert status == "BLOCKED"
    assert any("waterproofing is not a predecessor" in finding.issue for finding in findings)


def test_schedule_qa_blocks_material_boq_system_without_distinct_activity() -> None:
    items = [BoqItem("1.1", "Supply two heating boilers", "No.", None, 4)]
    tasks = [_governed_task(1, "General mechanical works", ["p4:1.1"], [])]
    status, findings, _ = validate_schedule_proposal(tasks, items)
    assert status == "BLOCKED"
    assert any("Heating system" in finding.issue for finding in findings)


def test_governed_workbook_keeps_import_sheet_and_adds_review_evidence() -> None:
    tasks = [_governed_task(1, "Mobilization", [], [])]
    status, findings, coverage = validate_schedule_proposal(tasks, [])
    content = build_openai_schedule_workbook(
        tasks,
        coverage=coverage,
        findings=findings,
        qa_status=status,
    )
    workbook = load_workbook(io.BytesIO(content), read_only=True)
    assert workbook.sheetnames == ["Progress Report", "Tasks", "BOQ Coverage", "QA Review"]
    assert tuple(next(workbook["Tasks"].values)) == HEADERS
    assert workbook["QA Review"]["B1"].value == status
    progress = list(workbook["Progress Report"].values)
    assert progress[3][1:7] == (
        "Task ID",
        "Task Name",
        "Proposed Duration",
        "Planned Start",
        "Planned Finish",
        "Work Complete %",
    )
    assert progress[4][1:6] == (1, "Mobilization", "2 days", "2026-10-12", "2026-10-13")


def test_customer_ai_endpoint_accepts_neither_api_key_nor_template() -> None:
    import inspect

    from app.routers.ms_project import generate_ai_schedule

    parameters = inspect.signature(generate_ai_schedule).parameters
    assert "openai_api_key" not in parameters
    assert "template" not in parameters
