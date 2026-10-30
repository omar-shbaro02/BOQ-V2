import hashlib
import io
import uuid
from datetime import date

import pytest
from app.services.governed_workbook import parse_rows
from app.services.openai_schedule import build_openai_schedule_workbook
from app.services.progress_template import build_progress_template
from app.services.progress_workbook import analyze_progress_workbook
from fastapi import HTTPException
from openpyxl import load_workbook
from test_project_context_api import context_api as context_api


def source(legacy=False):
    content = build_openai_schedule_workbook(
        [
            {
                "id": 1,
                "name": "Excavation",
                "duration_days": 10,
                "start": "2026-10-01",
                "finish": "2026-10-10",
                "dependency_ids": [],
                "boq_references": ["p1:1.1"],
            },
            {
                "id": 2,
                "name": "Concrete",
                "duration_days": 10,
                "start": "2026-10-11",
                "finish": "2026-10-20",
                "dependency_ids": [1],
                "boq_references": ["p2:2.1"],
            },
        ]
    )
    workbook = load_workbook(io.BytesIO(content))
    workbook["Progress Report"]["G5"] = 99
    workbook["Progress Report"]["J5"] = "Prior company-specific observation"
    if legacy:
        del workbook["Progress Report"]
        coverage = workbook.create_sheet("BOQ Coverage")
        coverage.append(["BOQ Reference", "Mapped Activity IDs"])
        coverage.append(["p1:1.1", "1"])
        coverage.append(["p2:2.1", "2"])
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


@pytest.mark.parametrize("legacy", [False, True])
def test_blank_template_preserves_exact_schedule_subdivisions_and_source_hash(legacy):
    content = source(legacy)
    output = build_progress_template(content, project_id="project-a", as_of=date(2026, 10, 5))
    original = load_workbook(io.BytesIO(content))
    template = load_workbook(io.BytesIO(output))
    assert list(original["Tasks"].values) == list(template["Tasks"].values)
    assert template["Progress Report"]["A5"].value == 1
    assert template["Progress Report"]["B5"].value == "Excavation"
    assert template["Progress Report"]["F5"].value is None
    assert template["Progress Report"]["I5"].value is None
    for title in ("Progress Report", "Schedule Comparison"):
        assert "Subdivision" not in [cell.value for cell in template[title][4]]
    metadata = dict(template["Progress Template Metadata"].iter_rows(min_row=2, values_only=True))
    assert metadata["source_sha256"] == hashlib.sha256(content).hexdigest()
    assert metadata["project_id"] == "project-a"
    assert metadata["mode"] == "BLANK_TEMPLATE"
    assert len(parse_rows(output)) == 2
    assert analyze_progress_workbook(output, date(2026, 10, 5))["counts"]["NOT REPORTED"] == 2
    assert "ISNUMBER" in template["Schedule Comparison"]["G5"].value
    assert "ROUND" in template["Schedule Comparison"]["I5"].value


def test_sample_is_explicit_previewable_and_denied_governed_import():
    output = build_progress_template(
        source(), project_id="project-a", as_of=date(2026, 10, 5), sample=True
    )
    result = analyze_progress_workbook(output, date(2026, 10, 5))
    assert result["rows"][0]["reported_percent"] == 25
    assert result["rows"][1]["reported_percent"] == 0
    assert result["rows"][0]["evidence_state"] == "REPORTED ONLY"
    with pytest.raises(HTTPException, match="Fictional examples"):
        parse_rows(output)


def test_task_only_template_preserves_progress_input_positions():
    content = build_progress_template(source(), project_id="project-a", as_of=date(2026, 10, 5))
    workbook = load_workbook(io.BytesIO(content))
    sheet = workbook["Progress Report"]
    for column, value in {
        "F": 25,
        "G": "2026-10-05",
        "H": "Site update",
        "I": "Photo 12",
        "J": "Site engineer",
    }.items():
        sheet[f"{column}5"] = value
    output = io.BytesIO()
    workbook.save(output)
    row = parse_rows(output.getvalue())[0]
    assert row["task_id"] == 1
    assert row["task_name"] == "Excavation"
    assert row["reported_percent"] == 25
    assert row["report_date"] == date(2026, 10, 5)
    assert row["remarks"] == "Site update"
    assert row["reported_evidence_reference"] == "Photo 12"
    assert row["responsible_party"] == "Site engineer"
    preview = analyze_progress_workbook(output.getvalue(), date(2026, 10, 5))["rows"][0]
    assert preview["reported_percent"] == 25
    assert preview["evidence_reference"] == "Photo 12"


@pytest.mark.parametrize(
    "mutation,reason",
    [
        (lambda w: setattr(w["Tasks"]["A3"], "value", 1), "unique positive"),
        (lambda w: setattr(w["Tasks"]["D2"], "value", "=TODAY()"), "not formulas"),
        (lambda w: setattr(w["Tasks"]["E2"], "value", "2026-09-30"), "precede"),
    ],
)
def test_invalid_source_fails_closed(mutation, reason):
    workbook = load_workbook(io.BytesIO(source()))
    mutation(workbook)
    output = io.BytesIO()
    workbook.save(output)
    with pytest.raises(ValueError, match=reason):
        build_progress_template(output.getvalue(), project_id="project-a", as_of=date(2026, 10, 5))


def test_project_template_endpoint_is_scoped_and_does_not_reuse_other_downloads(context_api):
    client, _, project_id, headers = context_api
    endpoint = f"/api/v1/projects/{project_id}/progress/template"

    def request(url=endpoint, auth=headers, **extra):
        return client.post(
            url,
            headers=auth,
            files={"schedule": ("schedule.xlsx", source())},
            data={"as_of": "2026-10-05", **extra},
        )

    sample = request(sample="true")
    assert sample.status_code == 200, sample.text
    blank = request()
    assert blank.status_code == 200, blank.text
    assert "no-store" in blank.headers["cache-control"]
    assert (
        analyze_progress_workbook(blank.content, date(2026, 10, 5))["counts"]["NOT REPORTED"] == 2
    )
    assert request(auth={**headers, "X-VAI-Actor-ID": "outsider@test"}).status_code == 403
    assert request(url=f"/api/v1/projects/{uuid.uuid4()}/progress/template").status_code == 404
    assert request(auth={}).status_code in {401, 403}
