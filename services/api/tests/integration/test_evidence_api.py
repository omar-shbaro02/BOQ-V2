import uuid
from collections.abc import Generator
from datetime import date as dt_date
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from typing import Any

import pytest
from app.database import Base, get_db
from app.main import app
from app.models import (
    BoqLine,
    BoqNormalizationRun,
    BoqSourceRow,
    BoqSourceVersion,
    EvidenceArtifact,
    EvidenceItem,
    EvidenceRelation,
    ProposedScheduleActivity,
    ScheduleDraftGeneration,
    ScheduleLogicProposal,
    VerificationEvent,
)
from app.storage import LocalEvidenceStore, get_evidence_store
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool


@pytest.fixture
def evidence_api(
    tmp_path: Path,
) -> Generator[tuple[TestClient, sessionmaker[Session]], None, None]:
    engine = create_engine(
        "sqlite+pysqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    test_sessions = sessionmaker(bind=engine, expire_on_commit=False)

    def override_db() -> Generator[Session, None, None]:
        with test_sessions() as session:
            yield session

    app.dependency_overrides[get_db] = override_db
    app.dependency_overrides[get_evidence_store] = lambda: LocalEvidenceStore(tmp_path)
    client = TestClient(app)
    try:
        yield client, test_sessions
    finally:
        client.close()
        app.dependency_overrides.clear()
        Base.metadata.drop_all(engine)


def bootstrap(client: TestClient) -> tuple[str, str, dict[str, str], str]:
    actor = "evidence-admin@example.test"
    organization = client.post(
        "/api/v1/organizations",
        headers={"X-VAI-Actor-ID": actor, "X-VAI-Roles": "ORGANIZATION_ADMIN"},
        json={"name": "Evidence Org", "slug": "evidence-org"},
    )
    assert organization.status_code == 201, organization.text
    organization_id = organization.json()["id"]
    headers = {"X-VAI-Actor-ID": actor, "X-VAI-Organization-ID": organization_id}
    project = client.post(
        f"/api/v1/organizations/{organization_id}/projects",
        headers=headers,
        json={
            "code": "EVD-001",
            "name": "Evidence Pilot",
            "timezone": "Asia/Beirut",
            "currency": "USD",
            "delivery_model": "DESIGN_BID_BUILD",
            "reporting_cadence": "WEEKLY",
        },
    )
    assert project.status_code == 201, project.text
    project_id = project.json()["id"]
    controlled_object = client.post(
        f"/api/v1/projects/{project_id}/controlled-objects",
        headers=headers,
        json={"object_type": "BOQ_ITEM", "code": "BOQ-100", "name": "Concrete"},
    )
    assert controlled_object.status_code == 201, controlled_object.text
    return organization_id, project_id, headers, controlled_object.json()["id"]


def evidence_item(
    *,
    controlled_object_id: str,
    field_name: str = "installed_quantity",
    value: Any = 12,
    measurement_basis: str | None = None,
    supersedes_item_id: str | None = None,
) -> dict[str, Any]:
    return {
        "controlled_object_id": controlled_object_id,
        "field_name": field_name,
        "value": value,
        "unit": "m3",
        "measurement_basis": measurement_basis,
        "semantic_state": "REPORTED",
        "truth_type": "REPORTED_CLAIM",
        "as_of": "2026-09-04T08:00:00Z",
        "confidence": "0.65",
        "supersedes_item_id": supersedes_item_id,
    }


def test_artifact_upload_hash_download_and_csv_import_are_idempotent(
    evidence_api: tuple[TestClient, sessionmaker[Session]],
) -> None:
    client, sessions = evidence_api
    _, project_id, headers, _ = bootstrap(client)
    csv_content = (
        b"field_name,value,semantic_state,truth_type,as_of,object_code,unit,"
        b"measurement_basis,confidence\n"
        b"installed_quantity,0,REPORTED,REPORTED_CLAIM,2026-09-04T08:00:00Z,"
        b"BOQ-100,m3,,0.7\n"
    )
    uploaded = client.post(
        f"/api/v1/projects/{project_id}/evidence/artifacts",
        headers=headers,
        data={"source_type": "SITE_EXPORT", "source_id": "export-42"},
        files={"upload": ("evidence.csv", csv_content, "text/csv")},
    )
    assert uploaded.status_code == 201, uploaded.text
    artifact = uploaded.json()
    assert artifact["size_bytes"] == len(csv_content)
    assert len(artifact["sha256"]) == 64
    assert artifact["scan_result"] == "CLEAN"
    downloaded = client.get(
        f"/api/v1/projects/{project_id}/evidence/artifacts/{artifact['id']}/content",
        headers=headers,
    )
    assert downloaded.content == csv_content

    preview_url = f"/api/v1/projects/{project_id}/evidence/imports/preview"
    preview_headers = {**headers, "Idempotency-Key": "csv-run-001"}
    preview = client.post(
        preview_url,
        headers=preview_headers,
        json={"artifact_id": artifact["id"], "mapping": {}},
    )
    assert preview.status_code == 201, preview.text
    assert preview.json()["status"] == "PREVIEW"
    assert preview.json()["valid_rows"] == 1
    repeated = client.post(
        preview_url,
        headers=preview_headers,
        json={"artifact_id": artifact["id"], "mapping": {}},
    )
    assert repeated.json()["id"] == preview.json()["id"]

    commit_url = f"/api/v1/projects/{project_id}/evidence/imports/{preview.json()['id']}/commit"
    committed = client.post(commit_url, headers=headers)
    repeated_commit = client.post(commit_url, headers=headers)
    assert committed.status_code == repeated_commit.status_code == 200
    assert committed.json()["evidence_item_ids"] == repeated_commit.json()["evidence_item_ids"]
    with sessions() as session:
        assert len(list(session.scalars(select(EvidenceArtifact)))) == 1
        imported = list(session.scalars(select(EvidenceItem)))
        assert len(imported) == 1
        assert imported[0].value == 0


def test_boq_source_preserves_workbook_structure_rows_and_revision_lineage(
    evidence_api: tuple[TestClient, sessionmaker[Session]],
) -> None:
    client, sessions = evidence_api
    _, project_id, headers, _ = bootstrap(client)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Main BOQ"
    sheet.merge_cells("A1:D1")
    sheet["A1"] = "Bill of Quantities"
    sheet.append(["Item", "Description", "Quantity", "Amount", "Unit"])
    sheet.append(["1.01", "Concrete", 125, "=C3*80", "m3"])
    sheet.append(["1.02", "Concrete walls", 75, 6000, "m3"])
    sheet.append([None, "Subtotal", None, 10000, None])
    sheet.append(["2.01", "Supply of air handling unit", 2, 5000, "nr"])
    sheet.append(["2.02", "Consultant inspection and approval", 1, 500, "item"])
    sheet.append(["2.03", "Testing and commissioning", 1, 600, "item"])
    sheet.append(["2.04", "Site establishment and mobilization", 1, 700, "item"])
    sheet.append(["2.05", "Supply only of cables", 10, 800, "m"])
    sheet.append(["X", "Unclear commercial entry", None, None])
    sheet.column_dimensions["D"].hidden = True
    second = workbook.create_sheet("Provisional")
    second.append(["P1", "Allowance", None, 10000])
    content = BytesIO()
    workbook.save(content)

    upload_url = f"/api/v1/projects/{project_id}/evidence/artifacts"
    uploaded = client.post(
        upload_url,
        headers=headers,
        data={"source_type": "BOQ", "source_id": "contract-boq-v1"},
        files={
            "upload": (
                "project-boq.xlsx",
                content.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert uploaded.status_code == 201, uploaded.text
    source_url = f"/api/v1/projects/{project_id}/bootstrap/boq-sources"
    created = client.post(
        source_url,
        headers={**headers, "Idempotency-Key": "boq-source-v1"},
        json={"artifact_id": uploaded.json()["id"]},
    )
    assert created.status_code == 201, created.text
    source = created.json()
    assert source["version_number"] == 1
    assert source["content_sha256"] == uploaded.json()["sha256"]
    assert source["extraction_status"] == "EXTRACTED"
    assert source["extracted_row_count"] == 12
    assert [item["sheet_name"] for item in source["structure_manifest"]] == [
        "Main BOQ",
        "Provisional",
    ]
    assert any(warning["code"] == "MERGED_CELLS_PRESERVED" for warning in source["warnings"])
    assert any(
        warning["code"] == "HIDDEN_WORKBOOK_CONTENT_PRESERVED" for warning in source["warnings"]
    )

    repeated = client.post(
        source_url,
        headers={**headers, "Idempotency-Key": "boq-source-v1"},
        json={"artifact_id": uploaded.json()["id"]},
    )
    assert repeated.status_code == 201
    assert repeated.json()["id"] == source["id"]
    mismatch = client.post(
        source_url,
        headers={**headers, "Idempotency-Key": "boq-source-v1"},
        json={
            "artifact_id": uploaded.json()["id"],
            "prior_source_version_id": str(uuid.uuid4()),
        },
    )
    assert mismatch.status_code == 409

    detail = client.get(f"{source_url}/{source['id']}", headers=headers)
    assert detail.status_code == 200, detail.text
    formula_row = next(row for row in detail.json()["rows"] if row["row_number"] == 3)
    assert formula_row["raw_values"][3] == "=C3*80"
    assert formula_row["warnings"][0]["code"] == "FORMULA_PRESERVED_NOT_EVALUATED"
    with sessions() as session:
        assert len(list(session.scalars(select(BoqSourceVersion)))) == 1
        assert len(list(session.scalars(select(BoqSourceRow)))) == 12

    normalization_url = f"{source_url}/{source['id']}/normalizations"
    normalized = client.post(
        normalization_url,
        headers={**headers, "Idempotency-Key": "normalize-boq-v1"},
        json={},
    )
    assert normalized.status_code == 201, normalized.text
    summary = normalized.json()
    assert summary["total_rows"] == 12
    assert summary["schedule_relevant_rows"] == 6
    assert summary["review_required_rows"] == 2
    assert summary["header_rows"] == {"Main BOQ": 2, "Provisional": 0}
    replay = client.post(
        normalization_url,
        headers={**headers, "Idempotency-Key": "normalize-boq-v1"},
        json={},
    )
    assert replay.status_code == 201
    assert replay.json()["id"] == summary["id"]
    mismatch = client.post(
        normalization_url,
        headers={**headers, "Idempotency-Key": "normalize-boq-v1"},
        json={"header_rows": {"Main BOQ": 1}},
    )
    assert mismatch.status_code == 409

    normalization = client.get(f"{source_url}/{source['id']}/normalization", headers=headers)
    assert normalization.status_code == 200, normalization.text
    lines = normalization.json()["lines"]
    by_description = {line["description"]: line for line in lines if line["description"]}
    assert by_description["Concrete"]["classification"] == "DIRECT_EXECUTION_SCOPE"
    assert by_description["Concrete"]["quantity"] == "125.000000"
    assert by_description["Concrete"]["schedule_relevant"] is True
    assert by_description["Concrete"]["review_state"] == "REVIEW_REQUIRED"
    assert by_description["Subtotal"]["classification"] == "SUBTOTAL_TOTAL"
    assert by_description["Subtotal"]["schedule_relevant"] is False
    assert (
        by_description["Supply of air handling unit"]["classification"] == "PROCUREMENT_OR_SUPPLY"
    )
    assert by_description["Consultant inspection and approval"]["classification"] == (
        "APPROVAL_INSPECTION"
    )
    assert by_description["Testing and commissioning"]["classification"] == (
        "TESTING_COMMISSIONING"
    )
    assert by_description["Site establishment and mobilization"]["classification"] == (
        "PRELIMINARIES_GENERAL"
    )
    assert by_description["Supply only of cables"]["classification"] == (
        "MATERIAL_ONLY_NON_SCHEDULE"
    )
    assert by_description["Supply only of cables"]["schedule_relevant"] is False
    assert by_description["Unclear commercial entry"]["classification"] == (
        "UNKNOWN_REVIEW_REQUIRED"
    )
    assert by_description["Unclear commercial entry"]["review_state"] == "REVIEW_REQUIRED"
    assert any(
        line["classification"] == "SUMMARY_HEADER" and not line["schedule_relevant"]
        for line in lines
    )
    provisional_line = next(
        line for line in lines if line["classification"] == "PROVISIONAL_OR_ALLOWANCE"
    )
    assert provisional_line["schedule_relevant"] is False
    with sessions() as session:
        assert len(list(session.scalars(select(BoqNormalizationRun)))) == 1
        assert len(list(session.scalars(select(BoqLine)))) == 12

    structure_url = f"{source_url}/{source['id']}/planning-structures"
    structure_response = client.post(
        structure_url,
        headers={**headers, "Idempotency-Key": "structure-v1"},
        json={},
    )
    assert structure_response.status_code == 201, structure_response.text
    structure = structure_response.json()
    assert structure["status"] == "PROPOSED"
    assert structure["version_number"] == 1
    assert len(structure["line_mappings"]) == 6
    assert len(structure["unmapped_lines"]) == 6
    assert all(node["node_type"] != "AUTHORIZED" for node in structure["wbs_nodes"])
    repeated_structure = client.post(
        structure_url,
        headers={**headers, "Idempotency-Key": "structure-v1"},
        json={},
    )
    assert repeated_structure.json()["id"] == structure["id"]

    direct_package = next(
        package
        for package in structure["work_packages"]
        if package["classification"] == "DIRECT_EXECUTION_SCOPE"
    )
    direct_lines = [
        mapping
        for mapping in structure["line_mappings"]
        if mapping["work_package_id"] == direct_package["id"]
    ]
    assert len(direct_lines) == 2
    revision_url = f"/api/v1/projects/{project_id}/bootstrap/planning-structures"
    split = client.post(
        f"{revision_url}/{structure['id']}/revisions",
        headers={**headers, "Idempotency-Key": "structure-split"},
        json={
            "action": "SPLIT_PACKAGE",
            "package_ids": [direct_package["id"]],
            "line_ids": [direct_lines[0]["boq_line_id"]],
            "new_package_names": ["Concrete base", "Concrete walls"],
            "reason": "Planner separated crews and measurable scope",
        },
    )
    assert split.status_code == 201, split.text
    assert split.json()["version_number"] == 2
    assert split.json()["supersedes_version_id"] == structure["id"]
    split_ids = split.json()["change_summary"]["new_package_ids"]

    stale = client.post(
        f"{revision_url}/{structure['id']}/revisions",
        headers={**headers, "Idempotency-Key": "structure-stale"},
        json={
            "action": "REJECT_PROPOSAL",
            "reason": "This stale proposal should not be revisable",
        },
    )
    assert stale.status_code == 409

    merged = client.post(
        f"{revision_url}/{split.json()['id']}/revisions",
        headers={**headers, "Idempotency-Key": "structure-merge"},
        json={
            "action": "MERGE_PACKAGES",
            "package_ids": split_ids,
            "new_package_names": ["Concrete works"],
            "reason": "Planner confirmed one shared delivery package",
        },
    )
    assert merged.status_code == 201, merged.text
    merged_id = merged.json()["change_summary"]["new_package_id"]
    approval_mapping = next(
        mapping
        for mapping in merged.json()["line_mappings"]
        if next(line for line in lines if line["id"] == mapping["boq_line_id"])["classification"]
        == "APPROVAL_INSPECTION"
    )
    approval_package = next(
        package
        for package in merged.json()["work_packages"]
        if package["classification"] == "APPROVAL_INSPECTION"
    )
    remapped = client.post(
        f"{revision_url}/{merged.json()['id']}/revisions",
        headers={**headers, "Idempotency-Key": "structure-remap"},
        json={
            "action": "REMAP_LINE",
            "line_ids": [approval_mapping["boq_line_id"]],
            "target_package_id": approval_package["id"],
            "reason": "Planner confirmed approval scope package mapping",
        },
    )
    assert remapped.status_code == 201, remapped.text
    rejected = client.post(
        f"{revision_url}/{remapped.json()['id']}/revisions",
        headers={**headers, "Idempotency-Key": "structure-reject"},
        json={
            "action": "REJECT_PROPOSAL",
            "reason": "Planner rejected the proposal pending final review",
        },
    )
    assert rejected.status_code == 201, rejected.text
    assert rejected.json()["status"] == "REJECTED"
    accepted = client.post(
        f"{revision_url}/{rejected.json()['id']}/revisions",
        headers={**headers, "Idempotency-Key": "structure-accept"},
        json={
            "action": "ACCEPT_PROPOSAL",
            "reason": "Planner reviewed all proposed mappings and packages",
        },
    )
    assert accepted.status_code == 201, accepted.text
    assert accepted.json()["status"] == "REVIEWED"
    history = client.get(structure_url, headers=headers)
    assert [item["version_number"] for item in history.json()] == [1, 2, 3, 4, 5, 6]

    draft_url = f"{revision_url}/{accepted.json()['id']}/schedule-drafts"
    draft = client.post(
        draft_url,
        headers={**headers, "Idempotency-Key": "schedule-draft-v1"},
        json={
            "validation_owner": "planner@example.test",
            "productivity_inputs": [
                {
                    "work_package_id": merged_id,
                    "rate_per_working_day": "50",
                    "quantity_unit": "m3",
                    "source_type": "PROJECT_HISTORICAL_ACTUAL",
                    "source_reference": "Project productivity record PR-01",
                    "source_date": "2026-09-01",
                    "source_version": "1",
                    "confidence": "HIGH",
                    "review_state": "ACCEPTED",
                }
            ],
        },
    )
    assert draft.status_code == 201, draft.text
    assert draft.json()["draft_state"] == "VERIFICATION_REQUIRED"
    assert draft.json()["activity_count"] == 7
    assert draft.json()["unresolved_duration_count"] == 6
    repeated_draft = client.post(
        draft_url,
        headers={**headers, "Idempotency-Key": "schedule-draft-v1"},
        json={
            "validation_owner": "planner@example.test",
            "productivity_inputs": [
                {
                    "work_package_id": merged_id,
                    "rate_per_working_day": "50",
                    "quantity_unit": "m3",
                    "source_type": "PROJECT_HISTORICAL_ACTUAL",
                    "source_reference": "Project productivity record PR-01",
                    "source_date": "2026-09-01",
                    "source_version": "1",
                    "confidence": "HIGH",
                    "review_state": "ACCEPTED",
                }
            ],
        },
    )
    assert repeated_draft.json()["id"] == draft.json()["id"]
    draft_detail = client.get(
        f"/api/v1/projects/{project_id}/bootstrap/schedule-drafts/{draft.json()['id']}",
        headers=headers,
    )
    assert draft_detail.status_code == 200, draft_detail.text
    activities = draft_detail.json()["activities"]
    execution = next(item for item in activities if item["activity_type"] == "EXECUTION")
    assert execution["quantity"] == "200.000000"
    assert execution["unit"] == "m3"
    assert execution["duration_unrounded"] == "4.0000000000"
    assert execution["duration_working_days"] == "4.000000"
    assert execution["duration_basis"] == "PROJECT_PRODUCTIVITY"
    assert execution["duration_status"] == "CALCULATED"
    assert len(execution["boq_line_refs"]) == 2
    assert len(draft_detail.json()["assumptions"]) == 6
    assert all(
        item["duration_status"] == "VERIFICATION_REQUIRED"
        for item in activities
        if item["id"] != execution["id"]
    )
    inspection = next(item for item in activities if item["activity_type"] == "INSPECTION_RELEASE")
    testing = next(item for item in activities if item["activity_type"] == "TESTING_COMMISSIONING")
    delivery = next(item for item in activities if item["activity_type"] == "DELIVERY")
    logic_url = (
        f"/api/v1/projects/{project_id}/bootstrap/schedule-drafts/{draft.json()['id']}/logic"
    )
    logic_payload = {
        "validation_owner": "planner@example.test",
        "calendar": {
            "calendar_id": "PROJECT-6D",
            "name": "Reviewed six-day project calendar",
            "working_weekdays": [0, 1, 2, 3, 4, 5],
            "working_hours_per_day": "8",
            "holidays": ["2026-12-25"],
            "review_state": "ACCEPTED",
        },
        "dependencies": [
            {
                "predecessor_activity_id": inspection["id"],
                "successor_activity_id": testing["id"],
                "relation_type": "FINISH_TO_START",
                "lag_working_days": "0",
                "basis": "APPROVAL_RELEASE_PREREQUISITE",
                "source_reference": "Planner rule PR-INSPECT-01",
                "confidence": "HIGH",
                "review_state": "ACCEPTED",
            }
        ],
        "sequence_templates": [
            {
                "predecessor_activity_id": delivery["id"],
                "successor_activity_id": execution["id"],
                "relation_type": "FINISH_TO_START",
                "lag_working_days": "0",
                "basis": "VALIDATED_SEQUENCE_TEMPLATE",
                "source_reference": "Template PROCURE-INSTALL v1",
                "confidence": "MEDIUM",
                "review_state": "REVIEW_REQUIRED",
            }
        ],
        "milestones": [
            {
                "code": "MS-COMMISSION",
                "name": "Commissioning complete",
                "target_date": "2027-03-31",
                "source_type": "CONTRACT_EVIDENCE",
                "source_reference": "Contract milestone schedule C-01",
                "feeds_from_activity_ids": [testing["id"]],
                "material": True,
            }
        ],
        "constraints": [
            {
                "activity_id": execution["id"],
                "constraint_type": "START_NO_EARLIER_THAN",
                "constraint_date": "2026-10-01",
                "source_reference": "Project start notice NTP-01",
            }
        ],
    }
    logic = client.post(
        logic_url,
        headers={**headers, "Idempotency-Key": "schedule-logic-v1"},
        json=logic_payload,
    )
    assert logic.status_code == 201, logic.text
    logic_result = logic.json()
    assert len(logic_result["dependencies"]) == 4
    procurement_dependencies = [
        item for item in logic_result["dependencies"] if item["basis"] == "PROCUREMENT_PREREQUISITE"
    ]
    assert len(procurement_dependencies) == 2
    assert any(
        item["predecessor_activity_id"] == inspection["id"]
        and item["successor_activity_id"] == testing["id"]
        for item in logic_result["dependencies"]
    )
    assert logic_result["calendar"]["review_state"] == "ACCEPTED"
    assert logic_result["milestones"][0]["status"] == "PROPOSED"
    assert logic_result["milestones"][0]["authorized"] is False
    assert logic_result["constraints"][0]["status"] == "PROPOSED"
    replayed_logic = client.post(
        logic_url,
        headers={**headers, "Idempotency-Key": "schedule-logic-v1"},
        json=logic_payload,
    )
    assert replayed_logic.json()["id"] == logic_result["id"]
    loaded_logic = client.get(logic_url, headers=headers)
    assert loaded_logic.json()["id"] == logic_result["id"]

    workbook["Main BOQ"].append(["1.02", "Rebar", 15, 12000])
    revised_content = BytesIO()
    workbook.save(revised_content)
    revised_upload = client.post(
        upload_url,
        headers=headers,
        data={
            "source_type": "BOQ",
            "source_id": "contract-boq-v2",
            "supersedes_artifact_id": uploaded.json()["id"],
        },
        files={"upload": ("project-boq-v2.xlsx", revised_content.getvalue())},
    )
    assert revised_upload.status_code == 201, revised_upload.text
    revised = client.post(
        source_url,
        headers={**headers, "Idempotency-Key": "boq-source-v2"},
        json={
            "artifact_id": revised_upload.json()["id"],
            "prior_source_version_id": source["id"],
        },
    )
    assert revised.status_code == 201, revised.text
    assert revised.json()["version_number"] == 2
    assert revised.json()["prior_source_version_id"] == source["id"]
    revised_normalization = client.post(
        f"{source_url}/{revised.json()['id']}/normalizations",
        headers={**headers, "Idempotency-Key": "normalize-boq-v2"},
        json={},
    )
    assert revised_normalization.status_code == 201, revised_normalization.text
    delta = client.post(
        f"{source_url}/{revised.json()['id']}/revision-delta",
        headers={**headers, "Idempotency-Key": "delta-boq-v2"},
        json={},
    )
    assert delta.status_code == 201, delta.text
    assert delta.json()["prior_source_version_id"] == source["id"]
    assert len(delta.json()["added_scope"]) == 1
    assert delta.json()["schedule_effects"]["authorized_schedule_unchanged"] is True
    loaded_delta = client.get(
        f"{source_url}/{revised.json()['id']}/revision-delta", headers=headers
    )
    assert loaded_delta.status_code == 200
    assert loaded_delta.json()["id"] == delta.json()["id"]


@pytest.mark.parametrize(
    ("relation_type", "expected_finish", "expected_second_float"),
    [
        ("FINISH_TO_START", "2026-10-09", "0"),
        ("START_TO_START", "2026-10-07", "1"),
    ],
)
def test_bootstrap_cpm_review_authority_and_exports(
    evidence_api: tuple[TestClient, sessionmaker[Session]],
    relation_type: str,
    expected_finish: str,
    expected_second_float: str,
) -> None:
    client, sessions = evidence_api
    organization_id, project_id, headers, _ = bootstrap(client)
    generation_id = uuid.uuid4()
    first_id, second_id = uuid.uuid4(), uuid.uuid4()
    with sessions() as session:
        session.add(
            ScheduleDraftGeneration(
                id=generation_id,
                organization_id=uuid.UUID(organization_id),
                project_id=uuid.UUID(project_id),
                planning_structure_id=uuid.uuid4(),
                generator_version="test",
                draft_state="DRAFT_GENERATED",
                productivity_inputs=[],
                duration_inputs=[],
                warnings=[],
                activity_count=2,
                unresolved_duration_count=0,
                idempotency_key="seed-generation",
                request_hash="a" * 64,
                created_by="evidence-admin@example.test",
            )
        )
        for activity_id, code, duration in (
            (first_id, "ACT-001", "3"),
            (second_id, "ACT-002", "2"),
        ):
            session.add(
                ProposedScheduleActivity(
                    id=activity_id,
                    organization_id=uuid.UUID(organization_id),
                    project_id=uuid.UUID(project_id),
                    generation_id=generation_id,
                    activity_code=code,
                    wbs_node_id="WBS-1",
                    work_package_id=f"WP-{code}",
                    boq_line_refs=[str(uuid.uuid4())],
                    activity_name=f"Test {code}",
                    activity_type="EXECUTION",
                    description="Traceable execution task",
                    duration_working_days=Decimal(duration),
                    duration_unrounded=Decimal(duration),
                    duration_status="EXPLICIT",
                    duration_basis="EXPLICIT_PROJECT_INPUT",
                    confidence="HIGH",
                    review_state="ACCEPTED",
                    warnings=[],
                )
            )
        logic_id = uuid.uuid4()
        session.add(
            ScheduleLogicProposal(
                id=logic_id,
                organization_id=uuid.UUID(organization_id),
                project_id=uuid.UUID(project_id),
                generation_id=generation_id,
                logic_version="test",
                dependencies=[
                    {
                        "predecessor_activity_id": str(first_id),
                        "successor_activity_id": str(second_id),
                        "relation_type": relation_type,
                        "lag_working_days": "0",
                        "basis": "PHYSICAL_PREREQUISITE",
                        "source_reference": "test",
                        "confidence": "HIGH",
                        "review_state": "ACCEPTED",
                        "generated": False,
                    }
                ],
                milestones=[],
                constraints=[],
                calendar={
                    "calendar_id": "PROJECT",
                    "name": "Reviewed calendar",
                    "working_weekdays": [0, 1, 2, 3, 4],
                    "working_hours_per_day": "8",
                    "holidays": [],
                    "shift_pattern": None,
                    "review_state": "ACCEPTED",
                },
                sequence_templates=[],
                assumptions=[],
                warnings=[],
                review_state="ACCEPTED",
                idempotency_key="seed-logic",
                request_hash="b" * 64,
                created_by="evidence-admin@example.test",
            )
        )
        session.commit()

    calculation = client.post(
        f"/api/v1/projects/{project_id}/bootstrap/schedule-drafts/{generation_id}/calculations",
        headers={**headers, "Idempotency-Key": "calculate-valid"},
        json={"project_start": "2026-10-05"},
    )
    assert calculation.status_code == 201, calculation.text
    calculated = calculation.json()
    assert calculated["readiness"] == "PLANNER_REVIEW_REQUIRED"
    assert calculated["proposed_finish"] == expected_finish
    assert calculated["activity_results"][1]["total_float_working_days"] == expected_second_float
    assert calculated["activity_results"][0]["free_float_working_days"] == "0"
    assert calculated["input_hash"]
    replay = client.post(
        f"/api/v1/projects/{project_id}/bootstrap/schedule-drafts/{generation_id}/calculations",
        headers={**headers, "Idempotency-Key": "calculate-valid"},
        json={"project_start": "2026-10-05"},
    )
    assert replay.status_code == 201
    assert replay.json()["id"] == calculated["id"]
    assert replay.json()["input_hash"] == calculated["input_hash"]
    mismatch = client.post(
        f"/api/v1/projects/{project_id}/bootstrap/schedule-drafts/{generation_id}/calculations",
        headers={**headers, "Idempotency-Key": "calculate-valid"},
        json={"project_start": "2026-10-06"},
    )
    assert mismatch.status_code == 409

    review = client.post(
        f"/api/v1/projects/{project_id}/bootstrap/schedule-calculations/{calculated['id']}/reviews",
        headers={**headers, "Idempotency-Key": "review-valid"},
        json={
            "reason": "Planner checked logic, dates, calendar, and traceability.",
            "edits": [
                {
                    "field_path": "activities.ACT-001.responsible_owner",
                    "original_value": "PROJECT_DELIVERY_MANAGER",
                    "revised_value": "Site manager",
                    "reason": "Assign the accountable delivery owner.",
                }
            ],
        },
    )
    assert review.status_code == 201, review.text
    assert review.json()["state"] == "PM_APPROVAL_REQUIRED"
    assert review.json()["edit_history"][0]["actor_id"] == "evidence-admin@example.test"

    denied = client.post(
        f"/api/v1/projects/{project_id}/bootstrap/schedule-releases/{review.json()['id']}/approve",
        headers={**headers, "Idempotency-Key": "approve-denied"},
        json={
            "authority_grant_id": str(uuid.uuid4()),
            "approval_reference": "PM-001",
            "reason": "Approve reviewed project schedule for control.",
        },
    )
    assert denied.status_code == 403
    grant = client.post(
        f"/api/v1/projects/{project_id}/authority-grants",
        headers=headers,
        json={
            "actor_id": "evidence-admin@example.test",
            "authority_type": "CURRENT_SCHEDULE_APPROVAL",
            "valid_from": dt_date.today().isoformat(),
        },
    )
    assert grant.status_code == 201, grant.text
    approved = client.post(
        f"/api/v1/projects/{project_id}/bootstrap/schedule-releases/{review.json()['id']}/approve",
        headers={**headers, "Idempotency-Key": "approve-valid"},
        json={
            "authority_grant_id": grant.json()["id"],
            "approval_reference": "PM-001",
            "reason": "Approve reviewed project schedule for control.",
        },
    )
    assert approved.status_code == 201, approved.text
    assert approved.json()["state"] == "CURRENT_AUTHORIZED"
    duplicate_approval = client.post(
        f"/api/v1/projects/{project_id}/bootstrap/schedule-releases/{review.json()['id']}/approve",
        headers={**headers, "Idempotency-Key": "approve-again"},
        json={
            "authority_grant_id": grant.json()["id"],
            "approval_reference": "PM-002",
            "reason": "Attempt to approve the same reviewed release again.",
        },
    )
    assert duplicate_approval.status_code == 409
    current = client.get(
        f"/api/v1/projects/{project_id}/authorized-context/current/SCHEDULE", headers=headers
    )
    assert current.status_code == 200
    assert current.json()["id"] == approved.json()["authorized_context_id"]
    for export_format, media_type in (
        ("json", "application/json"),
        ("csv", "text/csv"),
        ("xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    ):
        exported = client.get(
            f"/api/v1/projects/{project_id}/bootstrap/schedule-releases/{approved.json()['id']}/export/{export_format}",
            headers=headers,
        )
        assert exported.status_code == 200
        assert exported.headers["content-type"].startswith(media_type)
        if export_format == "json":
            assert exported.json()["export_metadata"]["semantic_state"] == "CURRENT_AUTHORIZED"
            assert (
                exported.json()["export_metadata"]["authorized_context_id"]
                == approved.json()["authorized_context_id"]
            )
        if export_format == "csv":
            assert b"semantic_notice" in exported.content
            assert b"CURRENT_AUTHORIZED" in exported.content


@pytest.mark.parametrize(
    ("conflict_mode", "expected_finding"),
    [
        ("cycle", "DEPENDENCY_CYCLE"),
        ("constraint", "CONSTRAINT_CONFLICT"),
        ("milestone", "NEGATIVE_FLOAT"),
        ("invalid_lag", "INVALID_LAG"),
        ("invalid_duration", "INVALID_DURATION"),
        ("weak_critical", "WEAK_CRITICAL_ASSUMPTION"),
        ("missing_prerequisite", "MISSING_PREREQUISITE"),
    ],
)
def test_bootstrap_dependency_cycle_blocks_review(
    evidence_api: tuple[TestClient, sessionmaker[Session]],
    conflict_mode: str,
    expected_finding: str,
) -> None:
    client, sessions = evidence_api
    organization_id, project_id, headers, _ = bootstrap(client)
    generation_id = uuid.uuid4()
    ids = [uuid.uuid4(), uuid.uuid4()]
    with sessions() as session:
        session.add(
            ScheduleDraftGeneration(
                id=generation_id,
                organization_id=uuid.UUID(organization_id),
                project_id=uuid.UUID(project_id),
                planning_structure_id=uuid.uuid4(),
                generator_version="test",
                draft_state="DRAFT_GENERATED",
                productivity_inputs=[],
                duration_inputs=[],
                warnings=[],
                activity_count=2,
                unresolved_duration_count=0,
                idempotency_key="cycle-generation",
                request_hash="c" * 64,
                created_by="evidence-admin@example.test",
            )
        )
        for index, activity_id in enumerate(ids):
            session.add(
                ProposedScheduleActivity(
                    id=activity_id,
                    organization_id=uuid.UUID(organization_id),
                    project_id=uuid.UUID(project_id),
                    generation_id=generation_id,
                    activity_code=f"CYCLE-{index}",
                    wbs_node_id="WBS",
                    work_package_id="WP",
                    boq_line_refs=[str(uuid.uuid4())],
                    activity_name=f"Cycle {index}",
                    activity_type=(
                        ("SUBMITTAL", "PROCUREMENT")[index]
                        if conflict_mode == "missing_prerequisite"
                        else "EXECUTION"
                    ),
                    description="Cycle test",
                    duration_working_days=Decimal(
                        "1.5" if conflict_mode == "invalid_duration" and index == 1 else "1"
                    ),
                    duration_unrounded=Decimal(
                        "1.5" if conflict_mode == "invalid_duration" and index == 1 else "1"
                    ),
                    duration_status="EXPLICIT",
                    duration_basis="EXPLICIT_PROJECT_INPUT",
                    confidence="LOW" if conflict_mode == "weak_critical" else "HIGH",
                    review_state="ACCEPTED",
                    warnings=[],
                )
            )
        dependencies = [
            {
                "predecessor_activity_id": str(ids[index]),
                "successor_activity_id": str(ids[1 - index]),
                "relation_type": "FINISH_TO_START",
                "lag_working_days": "0.5" if conflict_mode == "invalid_lag" else "0",
                "basis": "EXPLICIT_IMPORTED",
                "source_reference": "cycle",
                "confidence": "HIGH",
                "review_state": "ACCEPTED",
                "generated": False,
            }
            for index in range(
                2
                if conflict_mode == "cycle"
                else 0
                if conflict_mode == "missing_prerequisite"
                else 1
            )
        ]
        session.add(
            ScheduleLogicProposal(
                id=uuid.uuid4(),
                organization_id=uuid.UUID(organization_id),
                project_id=uuid.UUID(project_id),
                generation_id=generation_id,
                logic_version="test",
                dependencies=dependencies,
                milestones=(
                    [
                        {
                            "code": "MS-EARLY",
                            "name": "Required early completion",
                            "target_date": "2026-10-05",
                            "source_type": "CONTRACT_EVIDENCE",
                            "source_reference": "test milestone",
                            "feeds_from_activity_ids": [str(ids[1])],
                            "material": True,
                            "status": "PROPOSED",
                            "authorized": False,
                        }
                    ]
                    if conflict_mode == "milestone"
                    else []
                ),
                constraints=(
                    [
                        {
                            "activity_id": str(ids[1]),
                            "constraint_type": "MUST_FINISH_ON",
                            "constraint_date": "2026-10-05",
                            "source_reference": "test constraint",
                        }
                    ]
                    if conflict_mode == "constraint"
                    else []
                ),
                calendar={
                    "calendar_id": "PROJECT",
                    "name": "Reviewed",
                    "working_weekdays": [0, 1, 2, 3, 4],
                    "working_hours_per_day": "8",
                    "holidays": [],
                    "review_state": "ACCEPTED",
                },
                sequence_templates=[],
                assumptions=[],
                warnings=[],
                review_state="ACCEPTED",
                idempotency_key="cycle-logic",
                request_hash="d" * 64,
                created_by="evidence-admin@example.test",
            )
        )
        session.commit()
    calculation = client.post(
        f"/api/v1/projects/{project_id}/bootstrap/schedule-drafts/{generation_id}/calculations",
        headers={**headers, "Idempotency-Key": "calculate-cycle"},
        json={"project_start": "2026-10-05"},
    )
    assert calculation.status_code == 201, calculation.text
    assert calculation.json()["readiness"] == "VALIDATION_BLOCKED"
    assert any(
        item["code"] == expected_finding for item in calculation.json()["validation_findings"]
    )
    review = client.post(
        f"/api/v1/projects/{project_id}/bootstrap/schedule-calculations/{calculation.json()['id']}/reviews",
        headers={**headers, "Idempotency-Key": "review-cycle"},
        json={"reason": "Attempt review despite the dependency cycle."},
    )
    assert review.status_code == 409


def test_bootstrap_real_xlsx_to_authorized_schedule_and_revision_isolation(
    evidence_api: tuple[TestClient, sessionmaker[Session]],
) -> None:
    client, _ = evidence_api
    _, project_id, headers, _ = bootstrap(client)
    prefix = f"/api/v1/projects/{project_id}"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Concrete works"
    sheet.append(["Item", "Description", "Quantity", "Unit"])
    sheet.append(["C-01", "Concrete foundations", 120, "m3"])
    sheet.append(["C-02", "Concrete columns", 80, "m3"])
    sheet.append([None, "Subtotal", 200, "m3"])
    content = BytesIO()
    workbook.save(content)
    upload = client.post(
        f"{prefix}/evidence/artifacts",
        headers=headers,
        data={"source_type": "BOQ", "source_id": "pilot-boq-1"},
        files={"upload": ("pilot-boq.xlsx", content.getvalue())},
    )
    assert upload.status_code == 201, upload.text
    source = client.post(
        f"{prefix}/bootstrap/boq-sources",
        headers={**headers, "Idempotency-Key": "pilot-source"},
        json={"artifact_id": upload.json()["id"]},
    )
    assert source.status_code == 201, source.text
    source_id = source.json()["id"]
    normalization = client.post(
        f"{prefix}/bootstrap/boq-sources/{source_id}/normalizations",
        headers={**headers, "Idempotency-Key": "pilot-normalize"},
        json={},
    )
    assert normalization.status_code == 201, normalization.text
    assert normalization.json()["schedule_relevant_rows"] == 2
    structure = client.post(
        f"{prefix}/bootstrap/boq-sources/{source_id}/planning-structures",
        headers={**headers, "Idempotency-Key": "pilot-structure"},
        json={},
    )
    assert structure.status_code == 201, structure.text
    assert len(structure.json()["line_mappings"]) == 2
    package_id = structure.json()["work_packages"][0]["id"]
    reviewed_structure = client.post(
        f"{prefix}/bootstrap/planning-structures/{structure.json()['id']}/revisions",
        headers={**headers, "Idempotency-Key": "pilot-structure-review"},
        json={"action": "ACCEPT_PROPOSAL", "reason": "Planner checked BOQ grouping and scope."},
    )
    assert reviewed_structure.status_code == 201, reviewed_structure.text
    draft = client.post(
        f"{prefix}/bootstrap/planning-structures/{reviewed_structure.json()['id']}/schedule-drafts",
        headers={**headers, "Idempotency-Key": "pilot-draft"},
        json={
            "validation_owner": "planner@example.test",
            "productivity_inputs": [
                {
                    "work_package_id": package_id,
                    "rate_per_working_day": "50",
                    "quantity_unit": "m3",
                    "source_type": "PROJECT_HISTORICAL_ACTUAL",
                    "source_reference": "Measured project record PR-01",
                    "source_version": "1",
                    "confidence": "HIGH",
                    "review_state": "ACCEPTED",
                }
            ],
        },
    )
    assert draft.status_code == 201, draft.text
    detail = client.get(f"{prefix}/bootstrap/schedule-drafts/{draft.json()['id']}", headers=headers)
    assert detail.status_code == 200
    assert len(detail.json()["activities"]) == 1
    assert detail.json()["activities"][0]["duration_working_days"] == "4.000000"
    assert len(detail.json()["activities"][0]["boq_line_refs"]) == 2
    logic = client.post(
        f"{prefix}/bootstrap/schedule-drafts/{draft.json()['id']}/logic",
        headers={**headers, "Idempotency-Key": "pilot-logic"},
        json={
            "validation_owner": "planner@example.test",
            "calendar": {
                "calendar_id": "PROJECT",
                "name": "Reviewed six-day calendar",
                "working_weekdays": [0, 1, 2, 3, 4, 5],
                "working_hours_per_day": 8,
                "holidays": [],
                "review_state": "ACCEPTED",
            },
        },
    )
    assert logic.status_code == 201, logic.text
    calculation = client.post(
        f"{prefix}/bootstrap/schedule-drafts/{draft.json()['id']}/calculations",
        headers={**headers, "Idempotency-Key": "pilot-cpm"},
        json={"project_start": "2026-10-05"},
    )
    assert calculation.status_code == 201, calculation.text
    assert calculation.json()["readiness"] == "PLANNER_REVIEW_REQUIRED"
    assert calculation.json()["proposed_finish"] == "2026-10-08"
    review = client.post(
        f"{prefix}/bootstrap/schedule-calculations/{calculation.json()['id']}/reviews",
        headers={**headers, "Idempotency-Key": "pilot-review"},
        json={"reason": "Planner validated four working days and the six-day calendar."},
    )
    assert review.status_code == 201, review.text
    grant = client.post(
        f"{prefix}/authority-grants",
        headers=headers,
        json={
            "actor_id": "evidence-admin@example.test",
            "authority_type": "CURRENT_SCHEDULE_APPROVAL",
            "valid_from": dt_date.today().isoformat(),
        },
    )
    assert grant.status_code == 201, grant.text
    approved = client.post(
        f"{prefix}/bootstrap/schedule-releases/{review.json()['id']}/approve",
        headers={**headers, "Idempotency-Key": "pilot-approve"},
        json={
            "authority_grant_id": grant.json()["id"],
            "approval_reference": "PM-SCH-001",
            "reason": "Approve reviewed BOQ-derived project control schedule.",
        },
    )
    assert approved.status_code == 201, approved.text
    context_id = approved.json()["authorized_context_id"]
    assert context_id
    revised_sheet = workbook.active
    revised_sheet.append(["C-03", "Concrete beams", 20, "m3"])
    revised_content = BytesIO()
    workbook.save(revised_content)
    revised_upload = client.post(
        f"{prefix}/evidence/artifacts",
        headers=headers,
        data={
            "source_type": "BOQ",
            "source_id": "pilot-boq-2",
            "supersedes_artifact_id": upload.json()["id"],
        },
        files={"upload": ("pilot-boq-v2.xlsx", revised_content.getvalue())},
    )
    assert revised_upload.status_code == 201, revised_upload.text
    revised_source = client.post(
        f"{prefix}/bootstrap/boq-sources",
        headers={**headers, "Idempotency-Key": "pilot-source-2"},
        json={
            "artifact_id": revised_upload.json()["id"],
            "prior_source_version_id": source_id,
        },
    )
    assert revised_source.status_code == 201, revised_source.text
    revised_id = revised_source.json()["id"]
    revised_normalization = client.post(
        f"{prefix}/bootstrap/boq-sources/{revised_id}/normalizations",
        headers={**headers, "Idempotency-Key": "pilot-normalize-2"},
        json={},
    )
    assert revised_normalization.status_code == 201, revised_normalization.text
    revised_structure = client.post(
        f"{prefix}/bootstrap/boq-sources/{revised_id}/planning-structures",
        headers={**headers, "Idempotency-Key": "pilot-structure-2"},
        json={},
    )
    assert revised_structure.status_code == 201, revised_structure.text
    unapproved_prior = client.post(
        f"{prefix}/bootstrap/boq-sources/{revised_id}/revision-delta",
        headers={**headers, "Idempotency-Key": "pilot-delta-unapproved"},
        json={"prior_release_id": review.json()["id"]},
    )
    assert unapproved_prior.status_code == 422
    delta = client.post(
        f"{prefix}/bootstrap/boq-sources/{revised_id}/revision-delta",
        headers={**headers, "Idempotency-Key": "pilot-delta"},
        json={"prior_release_id": approved.json()["id"]},
    )
    assert delta.status_code == 201, delta.text
    assert len(delta.json()["added_scope"]) == 1
    assert len(delta.json()["mapping_changes"]) >= 1
    assert delta.json()["current_authorized_context_id"] == context_id
    current = client.get(f"{prefix}/authorized-context/current/SCHEDULE", headers=headers)
    assert current.status_code == 200
    assert current.json()["id"] == context_id


def test_pdf_boq_is_preserved_but_requires_governed_extraction_adapter(
    evidence_api: tuple[TestClient, sessionmaker[Session]],
) -> None:
    client, _ = evidence_api
    _, project_id, headers, _ = bootstrap(client)
    uploaded = client.post(
        f"/api/v1/projects/{project_id}/evidence/artifacts",
        headers=headers,
        data={"source_type": "BOQ", "source_id": "scanned-boq"},
        files={"upload": ("scanned-boq.pdf", b"%PDF-1.4\n%%EOF", "application/pdf")},
    )
    assert uploaded.status_code == 201, uploaded.text
    source = client.post(
        f"/api/v1/projects/{project_id}/bootstrap/boq-sources",
        headers={**headers, "Idempotency-Key": "pdf-boq-1"},
        json={"artifact_id": uploaded.json()["id"]},
    )
    assert source.status_code == 201, source.text
    assert source.json()["extraction_status"] == "VERIFICATION_REQUIRED"
    assert source.json()["extracted_row_count"] == 0
    assert source.json()["warnings"][0]["code"] == "PDF_EXTRACTION_ADAPTER_REQUIRED"


def test_artifact_type_and_malware_gates_reject_before_storage(
    evidence_api: tuple[TestClient, sessionmaker[Session]],
) -> None:
    client, sessions = evidence_api
    _, project_id, headers, _ = bootstrap(client)
    url = f"/api/v1/projects/{project_id}/evidence/artifacts"
    metadata = {"source_type": "UNTRUSTED_UPLOAD", "source_id": "scan-test"}
    unsupported = client.post(
        url,
        headers=headers,
        data=metadata,
        files={"upload": ("payload.exe", b"not executable", "application/octet-stream")},
    )
    assert unsupported.status_code == 415
    eicar = client.post(
        url,
        headers=headers,
        data=metadata,
        files={
            "upload": (
                "eicar.txt",
                b"X5O!P%@AP[4\\PZX54(P^)7CC)7}$EICAR-STANDARD-ANTIVIRUS-TEST-FILE!$H+H*",
                "text/plain",
            )
        },
    )
    assert eicar.status_code == 422
    with sessions() as session:
        assert session.scalar(select(EvidenceArtifact.id)) is None


def test_xlsx_preview_and_invalid_rows_are_reported_before_commit(
    evidence_api: tuple[TestClient, sessionmaker[Session]],
) -> None:
    client, _ = evidence_api
    _, project_id, headers, _ = bootstrap(client)
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(
        [
            "field_name",
            "value",
            "semantic_state",
            "truth_type",
            "as_of",
            "object_code",
            "confidence",
        ]
    )
    sheet.append(
        [
            "installed_quantity",
            21,
            "REPORTED",
            "REPORTED_CLAIM",
            "2026-09-04T08:00:00Z",
            "BOQ-100",
            0.75,
        ]
    )
    content = BytesIO()
    workbook.save(content)
    upload = client.post(
        f"/api/v1/projects/{project_id}/evidence/artifacts",
        headers=headers,
        data={"source_type": "QS_WORKBOOK", "source_id": "qs-8"},
        files={
            "upload": (
                "quantities.xlsx",
                content.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    preview = client.post(
        f"/api/v1/projects/{project_id}/evidence/imports/preview",
        headers={**headers, "Idempotency-Key": "xlsx-run-001"},
        json={"artifact_id": upload.json()["id"], "mapping": {}},
    )
    assert preview.status_code == 201, preview.text
    assert (preview.json()["status"], preview.json()["valid_rows"]) == ("PREVIEW", 1)

    invalid_csv = (
        b"field_name,value,semantic_state,truth_type,as_of,object_code,confidence\n"
        b"physical_progress,0,REPORTED,REPORTED_CLAIM,2026-09-04T08:00:00Z,"
        b"BOQ-100,0.5\n"
    )
    invalid_upload = client.post(
        f"/api/v1/projects/{project_id}/evidence/artifacts",
        headers=headers,
        data={"source_type": "SITE_EXPORT", "source_id": "invalid-progress"},
        files={"upload": ("invalid.csv", invalid_csv, "text/csv")},
    )
    rejected = client.post(
        f"/api/v1/projects/{project_id}/evidence/imports/preview",
        headers={**headers, "Idempotency-Key": "invalid-run-001"},
        json={"artifact_id": invalid_upload.json()["id"], "mapping": {}},
    )
    assert rejected.status_code == 201, rejected.text
    assert rejected.json()["status"] == "REJECTED"
    assert rejected.json()["validation_errors"][0]["row"] == 2
    denied_commit = client.post(
        f"/api/v1/projects/{project_id}/evidence/imports/{rejected.json()['id']}/commit",
        headers=headers,
    )
    assert denied_commit.status_code == 409


def test_verification_creates_a_derived_fact_without_truth_inflation(
    evidence_api: tuple[TestClient, sessionmaker[Session]],
) -> None:
    client, sessions = evidence_api
    _, project_id, headers, object_id = bootstrap(client)
    created = client.post(
        f"/api/v1/projects/{project_id}/evidence/items",
        headers=headers,
        json=evidence_item(controlled_object_id=object_id),
    )
    assert created.status_code == 201, created.text
    source_id = created.json()["id"]
    verified = client.post(
        f"/api/v1/projects/{project_id}/evidence/items/{source_id}/verify",
        headers=headers,
        json={
            "method": "Quantity surveyor measurement",
            "outcome": "VERIFIED",
            "rationale": "Checked against signed measurement sheet",
            "verified_confidence": "0.95",
        },
    )
    assert verified.status_code == 201, verified.text
    result_id = verified.json()["result_item_id"]
    assert result_id != source_id
    verifications = client.get(
        f"/api/v1/projects/{project_id}/evidence/verifications",
        headers=headers,
        params={"item_id": source_id},
    )
    relations = client.get(
        f"/api/v1/projects/{project_id}/evidence/relations",
        headers=headers,
        params={"item_id": source_id},
    )
    assert [event["id"] for event in verifications.json()] == [verified.json()["id"]]
    assert relations.json()[0]["relation_type"] == "DERIVED_FROM"

    with sessions() as session:
        source = session.get(EvidenceItem, uuid.UUID(source_id))
        result = session.get(EvidenceItem, uuid.UUID(result_id))
        assert source is not None and result is not None
        assert (source.semantic_state, source.truth_type) == ("REPORTED", "REPORTED_CLAIM")
        assert (result.semantic_state, result.truth_type) == ("VERIFIED", "VERIFIED_FACT")
        assert result.derived_from_item_id == source.id
        assert session.scalar(
            select(VerificationEvent).where(
                VerificationEvent.id == uuid.UUID(verified.json()["id"])
            )
        )
        relation = session.scalar(
            select(EvidenceRelation).where(EvidenceRelation.from_item_id == result.id)
        )
        assert relation is not None and relation.relation_type == "DERIVED_FROM"


def test_contradictions_requests_reliability_and_supersession_preserve_history(
    evidence_api: tuple[TestClient, sessionmaker[Session]],
) -> None:
    client, sessions = evidence_api
    _, project_id, headers, object_id = bootstrap(client)
    item_url = f"/api/v1/projects/{project_id}/evidence/items"
    left = client.post(
        item_url, headers=headers, json=evidence_item(controlled_object_id=object_id, value=12)
    )
    right = client.post(
        item_url, headers=headers, json=evidence_item(controlled_object_id=object_id, value=18)
    )
    contradiction = client.post(
        f"/api/v1/projects/{project_id}/evidence/contradictions",
        headers=headers,
        json={
            "left_item_id": left.json()["id"],
            "right_item_id": right.json()["id"],
            "field_name": "installed_quantity",
            "material": True,
        },
    )
    assert contradiction.status_code == 201, contradiction.text
    resolved = client.post(
        f"/api/v1/projects/{project_id}/evidence/contradictions/"
        f"{contradiction.json()['id']}/resolve",
        headers=headers,
        json={
            "chosen_item_id": right.json()["id"],
            "resolution_policy": "SIGNED_MEASUREMENT_PREVAILS",
            "resolution_reason": "Right item has signed measurement support",
        },
    )
    assert resolved.json()["status"] == "RESOLVED"

    request = client.post(
        f"/api/v1/projects/{project_id}/evidence/requests",
        headers=headers,
        json={
            "controlled_object_id": object_id,
            "requested_fields": ["installed_quantity"],
            "reason": "Close quantity variance",
            "urgency": "URGENT",
            "owner_actor_id": "qs@example.test",
            "due_at": "2026-09-06T12:00:00Z",
        },
    )
    satisfied = client.post(
        f"/api/v1/projects/{project_id}/evidence/requests/{request.json()['id']}/satisfy",
        headers=headers,
        json={"evidence_item_id": right.json()["id"]},
    )
    assert satisfied.json()["status"] == "SATISFIED"

    reliability_url = f"/api/v1/projects/{project_id}/evidence/reliability"
    for score, date in (("0.55", "2026-09-01T00:00:00Z"), ("0.80", "2026-09-04T00:00:00Z")):
        response = client.post(
            reliability_url,
            headers=headers,
            json={
                "source_type": "SITE_TEAM",
                "source_id": "team-a",
                "evidence_class": "FIELD_REPORT",
                "score": score,
                "rationale": "Periodic source review",
                "valid_from": date,
            },
        )
        assert response.status_code == 201, response.text
    history = client.get(reliability_url, headers=headers, params={"source_id": "team-a"})
    assert [entry["score"] for entry in history.json()] == ["0.8000", "0.5500"]

    replacement = client.post(
        item_url,
        headers=headers,
        json=evidence_item(
            controlled_object_id=object_id,
            value=20,
            supersedes_item_id=left.json()["id"],
        ),
    )
    assert replacement.status_code == 201, replacement.text
    with sessions() as session:
        assert session.get(EvidenceItem, uuid.UUID(left.json()["id"])).status == "SUPERSEDED"
        assert len(list(session.scalars(select(EvidenceItem)))) == 3


@pytest.mark.parametrize(
    ("payload", "expected_status"),
    [
        ({"value": None, "truth_type": "REPORTED_CLAIM"}, 422),
        ({"value": None, "truth_type": "UNKNOWN"}, 201),
        ({"field_name": "physical_progress", "value": 0}, 422),
        (
            {
                "field_name": "physical_progress",
                "value": 0,
                "measurement_basis": "PHYSICAL_VERIFIED",
            },
            201,
        ),
    ],
)
def test_semantic_null_unknown_zero_and_progress_basis(
    evidence_api: tuple[TestClient, sessionmaker[Session]],
    payload: dict[str, Any],
    expected_status: int,
) -> None:
    client, _ = evidence_api
    _, project_id, headers, object_id = bootstrap(client)
    data = evidence_item(controlled_object_id=object_id)
    data.update(payload)
    response = client.post(
        f"/api/v1/projects/{project_id}/evidence/items", headers=headers, json=data
    )
    assert response.status_code == expected_status, response.text


def test_project_membership_is_required_for_evidence_access(
    evidence_api: tuple[TestClient, sessionmaker[Session]],
) -> None:
    client, _ = evidence_api
    organization_id, project_id, _, _ = bootstrap(client)
    denied = client.get(
        f"/api/v1/projects/{project_id}/evidence/items",
        headers={
            "X-VAI-Actor-ID": "outsider@example.test",
            "X-VAI-Organization-ID": organization_id,
        },
    )
    assert denied.status_code == 403
