"""Document 19's connected scenario uses public commands, never seeded project facts."""

import hashlib
import json
from datetime import UTC, datetime, timedelta
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from zoneinfo import ZoneInfo

import pytest
from app.services.presentation_render import inspect_source
from openpyxl import Workbook
from test_decision_case_api import open_case
from test_evidence_api import evidence_api as evidence_api
from test_project_context_api import record_payload
from test_project_control_api import bootstrap_project
from test_project_health import POLICY


def exercise_connected_acceptance(api, artifact_dir: Path, package_count: int = 1):
    client, _ = api
    _, project_id, headers = bootstrap_project(client, f"connected-acceptance-{package_count}")
    base = f"/api/v1/projects/{project_id}"
    today = datetime.now(UTC).date()
    sequence = 0
    evidence_ids = []
    commands = []

    def post(path, body, status=201):
        nonlocal sequence
        sequence += 1
        response = client.post(
            f"{base}{path}",
            headers={**headers, "Idempotency-Key": f"connected-{sequence}"},
            json=body,
        )
        assert response.status_code == status, f"{path}: {response.text}"
        commands.append({"path": path, "status": response.status_code})
        return response.json()

    def get(path):
        response = client.get(f"{base}{path}", headers=headers)
        assert response.status_code == 200, response.text
        return response.json()

    def grant(kind):
        return post(
            "/authority-grants",
            {
                "actor_id": headers["X-VAI-Actor-ID"],
                "authority_type": kind,
            },
        )["id"]

    def authorize(kind, payload, authority):
        source = post(
            "/authorized-context/versions",
            {
                "context_type": kind,
                "semantic_state": "PROPOSED",
                "effective_from": datetime.now(UTC).isoformat(),
                "payload": payload,
            },
        )
        return post(
            "/authorized-context/activate",
            {
                "source_version_id": source["id"],
                "authority_grant_id": grant(authority),
                "expected_current_context_id": None,
                "approval_reference": f"CONNECTED-{kind}-1",
                "reason": "Explicit human approval of the synthetic reviewed source.",
            },
        )

    def evidence(
        field,
        value,
        *,
        unit=None,
        basis=None,
        state="REPORTED",
        truth="REPORTED_CLAIM",
        controlled=None,
    ):
        item = post(
            "/evidence/items",
            {
                "controlled_object_id": (controlled or package)["id"],
                "field_name": field,
                "value": value,
                "unit": unit,
                "measurement_basis": basis,
                "semantic_state": state,
                "truth_type": truth,
                "as_of": observation_time.isoformat(),
                "confidence": "0.9",
            },
        )
        evidence_ids.append(item["id"])
        return item

    def verify(item):
        event = post(
            f"/evidence/items/{item['id']}/verify",
            {
                "method": "Synthetic quantity or record inspection",
                "outcome": "VERIFIED",
                "rationale": "Human checked the exact synthetic record; no liability finding.",
                "verified_confidence": "0.95",
            },
        )
        items = get("/evidence/items")
        result = next(row for row in items if row["id"] == event["result_item_id"])
        assert result["semantic_state"] == "VERIFIED"
        assert next(row for row in items if row["id"] == item["id"])["semantic_state"] == "REPORTED"
        evidence_ids.append(result["id"])
        return result

    # 1–3: identity, exact charter authority, explicit WBS graph and active recipient.
    charter = post("/project-context/records", record_payload())
    post(
        f"/project-context/records/{charter['id']}/transitions?expected_version=0",
        {
            "to_status": "REVIEWED",
            "reason": "Human reviewed project objectives and boundaries.",
        },
    )
    post(
        f"/project-context/records/{charter['id']}/transitions?expected_version=1",
        {
            "to_status": "APPROVED",
            "reason": "Accountable human approves the charter.",
            "authority_grant_id": grant("CHARTER_APPROVAL"),
            "approval_reference": "CHARTER-1",
        },
    )
    assert get(f"/project-context/records/{charter['id']}")["current_status"] == "APPROVED"
    root = post(
        "/controlled-objects", {"object_type": "AREA", "code": "WBS-1", "name": "Site works"}
    )
    names = [
        "Concrete foundations",
        "Concrete columns",
        "Concrete beams",
        "Concrete slabs",
        "Concrete stairs",
    ]
    assert package_count in {1, 5}
    packages = [
        post(
            "/controlled-objects",
            {"object_type": "WORK_PACKAGE", "code": f"WP-{index + 1:02}", "name": name},
        )
        for index, name in enumerate(names[:package_count])
    ]
    package = packages[0]
    for item in packages:
        post(
            "/relations",
            {"source_id": root["id"], "target_id": item["id"], "relation_type": "CONTAINS"},
        )
    recipient = post("/project-context/records", record_payload("STAKEHOLDER"))

    # 4–6: preserve the real workbook, review structure, reconcile primary allocations.
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Item", "Description", "Quantity", "Unit", "Rate", "Amount", "Currency", "WBS"])
    for index, item in enumerate(packages):
        sheet.append(
            [f"C-{index}-01", item["name"] + " primary", 120, "m3", 10, 1200, "USD", item["code"]]
        )
        sheet.append(
            [f"C-{index}-02", item["name"] + " secondary", 80, "m3", 10, 800, "USD", item["code"]]
        )
    content = BytesIO()
    workbook.save(content)
    artifact_dir.mkdir(parents=True, exist_ok=True)
    (artifact_dir / "operator-boq.xlsx").write_bytes(content.getvalue())
    sheet.cell(2, 3, 130)
    sheet.cell(2, 6, 1300)
    sheet.append(["REV-NEW", "Concrete retaining wall", 50, "m3", 10, 500, "USD", "WP-06"])
    workbook.save(artifact_dir / "operator-boq-revision.xlsx")
    upload = client.post(
        f"{base}/evidence/artifacts",
        headers=headers,
        data={"source_type": "BOQ", "source_id": "CONNECTED-BOQ-1"},
        files={"upload": ("connected-boq.xlsx", content.getvalue())},
    )
    assert upload.status_code == 201, upload.text
    source = post("/bootstrap/boq-sources", {"artifact_id": upload.json()["id"]})
    post(f"/bootstrap/boq-sources/{source['id']}/normalizations", {})
    structure = post(f"/bootstrap/boq-sources/{source['id']}/planning-structures", {})
    reviewed_structure = post(
        f"/bootstrap/planning-structures/{structure['id']}/revisions",
        {
            "action": "ACCEPT_PROPOSAL",
            "reason": "Planner reconciled BOQ, WBS and scope allocations.",
        },
    )
    assert any(node.get("code") == "WP-01" for node in reviewed_structure["wbs_nodes"])
    assert len(reviewed_structure["work_packages"]) == package_count
    draft = post(
        f"/bootstrap/planning-structures/{reviewed_structure['id']}/schedule-drafts",
        {
            "validation_owner": headers["X-VAI-Actor-ID"],
            "productivity_inputs": [
                {
                    "work_package_id": item["id"],
                    "rate_per_working_day": "100" if package_count > 1 else "50",
                    "quantity_unit": "m3",
                    "source_version": "1",
                    "source_type": "PROJECT_HISTORICAL_ACTUAL",
                    "source_reference": "Synthetic reviewed crew output",
                    "confidence": "HIGH",
                    "review_state": "ACCEPTED",
                }
                for item in reviewed_structure["work_packages"]
            ],
        },
    )
    detail = get(f"/bootstrap/schedule-drafts/{draft['id']}")
    assert len(detail["activities"]) == package_count
    node_codes = {node["id"]: node["code"] for node in reviewed_structure["wbs_nodes"]}
    activity_codes = {node_codes[item["wbs_node_id"]]: item for item in detail["activities"]}
    bindings = [(item, activity_codes[item["code"]]) for item in packages]
    activity = bindings[0][1]
    allocations = activity["boq_allocations"]
    assert len(allocations) == 2 and all(a["primary_counting"] for a in allocations)
    assert len({a["boq_line_id"] for a in allocations}) == 2
    assert sum(Decimal(a["allocated_quantity"]) for a in allocations) == 200
    assert sum(Decimal(a["allocated_value"]) for a in allocations) == 2000
    assert all(len(item["boq_allocations"]) == 2 for _, item in bindings)
    all_allocations = [allocation for _, item in bindings for allocation in item["boq_allocations"]]
    assert len({item["boq_line_id"] for item in all_allocations}) == package_count * 2
    assert sum(Decimal(item["allocated_value"]) for item in all_allocations) == 2000 * package_count
    post(
        f"/bootstrap/schedule-drafts/{draft['id']}/logic",
        {
            "validation_owner": headers["X-VAI-Actor-ID"],
            "dependencies": [
                {
                    "predecessor_activity_id": previous[1]["id"],
                    "successor_activity_id": following[1]["id"],
                    "relation_type": "FINISH_TO_START",
                    "lag_working_days": "0",
                    "basis": "PHYSICAL_PREREQUISITE",
                    "source_reference": "Reviewed synthetic concrete sequence",
                    "confidence": "HIGH",
                    "review_state": "ACCEPTED",
                }
                for previous, following in zip(bindings, bindings[1:], strict=False)
            ],
            "calendar": {
                "name": "Reviewed working week",
                "working_weekdays": [0, 1, 2, 3, 4],
                "working_hours_per_day": 8,
                "review_state": "ACCEPTED",
            },
        },
    )
    calculation = post(
        f"/bootstrap/schedule-drafts/{draft['id']}/calculations",
        {
            "project_start": (today - timedelta(days=14)).isoformat(),
        },
    )
    assert calculation["readiness"] == "PLANNER_REVIEW_REQUIRED"

    # 7: explicit planned quantities, independent human review and baseline authority.
    reviewed = post(
        f"/bootstrap/schedule-calculations/{calculation['id']}/reviews",
        {
            "reason": "Planner confirms the source trace and exact reviewed quantity plan.",
            "control_inputs": [
                {
                    "activity_code": bound_activity["activity_code"],
                    "controlled_object_id": controlled["id"],
                    "measurement_basis": "PHYSICAL_VERIFIED",
                    "source_reference": "Reviewed planned quantity curve revision 1",
                    "progress_plan": [
                        {
                            "as_of": today.isoformat(),
                            "numerator": "200",
                            "denominator": "200",
                            "unit": "m3",
                            "measurement_basis": "PHYSICAL_VERIFIED",
                        }
                    ],
                }
                for controlled, bound_activity in bindings
            ],
        },
    )
    approved = post(
        f"/bootstrap/schedule-releases/{reviewed['id']}/approve",
        {
            "authority_grant_id": grant("BASELINE_SCHEDULE_APPROVAL"),
            "authorize_as_baseline": True,
            "approval_reference": "SCHEDULE-BASELINE-1",
            "reason": "Human authorizes the reviewed baseline and current control schedule.",
        },
    )
    schedule = get("/authorized-context/current/SCHEDULE")
    assert schedule["id"] == approved["authorized_context_id"]
    assert len(schedule["payload"]["activities"]) == package_count
    assert len(schedule["payload"]["dependencies"]) == package_count - 1
    assert (
        schedule["source_version_id"]
        == approved["schedule_payload"]["authorization_lineage"]["baseline_context_id"]
    )
    baseline = next(
        item
        for item in get("/authorized-context/versions")
        if item["id"] == schedule["source_version_id"]
    )
    assert baseline["semantic_state"] == "BASELINE"
    assert baseline["payload"]["activities"] == schedule["payload"]["activities"]
    weighting = authorize(
        "PROGRESS_WEIGHTING",
        {
            "data_date": today.isoformat(),
            "weighting_basis": "BOQ_VALUE",
            "measurement_basis": "PHYSICAL_VERIFIED",
            "source_reference": f"Reviewed BOQ {source['id']}",
            "entries": [
                {"controlled_object_code": item["code"], "weight": str(Decimal(1) / package_count)}
                for item in packages
            ],
        },
        "PROGRESS_WEIGHTING_APPROVAL",
    )
    budget = authorize(
        "BUDGET",
        {
            "currency": "USD",
            "approved_budget": "2000",
            "controlled_object_code": "WP-01",
            "measurement_basis": "COST_VALUE",
        },
        "BUDGET_APPROVAL",
    )

    # 8: reported actual quantity, separate verification, exact authorized plan, daily update.
    observation_time = datetime.now(UTC)
    observations = []
    project_inputs = []
    for controlled in packages:
        planned_evidence = evidence(
            "planned_progress",
            "200",
            unit="m3",
            basis="PHYSICAL_VERIFIED",
            state="CURRENT_AUTHORIZED",
            truth="JUDGMENT",
            controlled=controlled,
        )
        reported = evidence(
            "physical_progress", "80", unit="m3", basis="PHYSICAL_VERIFIED", controlled=controlled
        )
        verified = verify(reported)
        pair = []
        for item, kind in ((planned_evidence, "PLANNED_AUTHORIZED"), (verified, "VERIFIED")):
            pair.append(
                post(
                    "/progress/measurements",
                    {
                        "evidence_item_id": item["id"],
                        "authorized_context_id": schedule["id"]
                        if kind == "PLANNED_AUTHORIZED"
                        else None,
                        "measurement_kind": kind,
                        "numerator": item["value"],
                        "denominator": "200",
                        "unit": "m3",
                    },
                )
            )
        observations.append((reported, verified, pair))
        project_inputs.append(
            {
                "controlled_object_id": controlled["id"],
                "planned_measurement_id": pair[0]["id"],
                "actual_measurement_id": pair[1]["id"],
            }
        )
    reported, verified, measurements = observations[0]
    cutoff = datetime.now(UTC)
    daily = record_payload("DAILY_UPDATE", evidence_item_ids=[reported["id"], verified["id"]])
    daily["payload"].update(
        {
            "data_date": cutoff.astimezone(ZoneInfo("Asia/Beirut")).date().isoformat(),
            "reporting_cutoff": cutoff.isoformat(),
            "quantities": [{"quantity": "80", "unit": "m3"}],
            "blockers": [],
        }
    )
    daily_record = post("/project-context/records", daily)

    # 9–10: independent cost ledgers and a separately verified delay, no entitlement inference.
    cost_ids = []
    for kind, value in (
        ("APPROVED_BUDGET", "2000"),
        ("COMMITMENT", "1200"),
        ("ACTUAL", "900"),
        ("ACCRUAL", "100"),
        ("EARNED_VALUE", "800"),
        ("PHYSICAL_VALUE", "800"),
    ):
        item = evidence(
            kind.lower(),
            value,
            unit="USD",
            basis="COST_VALUE",
            state="CURRENT_AUTHORIZED" if kind == "APPROVED_BUDGET" else "REPORTED",
            truth="JUDGMENT" if kind == "APPROVED_BUDGET" else "REPORTED_CLAIM",
        )
        if kind != "APPROVED_BUDGET":
            item = verify(item)
        cost_ids.append(
            post(
                "/cost/records",
                {
                    "evidence_item_id": item["id"],
                    "authorized_context_id": budget["id"] if kind == "APPROVED_BUDGET" else None,
                    "record_kind": kind,
                    "amount": value,
                    "currency": "USD",
                    "measurement_basis": "COST_VALUE",
                    "reporting_period_start": today.isoformat(),
                    "reporting_period_end": today.isoformat(),
                },
            )["id"]
        )
    delay = verify(evidence("schedule_variance_days", 8, unit="day"))
    case = open_case(client, project_id, headers, delay, "connected-delay")
    case = post(
        f"/decision-cases/{case['id']}/evidence",
        {
            "expected_version": case["version"],
            "evidence_item_ids": [
                item["id"]
                for item in get("/evidence/items")
                if item["id"] in evidence_ids and item["controlled_object_id"] == package["id"]
            ],
            "attachment_reason": "Freeze exact quantity, cost and delay sources.",
        },
        status=200,
    )["case"]
    frozen = post(
        f"/decision-cases/{case['id']}/snapshots",
        {
            "expected_version": case["version"],
            "data_date": datetime.now(UTC).isoformat(),
        },
    )
    # Both value bases are retained, but selecting them together must fail closed.
    cost_command = {
        "expected_version": frozen["case"]["version"],
        "snapshot_id": frozen["snapshot"]["id"],
        "controlled_object_id": package["id"],
        "cost_record_ids": cost_ids,
    }
    post(f"/decision-cases/{case['id']}/cost-assessments", cost_command, status=422)
    cost = post(
        f"/decision-cases/{case['id']}/cost-assessments",
        {**cost_command, "cost_record_ids": cost_ids[:-1]},
    )
    physical_cost = post(
        f"/decision-cases/{case['id']}/cost-assessments",
        {
            **cost_command,
            "expected_version": cost["case"]["version"],
            "cost_record_ids": [*cost_ids[:4], cost_ids[-1]],
        },
    )
    schedule_assessment = post(
        f"/decision-cases/{case['id']}/schedule-assessments",
        {
            "expected_version": physical_cost["case"]["version"],
            "snapshot_id": frozen["snapshot"]["id"],
            "controlled_object_id": package["id"],
            "activity_code": activity["activity_code"],
            "delay_evidence_item_id": delay["id"],
        },
    )
    assert schedule_assessment["assessment"]["schedule_context_id"] == schedule["id"]
    assert schedule_assessment["assessment"]["input_evidence_ids"] == [delay["id"]]
    assert Decimal(schedule_assessment["assessment"]["project_completion_exposure_days"]) == 8
    for assessment in (cost["assessment"], physical_cost["assessment"]):
        assert assessment["assessment_status"] == "ASSESSED"
        for field, expected in (
            ("current_authorized_budget", "2000"),
            ("commitments", "1200"),
            ("actuals", "900"),
            ("accruals", "100"),
            ("recognized_cost", "1000"),
            ("progress_value_ratio", "0.4"),
        ):
            assert Decimal(assessment[field]) == Decimal(expected), field
    assert Decimal(cost["assessment"]["estimate_at_completion"]) == Decimal("2500")
    assert physical_cost["assessment"]["estimate_at_completion"] is None

    # 11–14: one frozen dashboard projection, explainable RAG and read-only grounded considerations.
    control = post(
        "/project-control/snapshots",
        {
            "data_date": datetime.now(UTC).isoformat(),
            "progress_weighting_context_id": weighting["id"],
            "schedule_context_id": schedule["id"],
            "packages": project_inputs,
            "cost_assessment_ids": [cost["assessment"]["id"]],
            "schedule_assessment_ids": [schedule_assessment["assessment"]["id"]],
            "health_policy": {**POLICY, "owner_actor_id": headers["X-VAI-Actor-ID"]},
        },
    )
    dashboard = get(f"/project-control/snapshots/{control['id']}")
    assert dashboard == control
    assert len(control["package_results"]) == package_count
    assert len(control["health"]["packages"]) == package_count
    assert Decimal(control["shared_metrics"]["actual"]["percent"]) == 40
    assert Decimal(control["shared_metrics"]["planned"]["percent"]) == 100
    assert control["health"]["status"] == "RED"
    assert control["health"]["dimensions"]["progress"]["status"] == "RED"
    assert control["health"]["dimensions"]["schedule"]["status"] == "RED"
    assert control["health"]["dimensions"]["cost"]["status"] == "RED"
    support = control["supporting_context"]
    assert {charter["id"], recipient["id"], daily_record["id"]} <= {
        item["id"] for item in support["project_context"]
    }
    assert (
        next(item for item in support["project_context"] if item["id"] == charter["id"])["status"]
        == "APPROVED"
    )
    assert support["schedule_assessments"][0]["id"] == schedule_assessment["assessment"]["id"]
    assistant_response = client.post(
        f"{base}/project-control/snapshots/{control['id']}/assistant",
        headers=headers,
        json={"topic": "OVERVIEW"},
    )
    assert assistant_response.status_code == 200, assistant_response.text
    assistant = assistant_response.json()
    assert assistant["considerations"] and all(
        c["status"] == "MANAGEMENT_CONSIDERATION" for c in assistant["considerations"]
    )
    assert all(fact["source_refs"] for fact in assistant["facts"])
    assert next(f for f in assistant["facts"] if f["key"] == "forecast.unknown")["value"] is None

    # 15–17: independent report/deck approvals, exact frozen facts in every consumer.
    report = post(
        "/decision-center/report-snapshots",
        {
            "report_type": "ON_DEMAND_STATUS_REPORT",
            "project_control_snapshot_id": control["id"],
        },
    )
    report_grant = grant("REPORT_APPROVAL")
    for status in ("REVIEWED", "APPROVED"):
        post(
            f"/decision-center/report-snapshots/{report['id']}/lifecycle-events",
            {
                "to_status": status,
                "reason": "Explicit human review of the frozen project report.",
                "authority_grant_id": report_grant if status == "APPROVED" else None,
            },
        )
    deck = post(
        "/decision-center/presentations",
        {
            "report_snapshot_id": report["id"],
            "recipient_record_ids": [recipient["id"]],
        },
    )
    denied_download = client.get(
        f"{base}/decision-center/approved-presentations/{deck['id']}/content", headers=headers
    )
    assert denied_download.status_code == 409, denied_download.text
    for version, status in ((1, "REVIEWED"), (2, "APPROVED")):
        post(
            f"/decision-center/presentations/{deck['id']}/lifecycle-events",
            {
                "to_status": status,
                "expected_version": version,
                "reason": "Human reviewed exact presentation and intended stakeholder.",
                "authority_grant_id": report_grant if status == "APPROVED" else None,
            },
        )
    content_response = client.get(
        f"{base}/decision-center/approved-presentations/{deck['id']}/content", headers=headers
    )
    assert content_response.status_code == 200, content_response.text
    embedded = inspect_source(content_response.content)
    assert (
        embedded["shared_metrics"]
        == report["payload"]["project_control"]["shared_metrics"]
        == dashboard["shared_metrics"]
    )
    actual = next(f for f in assistant["facts"] if f["key"] == "completion.actual")
    assert actual["value"] == dashboard["shared_metrics"]["actual"]["percent"]
    assert embedded["report"]["payload"] == report["payload"]
    assert embedded["report"]["content_hash"] == report["content_hash"]
    assert (
        hashlib.sha256(
            json.dumps(report["payload"], sort_keys=True, separators=(",", ":")).encode()
        ).hexdigest()
        == report["content_hash"]
    )
    assert hashlib.sha256(content_response.content).hexdigest() == deck["content_hash"]
    report_detail = get(f"/decision-center/report-snapshots/{report['id']}")
    deck_detail = get(f"/decision-center/presentations/{deck['id']}")
    assert report_detail["lifecycle_status"] == deck_detail["lifecycle_status"] == "APPROVED"
    assert (
        deck_detail["source_report_approval_event_id"]
        == report_detail["lifecycle_events"][-1]["id"]
    )
    assert deck_detail["recipients"][0]["id"] == recipient["id"]
    for field in control["shared_metrics"]["cost_assessments"][0]["values"]:
        fact = next(f for f in assistant["facts"] if f["key"] == f"cost.0.{field}")
        assert Decimal(fact["value"]) == Decimal(
            control["shared_metrics"]["cost_assessments"][0]["values"][field]
        )
        if field in {"forecast_to_complete", "estimate_at_completion"}:
            assert fact["semantic_state"] == "FORECAST"

    # Later reported work cannot rewrite the frozen snapshot, advice, report or approved deck.
    observation_time = datetime.now(UTC)
    later = evidence("physical_progress", "160", unit="m3", basis="PHYSICAL_VERIFIED")
    later_daily = record_payload("DAILY_UPDATE", evidence_item_ids=[later["id"]])
    later_daily["record_key"] = "DAILY-LATER"
    later_daily["payload"].update(daily["payload"])
    later_daily["payload"]["reporting_cutoff"] = datetime.now(UTC).isoformat()
    later_daily["payload"]["quantities"] = [{"quantity": "160", "unit": "m3"}]
    post("/project-context/records", later_daily)
    assert get(f"/project-control/snapshots/{control['id']}") == dashboard
    assert get(f"/decision-center/report-snapshots/{report['id']}") == report_detail
    assert get(f"/decision-center/presentations/{deck['id']}") == deck_detail
    again = client.post(
        f"{base}/project-control/snapshots/{control['id']}/assistant",
        headers=headers,
        json={"topic": "OVERVIEW"},
    )
    assert again.status_code == 200, again.text
    assert again.json() == assistant
    again_content = client.get(
        f"{base}/decision-center/approved-presentations/{deck['id']}/content", headers=headers
    )
    assert again_content.status_code == 200, again_content.text
    assert again_content.content == content_response.content

    step_refs = [
        ("Create project", [project_id]),
        ("Authorize charter", [charter["id"]]),
        ("Create WBS", [root["id"], *[item["id"] for item in packages]]),
        ("Import and validate BOQ", [source["id"]]),
        ("Generate proposed schedule", [draft["id"], calculation["id"]]),
        (
            "Reconcile BOQ/WBS/activity allocations",
            [a["boq_line_id"] for _, item in bindings for a in item["boq_allocations"]],
        ),
        ("Authorize schedule and baseline", [schedule["id"], baseline["id"]]),
        (
            "Record daily update and governed progress",
            [daily_record["id"], *[m["id"] for m in measurements]],
        ),
        ("Record separate cost ledgers and value bases", cost_ids),
        (
            "Verify delay and assess its schedule exposure",
            [delay["id"], schedule_assessment["assessment"]["id"]],
        ),
        ("Freeze project-control snapshot", [control["id"]]),
        ("Explain package and project health", [control["id"]]),
        ("Query frozen-snapshot assistant", [control["id"]]),
        ("Produce grounded management considerations", [control["id"]]),
        ("Generate and approve report", [report["id"]]),
        ("Generate and approve stakeholder presentation", [deck["id"], recipient["id"]]),
        (
            "Prove cross-artifact equality and later-update isolation",
            [control["id"], report["id"], deck["id"]],
        ),
    ]

    manifest = {
        "scope": (
            "Synthetic public-API acceptance; not a representative human session or deployment"
        ),
        "project_id": project_id,
        "boq_source_id": source["id"],
        "boq_artifact_id": upload.json()["id"],
        "work_packages": [
            {"id": item["id"], "code": item["code"], "name": item["name"]} for item in packages
        ],
        "dependency_count": package_count - 1,
        "control_snapshot_id": control["id"],
        "report_id": report["id"],
        "presentation_id": deck["id"],
        "schedule_context_id": schedule["id"],
        "baseline_context_id": schedule["source_version_id"],
        "report_hash": report["content_hash"],
        "presentation_hash": deck["content_hash"],
        "assistant_hash": assistant["answer_hash"],
        "steps": [
            {"step": index, "description": description, "record_ids": refs, "status": "PASS"}
            for index, (description, refs) in enumerate(step_refs, start=1)
        ],
        "commands": commands,
    }
    artifact_dir.mkdir(parents=True, exist_ok=True)
    for name, value in (
        ("manifest", manifest),
        ("control-snapshot", control),
        ("report", report_detail),
        ("presentation", deck_detail),
        ("assistant", assistant),
    ):
        (artifact_dir / f"connected-{name}.json").write_text(
            json.dumps(value, indent=2, sort_keys=True), encoding="utf-8"
        )
    (artifact_dir / "connected-presentation.html").write_bytes(content_response.content)
    return manifest


@pytest.mark.parametrize("package_count", [1, 5])
def test_document_19_connected_governed_acceptance(evidence_api, tmp_path, package_count):
    exercise_connected_acceptance(evidence_api, tmp_path, package_count)
