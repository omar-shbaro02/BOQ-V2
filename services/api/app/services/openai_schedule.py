"""OpenAI-backed BOQ-to-schedule proposal with strict local validation."""

from __future__ import annotations

import io
import json
import uuid
from dataclasses import dataclass
from datetime import date, timedelta
from typing import Any

import httpx
from app.services.ms_project_export import HEADERS, PROGRESS_HEADERS, BoqItem, _safe_text
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

RESPONSES_URL = "https://api.openai.com/v1/responses"


def _schedule_schema() -> dict[str, Any]:
    relationship = {
        "type": "object",
        "properties": {
            "predecessor_id": {"type": "integer", "minimum": 1},
            "type": {"type": "string", "enum": ["FS", "SS", "FF", "SF"]},
            "lag_days": {"type": "integer", "minimum": -30, "maximum": 365},
            "basis": {"type": "string", "minLength": 3},
        },
        "required": ["predecessor_id", "type", "lag_days", "basis"],
        "additionalProperties": False,
    }
    task = {
        "type": "object",
        "properties": {
            "id": {"type": "integer", "minimum": 1},
            "name": {"type": "string", "minLength": 3},
            "duration_days": {"type": "integer", "minimum": 1},
            "start": {"type": "string"},
            "finish": {"type": "string"},
            "relationships": {"type": "array", "items": relationship},
            "boq_references": {"type": "array", "items": {"type": "string"}},
            "scope_type": {
                "type": "string",
                "enum": ["BOQ_REQUIREMENT", "EXECUTION_SUPPORT"],
            },
            "duration_basis": {
                "type": "string",
                "enum": [
                    "BOQ_QUANTITY_AND_VERIFIED_RATE",
                    "DOCUMENTED_PROJECT_INPUT",
                    "PROVISIONAL_PLANNING_ASSUMPTION",
                ],
            },
            "duration_note": {"type": "string", "minLength": 3},
        },
        "required": [
            "id",
            "name",
            "duration_days",
            "start",
            "finish",
            "relationships",
            "boq_references",
            "scope_type",
            "duration_basis",
            "duration_note",
        ],
        "additionalProperties": False,
    }
    return {
        "type": "object",
        "properties": {"tasks": {"type": "array", "minItems": 5, "items": task}},
        "required": ["tasks"],
        "additionalProperties": False,
    }


def _boq_prompt(
    items: list[BoqItem],
    project_start: date,
    target_finish: date | None,
    working_days: tuple[int, ...],
) -> str:
    rows = "\n".join(
        f"{item.reference} | {' > '.join(item.hierarchy) or 'Unclassified'} | "
        f"{item.number} | {item.description} | {item.unit} | "
        f"{item.quantity if item.quantity is not None else 'unknown'}"
        for item in items
    )
    finish_instruction = (
        f"The required completion date is {target_finish.isoformat()}."
        if target_finish
        else "Propose a realistic completion date from the scope and sequencing."
    )
    day_names = ("Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday")
    calendar = ", ".join(day_names[index] for index in working_days)
    return f"""Create a construction time-schedule proposal from this BOQ.
Project start: {project_start.isoformat()}. {finish_instruction}
Working days: {calendar}. No holiday or shutdown calendar was supplied.

Return executable work packages, not BOQ chapters or one task per commercial line. Every BOQ
reference must appear exactly once across BOQ_REQUIREMENT tasks. Use EXECUTION_SUPPORT only for
necessary approvals, procurement, curing, inspection, testing, access, commissioning or handover.
Never invent verified productivity, crews, prices, lead times, holidays or approvals. Without an
explicit verified rate, label duration_basis PROVISIONAL_PLANNING_ASSUMPTION and explain it.

Apply physical construction prerequisites: excavation, blinding, structure, curing/inspection,
below-grade waterproofing, protection/inspection, then backfill; underground-service testing before
closure; structure clearance, blockwork, MEP first fix, concealed-work inspection, then plaster or
ceiling closure; roof substrate, waterproofing, testing and protection; wet-area first fix,
inspection, screed, waterproofing/testing, finishes, then final fix. Separate significant equipment
submittal/approval, purchase, fabrication/delivery, installation and commissioning when evidenced.
Use FS/SS/FF/SF relationships only with a physical basis. SS/FF overlap requires staged area or
handover justification. Predecessors must be earlier task IDs. Dates must use the supplied working
days and be ISO YYYY-MM-DD. This remains a proposed schedule requiring planner review.

BOQ measured scope:
{rows}"""


async def create_openai_schedule(
    items: list[BoqItem],
    *,
    api_key: str,
    model: str,
    project_start: date,
    target_finish: date | None,
    working_days: tuple[int, ...] = (0, 1, 2, 3, 4, 5),
) -> list[dict[str, Any]]:
    body = {
        "model": model,
        "instructions": (
            "You are a senior construction planner. Transform BOQ evidence into a concise, "
            "logically linked schedule proposal. Respect the requested structured schema exactly."
        ),
        "input": _boq_prompt(items, project_start, target_finish, working_days),
        "reasoning": {"effort": "minimal"},
        "text": {
            "verbosity": "low",
            "format": {
                "type": "json_schema",
                "name": "boq_schedule",
                "strict": True,
                "schema": _schedule_schema(),
            },
        },
        "max_output_tokens": 16000,
    }
    try:
        async with httpx.AsyncClient(timeout=300) as client:
            response = await client.post(
                RESPONSES_URL,
                headers={"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"},
                json=body,
            )
    except httpx.TimeoutException as exc:
        raise ValueError(
            "Schedule generation took longer than five minutes. The BOQ remains saved; "
            "try generating it again."
        ) from exc
    except httpx.RequestError as exc:
        raise ValueError(
            "The scheduling service could not be reached. The BOQ remains saved; "
            "check the internet connection and try again."
        ) from exc
    if response.status_code >= 400:
        detail = response.json().get("error", {}).get("message", response.text[:500])
        raise ValueError(f"OpenAI schedule generation failed: {detail}")
    payload = response.json()
    if payload.get("status") == "incomplete":
        reason = payload.get("incomplete_details", {}).get("reason", "output limit")
        raise ValueError(
            f"OpenAI returned an incomplete schedule ({reason}); no workbook was created"
        )
    output_text = next(
        (
            content.get("text")
            for item in payload.get("output", [])
            for content in item.get("content", [])
            if content.get("type") == "output_text"
        ),
        None,
    )
    if not output_text:
        raise ValueError("OpenAI returned no structured schedule")
    try:
        tasks = json.loads(output_text)["tasks"]
    except json.JSONDecodeError as exc:
        raise ValueError("OpenAI returned an incomplete schedule; no workbook was created") from exc
    _validate_tasks(tasks, items, project_start, target_finish, working_days)
    return tasks


def _validate_tasks(
    tasks: list[dict[str, Any]],
    items: list[BoqItem],
    project_start: date,
    target_finish: date | None,
    working_days: tuple[int, ...],
) -> None:
    ids = [task["id"] for task in tasks]
    if ids != list(range(1, len(tasks) + 1)):
        raise ValueError("OpenAI schedule IDs must be consecutive from 1")
    for task in tasks:
        date.fromisoformat(task["start"])
        date.fromisoformat(task["finish"])


@dataclass(frozen=True)
class ScheduleFinding:
    finding_id: str
    severity: str
    activity_ids: tuple[int, ...]
    issue: str
    violated_rule: str
    proposed_correction: str


def _dependency_text(task: dict[str, Any]) -> str | None:
    values = []
    for relationship in task["relationships"]:
        lag = relationship["lag_days"]
        lag_text = f"{lag:+d}d" if lag else ""
        values.append(f"{relationship['predecessor_id']}{relationship['type']}{lag_text}")
    return ",".join(values) or None


def _transitive_predecessors(tasks: list[dict[str, Any]], task_id: int) -> set[int]:
    by_id = {task["id"]: task for task in tasks}
    found: set[int] = set()
    pending = [relationship["predecessor_id"] for relationship in by_id[task_id]["relationships"]]
    while pending:
        predecessor = pending.pop()
        if predecessor in found:
            continue
        found.add(predecessor)
        if predecessor in by_id:
            pending.extend(
                relationship["predecessor_id"]
                for relationship in by_id[predecessor]["relationships"]
            )
    return found


def validate_schedule_proposal(
    tasks: list[dict[str, Any]],
    items: list[BoqItem],
    working_days: tuple[int, ...] = (0, 1, 2, 3, 4, 5),
    project_start: date | None = None,
    target_finish: date | None = None,
) -> tuple[str, list[ScheduleFinding], list[dict[str, str]]]:
    findings: list[ScheduleFinding] = []
    allocations: dict[str, list[int]] = {item.reference: [] for item in items}
    for task in tasks:
        start = date.fromisoformat(task["start"])
        finish = date.fromisoformat(task["finish"])
        if finish < start or (project_start and start < project_start):
            findings.append(
                ScheduleFinding(
                    f"CAL-{len(findings) + 1:04d}",
                    "CRITICAL",
                    (task["id"],),
                    "Activity dates fall outside a valid forward schedule",
                    "Finish must follow start and work cannot precede the project start",
                    "Correct the dates and recalculate the proposal.",
                )
            )
        if target_finish and finish > target_finish:
            findings.append(
                ScheduleFinding(
                    f"CAL-{len(findings) + 1:04d}",
                    "CRITICAL",
                    (task["id"],),
                    "Activity exceeds the required completion date",
                    "The proposed schedule must expose deadline conflicts",
                    "Revise logic or durations, or obtain authorized acceptance of the conflict.",
                )
            )
        if start.weekday() not in working_days or finish.weekday() not in working_days:
            findings.append(
                ScheduleFinding(
                    f"CAL-{len(findings) + 1:04d}",
                    "CRITICAL",
                    (task["id"],),
                    "Activity starts or finishes on a non-working day",
                    "All dates and durations must use the selected project calendar",
                    "Recalculate the activity against the reviewed calendar.",
                )
            )
        for reference in task["boq_references"]:
            if reference in allocations:
                allocations[reference].append(task["id"])
            else:
                findings.append(
                    ScheduleFinding(
                        f"COV-{len(findings) + 1:04d}",
                        "CRITICAL",
                        (task["id"],),
                        f"Activity cites unknown BOQ reference {reference}",
                        "Activity evidence must resolve to an extracted BOQ item",
                        "Replace the reference with exact source evidence or remove the claim.",
                    )
                )

    for reference, activity_ids in allocations.items():
        if not activity_ids:
            findings.append(
                ScheduleFinding(
                    f"COV-{len(findings) + 1:04d}",
                    "CRITICAL",
                    (),
                    f"BOQ item {reference} is not allocated to an activity",
                    "Every measurable BOQ item requires an accountable disposition",
                    "Map the BOQ item to an executable activity or record a reviewed exclusion.",
                )
            )
        elif len(activity_ids) > 1:
            findings.append(
                ScheduleFinding(
                    f"COV-{len(findings) + 1:04d}",
                    "HIGH",
                    tuple(activity_ids),
                    f"BOQ item {reference} is allocated more than once",
                    "Measured quantities cannot be duplicated without an explicit allocation",
                    "Partition the quantity explicitly or retain one accountable activity mapping.",
                )
            )

    names = {task["id"]: task["name"].casefold() for task in tasks}
    normalized_names: dict[str, list[int]] = {}
    for task in tasks:
        normalized = " ".join(task["name"].casefold().split())
        normalized_names.setdefault(normalized, []).append(task["id"])
    for name, activity_ids in normalized_names.items():
        if len(activity_ids) > 1:
            findings.append(
                ScheduleFinding(
                    f"SCP-{len(findings) + 1:04d}",
                    "HIGH",
                    tuple(activity_ids),
                    f"Duplicate activity name: {name}",
                    "Duplicate scope must be reconciled before approval",
                    "Clarify distinct scope or consolidate the duplicated activity.",
                )
            )

    scope_rules = (
        (("boiler", "radiator", "heating"), ("boiler", "radiator", "heating"), "Heating system"),
        (("solar water", "solar hot"), ("solar water", "solar hot"), "Solar hot-water system"),
        (("ups", "uninterruptible"), ("ups", "uninterruptible"), "UPS system"),
        (("road base", "base course"), ("road base", "base course"), "Road-base course"),
        (("asphalt",), ("asphalt",), "Asphalt courses"),
        (("curb", "kerb"), ("curb", "kerb"), "Curbs"),
        (("lift", "elevator"), ("lift", "elevator"), "Lift system"),
        (("fire pump",), ("fire pump",), "Fire-pump package"),
        (("photovoltaic", " pv "), ("photovoltaic", " pv "), "Photovoltaic system"),
    )
    boq_text = " ".join(item.description.casefold() for item in items)
    task_text = " ".join(f" {name} " for name in names.values())
    for evidence_terms, activity_terms, label in scope_rules:
        if any(term in boq_text for term in evidence_terms) and not any(
            term in task_text for term in activity_terms
        ):
            findings.append(
                ScheduleFinding(
                    f"SCP-{len(findings) + 1:04d}",
                    "CRITICAL",
                    (),
                    f"{label} is present in the BOQ but lacks a distinct activity",
                    "Material BOQ systems require traceable schedule coverage",
                    f"Add traceable procurement, installation and testing coverage for {label}.",
                )
            )
    task_by_id = {task["id"]: task for task in tasks}
    for task in tasks:
        for relationship in task["relationships"]:
            predecessor_id = relationship["predecessor_id"]
            if predecessor_id not in task_by_id or predecessor_id >= task["id"]:
                findings.append(
                    ScheduleFinding(
                        f"LOG-{len(findings) + 1:04d}",
                        "CRITICAL",
                        (predecessor_id, task["id"]),
                        "Dependency references a missing, same or later activity",
                        "Predecessors must exist and produce an acyclic graph",
                        "Correct the predecessor reference and recalculate the schedule.",
                    )
                )
                continue
            predecessor = task_by_id[predecessor_id]
            lag = timedelta(days=relationship["lag_days"])
            predecessor_start = date.fromisoformat(predecessor["start"])
            predecessor_finish = date.fromisoformat(predecessor["finish"])
            successor_start = date.fromisoformat(task["start"])
            successor_finish = date.fromisoformat(task["finish"])
            relationship_type = relationship["type"]
            violated = (
                (relationship_type == "FS" and successor_start < predecessor_finish + lag)
                or (relationship_type == "SS" and successor_start < predecessor_start + lag)
                or (relationship_type == "FF" and successor_finish < predecessor_finish + lag)
                or (relationship_type == "SF" and successor_finish < predecessor_start + lag)
            )
            if violated:
                findings.append(
                    ScheduleFinding(
                        f"LOG-{len(findings) + 1:04d}",
                        "CRITICAL",
                        (predecessor_id, task["id"]),
                        f"{relationship_type} relationship is violated by activity dates",
                        "Dates must satisfy relationship type and lag",
                        "Recalculate both activities against the reviewed calendar and logic.",
                    )
                )
    waterproof = [
        task_id
        for task_id, name in names.items()
        if "waterproof" in name
        and any(word in name for word in ("below", "foundation", "basement"))
    ]
    backfill = [task_id for task_id, name in names.items() if "backfill" in name]
    for backfill_id in backfill:
        predecessors = _transitive_predecessors(tasks, backfill_id)
        if waterproof and not any(task_id in predecessors for task_id in waterproof):
            findings.append(
                ScheduleFinding(
                    f"SCH-{len(findings) + 1:04d}",
                    "CRITICAL",
                    tuple(waterproof + [backfill_id]),
                    "Below-grade waterproofing is not a predecessor of backfilling",
                    "Waterproofing, inspection and protection must precede affected backfill",
                    "Add separate below-grade waterproofing and inspection/protection "
                    "gates before backfill.",
                )
            )

    ceiling_closures = [
        task_id
        for task_id, name in names.items()
        if "ceiling" in name and any(word in name for word in ("close", "closure", "board"))
    ]
    concealed = [
        task_id
        for task_id, name in names.items()
        if "concealed" in name and any(word in name for word in ("inspect", "test", "service"))
    ]
    for closure_id in ceiling_closures:
        if concealed and not any(
            task_id in _transitive_predecessors(tasks, closure_id) for task_id in concealed
        ):
            findings.append(
                ScheduleFinding(
                    f"SCH-{len(findings) + 1:04d}",
                    "CRITICAL",
                    tuple(concealed + [closure_id]),
                    "Ceiling closure lacks a concealed-services inspection predecessor",
                    "Concealed services must be completed, tested and inspected before closure",
                    "Add a concealed-services completion and inspection gate before "
                    "ceiling closure.",
                )
            )

    for task in tasks:
        if task["duration_basis"] == "PROVISIONAL_PLANNING_ASSUMPTION":
            findings.append(
                ScheduleFinding(
                    f"ASM-{len(findings) + 1:04d}",
                    "WARNING",
                    (task["id"],),
                    "Duration uses an unverified planning assumption",
                    "Durations require documented productivity or project input",
                    "Confirm the duration, crew/productivity basis and calendar during "
                    "planner review.",
                )
            )

    coverage = [
        {
            "boq_ref": item.reference,
            "description": item.description,
            "hierarchy": " > ".join(item.hierarchy),
            "quantity": str(item.quantity) if item.quantity is not None else "",
            "unit": item.unit,
            "mapped_activity_ids": ",".join(str(value) for value in allocations[item.reference]),
            "coverage_status": "Covered"
            if len(allocations[item.reference]) == 1
            else ("Missing" if not allocations[item.reference] else "Requires Clarification"),
            "source_page": str(item.page),
        }
        for item in items
    ]
    status = (
        "BLOCKED"
        if any(item.severity == "CRITICAL" for item in findings)
        else ("PASS WITH WARNINGS" if findings else "PASS")
    )
    return status, findings, coverage


def build_openai_schedule_workbook(
    tasks: list[dict[str, Any]],
    *,
    coverage: list[dict[str, str]] | None = None,
    findings: list[ScheduleFinding] | None = None,
    qa_status: str = "PROPOSED",
) -> bytes:
    """Create the application-owned MS Project import contract.

    Customers supply project data, never an executable workbook template. Keeping
    the six-column projection here makes the export identical across tenants.
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Tasks"
    sheet.append(HEADERS)
    for task in tasks:
        sheet.append(
            (
                task["id"],
                _safe_text(task["name"]),
                f"{task['duration_days']} days",
                task["start"],
                task["finish"],
                _dependency_text(task)
                if "relationships" in task
                else ",".join(str(value) for value in task.get("dependency_ids", [])) or None,
            )
        )
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:F{sheet.max_row}"
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="6E1E36")
        cell.alignment = Alignment(horizontal="center")
    for column, width in {"A": 10, "B": 62, "C": 16, "D": 18, "E": 18, "F": 20}.items():
        sheet.column_dimensions[column].width = width

    progress = workbook.create_sheet("Progress Report", 0)
    progress.append(
        (
            "PROGRESS INPUT",
            "Update Work Complete %, Reporting Date, Remarks, Evidence Reference and "
            "Responsible Party. Upload this workbook on the Progress page for analysis.",
        )
    )
    progress.append(
        (
            "Evidence boundary",
            "Entered percentages are reported information until their evidence is reviewed.",
        )
    )
    progress.append(())
    progress.append(PROGRESS_HEADERS)
    for task in tasks:
        references = task.get("boq_references", [])
        progress.append(
            (
                ", ".join(references) or "EXECUTION SUPPORT",
                task["id"],
                _safe_text(task["name"]),
                f"{task['duration_days']} days",
                task["start"],
                task["finish"],
                None,
                None,
                None,
                None,
                None,
                task.get("duration_basis", "PROVISIONAL_PLANNING_ASSUMPTION"),
            )
        )
    percent_validation = DataValidation(
        type="decimal", operator="between", formula1="0", formula2="100", allow_blank=True
    )
    percent_validation.error = "Enter a percentage from 0 to 100."
    percent_validation.errorTitle = "Invalid progress percentage"
    progress.add_data_validation(percent_validation)
    percent_validation.add(f"G5:G{progress.max_row}")
    progress.conditional_formatting.add(
        f"G5:G{progress.max_row}",
        CellIsRule(
            operator="greaterThan", formula=["100"], fill=PatternFill("solid", fgColor="FFC7CE")
        ),
    )
    progress.freeze_panes = "A5"
    progress.auto_filter.ref = f"A4:L{progress.max_row}"
    progress.sheet_view.showGridLines = False
    for cell in progress[4]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="6E1E36")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for column, width in {
        "A": 28,
        "B": 10,
        "C": 52,
        "D": 18,
        "E": 16,
        "F": 16,
        "G": 18,
        "H": 18,
        "I": 36,
        "J": 30,
        "K": 24,
        "L": 34,
    }.items():
        progress.column_dimensions[column].width = width
    for column in ("G", "H", "I", "J", "K"):
        for row in range(5, progress.max_row + 1):
            progress[f"{column}{row}"].fill = PatternFill("solid", fgColor="FFF2CC")
    if coverage is not None:
        review = workbook.create_sheet("BOQ Coverage")
        review.append(
            (
                "BOQ Reference",
                "Hierarchy",
                "Description",
                "Quantity",
                "Unit",
                "Mapped Activity IDs",
                "Coverage Status",
                "Source Page",
            )
        )
        for item in coverage:
            review.append(
                tuple(
                    item[key]
                    for key in (
                        "boq_ref",
                        "hierarchy",
                        "description",
                        "quantity",
                        "unit",
                        "mapped_activity_ids",
                        "coverage_status",
                        "source_page",
                    )
                )
            )
        review.freeze_panes = "A2"
        review.auto_filter.ref = f"A1:H{review.max_row}"
        qa = workbook.create_sheet("QA Review")
        qa.append(("QA Status", qa_status))
        qa.append(
            (
                "Finding ID",
                "Severity",
                "Activity IDs",
                "Issue",
                "Violated Rule",
                "Proposed Correction",
                "Human Approval Required",
            )
        )
        for finding in findings or []:
            qa.append(
                (
                    finding.finding_id,
                    finding.severity,
                    ",".join(str(value) for value in finding.activity_ids),
                    finding.issue,
                    finding.violated_rule,
                    finding.proposed_correction,
                    "Yes",
                )
            )
        for target in (review, qa):
            for cell in target[1 if target is review else 2]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="6E1E36")
            target.freeze_panes = "A2" if target is review else "A3"
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def generation_id() -> str:
    return str(uuid.uuid4())
