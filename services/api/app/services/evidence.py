from __future__ import annotations

import csv
import io
import json
import uuid
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from app.auth import ActorContext
from app.evidence_schemas import (
    ContradictionCreate,
    ContradictionResolve,
    EvidenceItemCreate,
    EvidenceRelationCreate,
    EvidenceRequestCreate,
    ReliabilityCreate,
    VerificationCreate,
)
from app.generated.taxonomies import (
    ContradictionStatus,
    EvidenceItemStatus,
    EvidenceRelationType,
    EvidenceRequestStatus,
    ImportBatchStatus,
    SemanticState,
    TruthType,
    VerificationOutcome,
)
from app.models import (
    Contradiction,
    ControlledObject,
    EvidenceArtifact,
    EvidenceItem,
    EvidenceRelation,
    EvidenceRequest,
    ImportBatch,
    ImportBatchItem,
    Project,
    SourceReliabilityAssessment,
    VerificationEvent,
)
from app.services.audit import record_audit
from app.storage import EvidenceStore
from fastapi import HTTPException
from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy import select
from sqlalchemy.orm import Session

CANONICAL_IMPORT_FIELDS = (
    "field_name",
    "value",
    "semantic_state",
    "truth_type",
    "as_of",
    "object_code",
    "unit",
    "currency",
    "measurement_basis",
    "confidence",
    "source_reliability",
    "expires_at",
)
REQUIRED_IMPORT_FIELDS = {"field_name", "value", "semantic_state", "truth_type", "as_of"}


def scoped_artifact(session: Session, project: Project, artifact_id: uuid.UUID) -> EvidenceArtifact:
    artifact = session.get(EvidenceArtifact, artifact_id)
    if artifact is None or artifact.project_id != project.id:
        raise HTTPException(status_code=404, detail="Evidence artifact not found")
    return artifact


def scoped_item(session: Session, project: Project, item_id: uuid.UUID) -> EvidenceItem:
    item = session.get(EvidenceItem, item_id)
    if item is None or item.project_id != project.id:
        raise HTTPException(status_code=404, detail="Evidence item not found")
    return item


def create_evidence_item(
    session: Session,
    actor: ActorContext,
    project: Project,
    data: EvidenceItemCreate,
) -> EvidenceItem:
    if data.artifact_id:
        scoped_artifact(session, project, data.artifact_id)
    if data.controlled_object_id:
        controlled_object = session.get(ControlledObject, data.controlled_object_id)
        if controlled_object is None or controlled_object.project_id != project.id:
            raise HTTPException(status_code=422, detail="Controlled object must belong to project")
    if data.supersedes_item_id:
        prior = scoped_item(session, project, data.supersedes_item_id)
        if prior.status != EvidenceItemStatus.ACTIVE:
            raise HTTPException(status_code=409, detail="Only active evidence can be superseded")
        if (
            prior.field_name != data.field_name
            or prior.controlled_object_id != data.controlled_object_id
        ):
            raise HTTPException(
                status_code=422,
                detail="Superseding evidence must preserve field and controlled-object identity",
            )
        prior.status = EvidenceItemStatus.SUPERSEDED

    item = EvidenceItem(
        id=uuid.uuid4(),
        organization_id=project.organization_id,
        project_id=project.id,
        created_by=actor.actor_id,
        status=EvidenceItemStatus.ACTIVE,
        derived_from_item_id=None,
        **data.model_dump(),
    )
    session.add(item)
    record_audit(
        session,
        actor,
        organization_id=project.organization_id,
        project_id=project.id,
        action="EVIDENCE_ITEM_CREATED",
        object_type="EVIDENCE_ITEM",
        object_id=str(item.id),
        details={
            "field_name": item.field_name,
            "semantic_state": item.semantic_state,
            "truth_type": item.truth_type,
            "supersedes_item_id": str(item.supersedes_item_id) if item.supersedes_item_id else None,
        },
    )
    return item


def create_relation(
    session: Session,
    actor: ActorContext,
    project: Project,
    data: EvidenceRelationCreate,
) -> EvidenceRelation:
    scoped_item(session, project, data.from_item_id)
    scoped_item(session, project, data.to_item_id)
    relation = EvidenceRelation(
        id=uuid.uuid4(),
        organization_id=project.organization_id,
        project_id=project.id,
        created_by=actor.actor_id,
        **data.model_dump(),
    )
    session.add(relation)
    record_audit(
        session,
        actor,
        organization_id=project.organization_id,
        project_id=project.id,
        action="EVIDENCE_RELATION_CREATED",
        object_type="EVIDENCE_RELATION",
        object_id=str(relation.id),
        details={"relation_type": relation.relation_type},
    )
    return relation


def verify_item(
    session: Session,
    actor: ActorContext,
    project: Project,
    item_id: uuid.UUID,
    data: VerificationCreate,
) -> VerificationEvent:
    source = scoped_item(session, project, item_id)
    result: EvidenceItem | None = None
    if data.outcome == VerificationOutcome.VERIFIED:
        result = EvidenceItem(
            id=uuid.uuid4(),
            organization_id=project.organization_id,
            project_id=project.id,
            artifact_id=source.artifact_id,
            controlled_object_id=source.controlled_object_id,
            field_name=source.field_name,
            value=source.value,
            unit=source.unit,
            currency=source.currency,
            measurement_basis=source.measurement_basis,
            semantic_state=SemanticState.VERIFIED,
            truth_type=TruthType.VERIFIED_FACT,
            as_of=source.as_of,
            expires_at=source.expires_at,
            confidence=data.verified_confidence,
            source_reliability=source.source_reliability,
            status=EvidenceItemStatus.ACTIVE,
            supersedes_item_id=None,
            derived_from_item_id=source.id,
            created_by=actor.actor_id,
        )
        session.add(result)
        session.add(
            EvidenceRelation(
                id=uuid.uuid4(),
                organization_id=project.organization_id,
                project_id=project.id,
                from_item_id=result.id,
                to_item_id=source.id,
                relation_type=EvidenceRelationType.DERIVED_FROM,
                rationale=f"Verified using {data.method}",
                created_by=actor.actor_id,
            )
        )

    event = VerificationEvent(
        id=uuid.uuid4(),
        organization_id=project.organization_id,
        project_id=project.id,
        evidence_item_id=source.id,
        result_item_id=result.id if result else None,
        method=data.method,
        outcome=data.outcome,
        rationale=data.rationale,
        reviewer_actor_id=actor.actor_id,
    )
    session.add(event)
    record_audit(
        session,
        actor,
        organization_id=project.organization_id,
        project_id=project.id,
        action="EVIDENCE_VERIFICATION_RECORDED",
        object_type="VERIFICATION_EVENT",
        object_id=str(event.id),
        details={
            "source_item_id": str(source.id),
            "result_item_id": str(result.id) if result else None,
            "outcome": data.outcome,
        },
    )
    return event


def assess_reliability(
    session: Session,
    actor: ActorContext,
    project: Project,
    data: ReliabilityCreate,
) -> SourceReliabilityAssessment:
    assessment = SourceReliabilityAssessment(
        id=uuid.uuid4(),
        organization_id=project.organization_id,
        project_id=project.id,
        assessed_by=actor.actor_id,
        **data.model_dump(),
    )
    session.add(assessment)
    record_audit(
        session,
        actor,
        organization_id=project.organization_id,
        project_id=project.id,
        action="SOURCE_RELIABILITY_ASSESSED",
        object_type="SOURCE_RELIABILITY_ASSESSMENT",
        object_id=str(assessment.id),
        details={"source_id": data.source_id, "score": str(data.score)},
    )
    return assessment


def create_contradiction(
    session: Session,
    actor: ActorContext,
    project: Project,
    data: ContradictionCreate,
) -> Contradiction:
    left = scoped_item(session, project, data.left_item_id)
    right = scoped_item(session, project, data.right_item_id)
    if left.field_name != data.field_name or right.field_name != data.field_name:
        raise HTTPException(status_code=422, detail="Contradiction field must match both items")
    if left.controlled_object_id != right.controlled_object_id:
        raise HTTPException(
            status_code=422, detail="Contradicting items must address the same object"
        )
    contradiction = Contradiction(
        id=uuid.uuid4(),
        organization_id=project.organization_id,
        project_id=project.id,
        controlled_object_id=left.controlled_object_id,
        status=ContradictionStatus.OPEN,
        created_by=actor.actor_id,
        **data.model_dump(),
    )
    session.add(contradiction)
    session.add(
        EvidenceRelation(
            id=uuid.uuid4(),
            organization_id=project.organization_id,
            project_id=project.id,
            from_item_id=left.id,
            to_item_id=right.id,
            relation_type=EvidenceRelationType.CONTRADICTS,
            rationale=f"Contradiction {contradiction.id}",
            created_by=actor.actor_id,
        )
    )
    record_audit(
        session,
        actor,
        organization_id=project.organization_id,
        project_id=project.id,
        action="CONTRADICTION_DETECTED",
        object_type="CONTRADICTION",
        object_id=str(contradiction.id),
        details={"material": data.material, "field_name": data.field_name},
    )
    return contradiction


def resolve_contradiction(
    session: Session,
    actor: ActorContext,
    project: Project,
    contradiction_id: uuid.UUID,
    data: ContradictionResolve,
) -> Contradiction:
    contradiction = session.get(Contradiction, contradiction_id)
    if contradiction is None or contradiction.project_id != project.id:
        raise HTTPException(status_code=404, detail="Contradiction not found")
    if contradiction.status == ContradictionStatus.RESOLVED:
        raise HTTPException(status_code=409, detail="Contradiction is already resolved")
    if data.chosen_item_id and data.chosen_item_id not in {
        contradiction.left_item_id,
        contradiction.right_item_id,
    }:
        raise HTTPException(status_code=422, detail="Chosen item must be part of contradiction")
    contradiction.status = ContradictionStatus.RESOLVED
    contradiction.chosen_item_id = data.chosen_item_id
    contradiction.resolution_policy = data.resolution_policy
    contradiction.resolution_reason = data.resolution_reason
    contradiction.resolved_by = actor.actor_id
    contradiction.resolved_at = datetime.now(UTC)
    record_audit(
        session,
        actor,
        organization_id=project.organization_id,
        project_id=project.id,
        action="CONTRADICTION_RESOLVED",
        object_type="CONTRADICTION",
        object_id=str(contradiction.id),
        details={
            "chosen_item_id": str(data.chosen_item_id) if data.chosen_item_id else None,
            "policy": data.resolution_policy,
        },
    )
    return contradiction


def create_evidence_request(
    session: Session,
    actor: ActorContext,
    project: Project,
    data: EvidenceRequestCreate,
) -> EvidenceRequest:
    if data.controlled_object_id:
        controlled_object = session.get(ControlledObject, data.controlled_object_id)
        if controlled_object is None or controlled_object.project_id != project.id:
            raise HTTPException(status_code=422, detail="Controlled object must belong to project")
    request = EvidenceRequest(
        id=uuid.uuid4(),
        organization_id=project.organization_id,
        project_id=project.id,
        status=EvidenceRequestStatus.OPEN,
        satisfied_by_item_id=None,
        created_by=actor.actor_id,
        **data.model_dump(),
    )
    session.add(request)
    record_audit(
        session,
        actor,
        organization_id=project.organization_id,
        project_id=project.id,
        action="EVIDENCE_REQUESTED",
        object_type="EVIDENCE_REQUEST",
        object_id=str(request.id),
        details={"requested_fields": request.requested_fields, "urgency": request.urgency},
    )
    return request


def satisfy_evidence_request(
    session: Session,
    actor: ActorContext,
    project: Project,
    request_id: uuid.UUID,
    item_id: uuid.UUID,
) -> EvidenceRequest:
    request = session.get(EvidenceRequest, request_id)
    if request is None or request.project_id != project.id:
        raise HTTPException(status_code=404, detail="Evidence request not found")
    if request.status != EvidenceRequestStatus.OPEN:
        raise HTTPException(status_code=409, detail="Evidence request is not open")
    item = scoped_item(session, project, item_id)
    if item.field_name not in request.requested_fields:
        raise HTTPException(status_code=422, detail="Evidence field does not satisfy request")
    if request.controlled_object_id and item.controlled_object_id != request.controlled_object_id:
        raise HTTPException(status_code=422, detail="Evidence object does not satisfy request")
    request.status = EvidenceRequestStatus.SATISFIED
    request.satisfied_by_item_id = item.id
    record_audit(
        session,
        actor,
        organization_id=project.organization_id,
        project_id=project.id,
        action="EVIDENCE_REQUEST_SATISFIED",
        object_type="EVIDENCE_REQUEST",
        object_id=str(request.id),
        details={"evidence_item_id": str(item.id)},
    )
    return request


def read_tabular_artifact(artifact: EvidenceArtifact, store: EvidenceStore) -> list[dict[str, Any]]:
    content = store.get(artifact.storage_key)
    suffix = Path(artifact.original_filename).suffix.lower()
    if suffix == ".csv":
        text = content.decode("utf-8-sig")
        return [dict(row) for row in csv.DictReader(io.StringIO(text))]
    if suffix in {".xlsx", ".xlsm"}:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        headers = [str(value).strip() if value is not None else "" for value in next(rows, ())]
        return [dict(zip(headers, values, strict=True)) for values in rows]
    raise HTTPException(status_code=422, detail="Import artifact must be CSV or XLSX")


def parse_value(raw: Any) -> Any:
    if not isinstance(raw, str):
        return raw
    stripped = raw.strip()
    try:
        return json.loads(stripped)
    except json.JSONDecodeError:
        return stripped


def preview_import(
    session: Session,
    actor: ActorContext,
    project: Project,
    artifact: EvidenceArtifact,
    mapping: dict[str, str],
    idempotency_key: str,
    store: EvidenceStore,
) -> ImportBatch:
    existing = session.scalar(
        select(ImportBatch).where(
            ImportBatch.project_id == project.id,
            ImportBatch.idempotency_key == idempotency_key,
        )
    )
    if existing:
        if existing.artifact_id != artifact.id:
            raise HTTPException(
                status_code=409, detail="Idempotency key belongs to another artifact"
            )
        return existing

    rows = read_tabular_artifact(artifact, store)
    effective_mapping = {field: mapping.get(field, field) for field in CANONICAL_IMPORT_FIELDS}
    missing_headers = REQUIRED_IMPORT_FIELDS - {
        field
        for field, source_header in effective_mapping.items()
        if rows and source_header in rows[0]
    }
    errors: list[dict[str, Any]] = []
    normalized: list[dict[str, Any]] = []
    if missing_headers:
        errors.append({"row": 0, "message": f"Missing mapped headers: {sorted(missing_headers)}"})

    object_by_code = {
        obj.code: obj.id
        for obj in session.scalars(
            select(ControlledObject).where(ControlledObject.project_id == project.id)
        )
    }
    for row_number, row in enumerate(rows, start=2):
        try:
            object_code = row.get(effective_mapping["object_code"])
            if object_code and object_code not in object_by_code:
                raise ValueError(f"Unknown controlled-object code: {object_code}")
            candidate = {
                "artifact_id": str(artifact.id),
                "controlled_object_id": (str(object_by_code[object_code]) if object_code else None),
                "field_name": row.get(effective_mapping["field_name"]),
                "value": parse_value(row.get(effective_mapping["value"])),
                "semantic_state": row.get(effective_mapping["semantic_state"]),
                "truth_type": row.get(effective_mapping["truth_type"]),
                "as_of": row.get(effective_mapping["as_of"]),
                "unit": row.get(effective_mapping["unit"]) or None,
                "currency": row.get(effective_mapping["currency"]) or None,
                "measurement_basis": row.get(effective_mapping["measurement_basis"]) or None,
                "confidence": row.get(effective_mapping["confidence"]) or "0.5",
                "source_reliability": (row.get(effective_mapping["source_reliability"]) or None),
                "expires_at": row.get(effective_mapping["expires_at"]) or None,
                "supersedes_item_id": None,
            }
            validated = EvidenceItemCreate.model_validate(candidate)
            normalized.append(validated.model_dump(mode="json"))
        except (ValidationError, ValueError) as exc:
            errors.append({"row": row_number, "message": str(exc)})

    batch = ImportBatch(
        id=uuid.uuid4(),
        organization_id=project.organization_id,
        project_id=project.id,
        artifact_id=artifact.id,
        import_format=Path(artifact.original_filename).suffix.lower().lstrip("."),
        mapping=effective_mapping,
        normalized_rows=normalized,
        validation_errors=errors,
        status=ImportBatchStatus.PREVIEW if not errors else ImportBatchStatus.REJECTED,
        total_rows=len(rows),
        valid_rows=len(normalized),
        rejected_rows=len(rows) - len(normalized),
        idempotency_key=idempotency_key,
        created_by=actor.actor_id,
    )
    session.add(batch)
    record_audit(
        session,
        actor,
        organization_id=project.organization_id,
        project_id=project.id,
        action="EVIDENCE_IMPORT_PREVIEWED",
        object_type="IMPORT_BATCH",
        object_id=str(batch.id),
        details={"valid_rows": batch.valid_rows, "rejected_rows": batch.rejected_rows},
    )
    return batch


def commit_import(
    session: Session,
    actor: ActorContext,
    project: Project,
    batch_id: uuid.UUID,
) -> tuple[ImportBatch, list[uuid.UUID]]:
    batch = session.get(ImportBatch, batch_id)
    if batch is None or batch.project_id != project.id:
        raise HTTPException(status_code=404, detail="Import batch not found")
    existing_ids = list(
        session.scalars(
            select(ImportBatchItem.evidence_item_id).where(ImportBatchItem.batch_id == batch.id)
        )
    )
    if batch.status == ImportBatchStatus.COMMITTED:
        return batch, existing_ids
    if batch.status != ImportBatchStatus.PREVIEW or batch.validation_errors:
        raise HTTPException(status_code=409, detail="Only a clean preview can be committed")

    item_ids: list[uuid.UUID] = []
    for normalized in batch.normalized_rows:
        data = EvidenceItemCreate.model_validate(normalized)
        item = create_evidence_item(session, actor, project, data)
        session.add(ImportBatchItem(batch_id=batch.id, evidence_item_id=item.id))
        item_ids.append(item.id)
    batch.status = ImportBatchStatus.COMMITTED
    batch.committed_at = datetime.now(UTC)
    record_audit(
        session,
        actor,
        organization_id=project.organization_id,
        project_id=project.id,
        action="EVIDENCE_IMPORT_COMMITTED",
        object_type="IMPORT_BATCH",
        object_id=str(batch.id),
        details={"evidence_item_ids": [str(item_id) for item_id in item_ids]},
    )
    return batch, item_ids
