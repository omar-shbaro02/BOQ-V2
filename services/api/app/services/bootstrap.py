from __future__ import annotations

import hashlib
import io
import json
import re
import uuid
from copy import deepcopy
from datetime import date, datetime
from decimal import ROUND_CEILING, Decimal
from pathlib import Path
from typing import Any

from app.auth import ActorContext
from app.bootstrap_schemas import (
    BoqNormalizeCreate,
    BoqSourceCreate,
    PlanningStructureCreate,
    PlanningStructureRevisionCreate,
    ScheduleDraftGenerateCreate,
    ScheduleLogicCreate,
)
from app.generated.taxonomies import (
    BoqExtractionStatus,
    BoqRowClass,
    DependencyBasis,
    DurationBasis,
    DurationStatus,
    PlanningConfidence,
    PlanningReviewState,
    PlanningStructureAction,
    PlanningStructureStatus,
    ProductivitySourceType,
    ScheduleActivityArchetype,
    ScheduleDraftState,
)
from app.models import (
    BoqLine,
    BoqNormalizationRun,
    BoqPlanningStructureVersion,
    BoqSourceRow,
    BoqSourceVersion,
    EvidenceArtifact,
    PlanningAssumption,
    Project,
    ProposedScheduleActivity,
    ScheduleDraftGeneration,
    ScheduleLogicProposal,
)
from app.services.audit import record_audit
from app.storage import EvidenceStore
from fastapi import HTTPException
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

PARSER_NAME = "VAI_BOQ_XLSX"
PARSER_VERSION = "1.0.0"
NORMALIZER_VERSION = "1.0.0"
ACTIVITY_GENERATOR_VERSION = "1.0.0"
SCHEDULE_LOGIC_VERSION = "1.0.0"

CANONICAL_FIELDS = (
    "item_number",
    "division",
    "source_wbs_code",
    "item_name",
    "description",
    "unit",
    "quantity",
    "unit_price",
    "total_price",
    "currency",
    "location",
    "trade",
    "package",
    "notes",
)
HEADER_ALIASES = {
    "item_number": {"item", "item no", "item number", "ref", "reference", "code"},
    "division": {"division", "section", "bill", "bill no"},
    "source_wbs_code": {"wbs", "wbs code", "work breakdown structure"},
    "item_name": {"item name", "name", "work item"},
    "description": {"description", "item description", "scope", "particulars", "details"},
    "unit": {"unit", "uom", "unit of measure"},
    "quantity": {"quantity", "qty", "estimated quantity"},
    "unit_price": {"rate", "unit rate", "unit price", "price"},
    "total_price": {"amount", "total amount", "total price", "value"},
    "currency": {"currency", "ccy"},
    "location": {"location", "area", "zone", "floor", "building"},
    "trade": {"trade", "discipline"},
    "package": {"package", "work package"},
    "notes": {"notes", "remarks", "comment", "comments"},
}
SCHEDULE_RELEVANT_CLASSES = {
    BoqRowClass.DIRECT_EXECUTION_SCOPE,
    BoqRowClass.PROCUREMENT_OR_SUPPLY,
    BoqRowClass.TESTING_COMMISSIONING,
    BoqRowClass.APPROVAL_INSPECTION,
    BoqRowClass.PRELIMINARIES_GENERAL,
}
GROUPING_DIMENSIONS = {
    "source_wbs_code",
    "division",
    "location",
    "trade",
    "package",
    "parent_section",
}


def _json_value(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def _request_hash(data: BoqSourceCreate) -> str:
    payload = data.model_dump(mode="json")
    return hashlib.sha256(
        json.dumps(payload, sort_keys=True, separators=(",", ":")).encode()
    ).hexdigest()


def _extract_xlsx(
    content: bytes,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=False, data_only=False)
    except Exception as exc:
        raise HTTPException(status_code=422, detail="BOQ workbook could not be parsed") from exc

    rows: list[dict[str, Any]] = []
    manifest: list[dict[str, Any]] = []
    source_warnings: list[dict[str, Any]] = []
    sequence = 0
    for sheet in workbook.worksheets:
        merged_ranges = [str(value) for value in sheet.merged_cells.ranges]
        hidden_columns = [key for key, value in sheet.column_dimensions.items() if value.hidden]
        hidden_rows = [key for key, value in sheet.row_dimensions.items() if value.hidden]
        formula_cells = 0
        retained_rows = 0
        for row_number, cells in enumerate(sheet.iter_rows(), start=1):
            values = [_json_value(cell.value) for cell in cells]
            if not any(value is not None and value != "" for value in values):
                continue
            formula_columns = [
                cell.column_letter
                for cell in cells
                if isinstance(cell.value, str) and cell.value.startswith("=")
            ]
            formula_cells += len(formula_columns)
            row_warnings: list[dict[str, Any]] = []
            if formula_columns:
                row_warnings.append(
                    {
                        "code": "FORMULA_PRESERVED_NOT_EVALUATED",
                        "columns": formula_columns,
                    }
                )
            sequence += 1
            retained_rows += 1
            rows.append(
                {
                    "sequence_number": sequence,
                    "sheet_name": sheet.title,
                    "page_number": None,
                    "row_number": row_number,
                    "region": {
                        "cell_range": f"A{row_number}:{get_column_letter(len(cells))}{row_number}"
                    },
                    "raw_values": values,
                    "original_text": None,
                    "extraction_confidence": Decimal("1.0000"),
                    "warnings": row_warnings,
                }
            )
        manifest.append(
            {
                "sheet_name": sheet.title,
                "sheet_state": sheet.sheet_state,
                "max_row": sheet.max_row,
                "max_column": sheet.max_column,
                "retained_nonempty_rows": retained_rows,
                "merged_ranges": merged_ranges,
                "hidden_columns": hidden_columns,
                "hidden_rows": hidden_rows,
                "formula_cell_count": formula_cells,
            }
        )
        if sheet.sheet_state != "visible" or hidden_columns or hidden_rows:
            source_warnings.append(
                {
                    "code": "HIDDEN_WORKBOOK_CONTENT_PRESERVED",
                    "sheet_name": sheet.title,
                    "hidden_columns": hidden_columns,
                    "hidden_rows": hidden_rows,
                    "sheet_state": sheet.sheet_state,
                }
            )
        if merged_ranges:
            source_warnings.append(
                {
                    "code": "MERGED_CELLS_PRESERVED",
                    "sheet_name": sheet.title,
                    "ranges": merged_ranges,
                }
            )
    workbook.close()
    if not rows:
        source_warnings.append({"code": "NO_NONEMPTY_ROWS", "message": "Workbook is empty"})
    return rows, manifest, source_warnings


def create_boq_source(
    session: Session,
    actor: ActorContext,
    project: Project,
    artifact: EvidenceArtifact,
    data: BoqSourceCreate,
    idempotency_key: str,
    store: EvidenceStore,
) -> BoqSourceVersion:
    digest = _request_hash(data)
    existing = session.scalar(
        select(BoqSourceVersion).where(
            BoqSourceVersion.project_id == project.id,
            BoqSourceVersion.idempotency_key == idempotency_key,
        )
    )
    if existing:
        if existing.request_hash != digest:
            raise HTTPException(status_code=409, detail="Idempotency key payload mismatch")
        return existing

    prior = None
    if data.prior_source_version_id:
        prior = session.get(BoqSourceVersion, data.prior_source_version_id)
        if prior is None or prior.project_id != project.id:
            raise HTTPException(status_code=422, detail="Prior BOQ source version not found")
        if artifact.supersedes_artifact_id != prior.artifact_id:
            raise HTTPException(
                status_code=422,
                detail="Revised BOQ artifact must supersede the prior source artifact",
            )
    latest = session.scalar(
        select(BoqSourceVersion)
        .where(BoqSourceVersion.project_id == project.id)
        .order_by(BoqSourceVersion.version_number.desc())
        .limit(1)
    )
    if latest and (prior is None or prior.id != latest.id):
        raise HTTPException(
            status_code=422,
            detail="New BOQ source must name the latest source version as its predecessor",
        )
    already_registered = session.scalar(
        select(BoqSourceVersion.id).where(
            BoqSourceVersion.project_id == project.id,
            BoqSourceVersion.artifact_id == artifact.id,
        )
    )
    if already_registered:
        raise HTTPException(status_code=409, detail="Artifact is already a BOQ source version")

    suffix = Path(artifact.original_filename).suffix.lower()
    content = store.get(artifact.storage_key)
    extracted: list[dict[str, Any]] = []
    manifest: list[dict[str, Any]] = []
    warnings: list[dict[str, Any]] = []
    parser_name = PARSER_NAME
    parser_version = PARSER_VERSION
    confidence = Decimal("1.0000")
    if suffix in {".xlsx", ".xlsm"}:
        extracted, manifest, warnings = _extract_xlsx(content)
        status = (
            BoqExtractionStatus.EXTRACTED
            if extracted
            else BoqExtractionStatus.VERIFICATION_REQUIRED
        )
    elif suffix == ".pdf":
        parser_name = "PRESERVATION_ONLY"
        parser_version = "1.0.0"
        confidence = Decimal("0.0000")
        status = BoqExtractionStatus.VERIFICATION_REQUIRED
        warnings = [
            {
                "code": "PDF_EXTRACTION_ADAPTER_REQUIRED",
                "message": "PDF is preserved but no governed text/OCR parser is configured",
            }
        ]
    else:
        raise HTTPException(status_code=422, detail="BOQ source must be XLSX, XLSM, or PDF")

    next_version = (latest.version_number if latest else 0) + 1
    source = BoqSourceVersion(
        id=uuid.uuid4(),
        organization_id=project.organization_id,
        project_id=project.id,
        artifact_id=artifact.id,
        prior_source_version_id=prior.id if prior else None,
        version_number=next_version,
        source_format=suffix.lstrip("."),
        parser_name=parser_name,
        parser_version=parser_version,
        content_sha256=artifact.sha256,
        extraction_status=status,
        extraction_confidence=confidence,
        structure_manifest=manifest,
        warnings=warnings,
        extracted_row_count=len(extracted),
        idempotency_key=idempotency_key,
        request_hash=digest,
        created_by=actor.actor_id,
    )
    session.add(source)
    for row in extracted:
        session.add(
            BoqSourceRow(
                id=uuid.uuid4(),
                organization_id=project.organization_id,
                project_id=project.id,
                source_version_id=source.id,
                **row,
            )
        )
    record_audit(
        session,
        actor,
        organization_id=project.organization_id,
        project_id=project.id,
        action="BOQ_SOURCE_EXTRACTED",
        object_type="BOQ_SOURCE_VERSION",
        object_id=str(source.id),
        object_version=source.version_number,
        details={
            "artifact_id": str(artifact.id),
            "content_sha256": artifact.sha256,
            "extraction_status": status,
            "extracted_row_count": len(extracted),
            "prior_source_version_id": str(prior.id) if prior else None,
        },
    )
    return source


def _header_token(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(value or "").strip().lower()).strip()


def _header_mapping(values: list[Any]) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for index, value in enumerate(values, start=1):
        token = _header_token(value)
        for field, aliases in HEADER_ALIASES.items():
            if field not in mapping and token in aliases:
                mapping[field] = index
    return mapping


def _text(value: Any) -> str | None:
    if value is None:
        return None
    result = str(value).strip()
    return result or None


def _decimal(value: Any, field: str, warnings: list[dict[str, Any]]) -> Decimal | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        warnings.append({"code": "INVALID_NUMERIC_VALUE", "field": field, "value": value})
        return None
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value))
    text = str(value).strip()
    if text.startswith("="):
        warnings.append({"code": "FORMULA_REQUIRES_VERIFICATION", "field": field})
        return None
    negative = text.startswith("(") and text.endswith(")")
    cleaned = re.sub(r"[^0-9.\-]", "", text.replace(",", ""))
    try:
        result = Decimal(cleaned)
    except Exception:
        warnings.append({"code": "INVALID_NUMERIC_VALUE", "field": field, "value": text})
        return None
    return -result if negative and result > 0 else result


def _value(values: list[Any], mapping: dict[str, int], field: str) -> Any:
    index = mapping.get(field)
    return values[index - 1] if index and index <= len(values) else None


def _classification(
    *,
    values: list[Any],
    normalized: dict[str, Any],
    is_header: bool,
) -> tuple[BoqRowClass, list[str], PlanningConfidence]:
    if is_header or len(_header_mapping(values)) >= 2:
        return BoqRowClass.SUMMARY_HEADER, ["HEADER_COLUMNS_MATCHED"], PlanningConfidence.HIGH
    searchable = " ".join(
        str(value).lower()
        for value in (
            normalized.get("item_name"),
            normalized.get("description"),
            normalized.get("notes"),
            *values,
        )
        if value not in (None, "")
    )
    if re.search(r"\b(grand total|sub[ -]?total|carried forward|brought forward)\b", searchable):
        return BoqRowClass.SUBTOTAL_TOTAL, ["FINANCIAL_TOTAL_KEYWORD"], PlanningConfidence.HIGH
    if re.search(r"\b(provisional sum|provisional|allowance|prime cost sum|pc sum)\b", searchable):
        return (
            BoqRowClass.PROVISIONAL_OR_ALLOWANCE,
            ["PROVISIONAL_ALLOWANCE_KEYWORD"],
            PlanningConfidence.HIGH,
        )
    if re.search(
        r"\b(test(?:ing)?|commission(?:ing)?|balanc(?:e|ing)|startup|start up)\b", searchable
    ):
        return (
            BoqRowClass.TESTING_COMMISSIONING,
            ["TESTING_COMMISSIONING_KEYWORD"],
            PlanningConfidence.MEDIUM,
        )
    if re.search(
        r"\b(inspection|approval|approve|submission|submittal|authority release)\b", searchable
    ):
        return (
            BoqRowClass.APPROVAL_INSPECTION,
            ["APPROVAL_INSPECTION_KEYWORD"],
            PlanningConfidence.MEDIUM,
        )
    if re.search(
        r"\b(preliminar(?:y|ies)|mobilization|mobilisation|site establishment|"
        r"temporary facilit(?:y|ies))\b",
        searchable,
    ):
        return (
            BoqRowClass.PRELIMINARIES_GENERAL,
            ["PRELIMINARIES_KEYWORD"],
            PlanningConfidence.MEDIUM,
        )
    if re.search(r"\b(supply only|material only|materials only)\b", searchable):
        return (
            BoqRowClass.MATERIAL_ONLY_NON_SCHEDULE,
            ["MATERIAL_ONLY_KEYWORD"],
            PlanningConfidence.HIGH,
        )
    if re.search(r"\b(supply|procure|procurement|manufacture|delivery|deliver)\b", searchable):
        return (
            BoqRowClass.PROCUREMENT_OR_SUPPLY,
            ["PROCUREMENT_SUPPLY_KEYWORD"],
            PlanningConfidence.MEDIUM,
        )
    nonempty = [value for value in values if value not in (None, "")]
    if normalized.get("quantity") is not None and (
        normalized.get("description") or normalized.get("item_name")
    ):
        return (
            BoqRowClass.DIRECT_EXECUTION_SCOPE,
            ["DESCRIPTION_AND_QUANTITY_PRESENT"],
            PlanningConfidence.MEDIUM,
        )
    if len(nonempty) == 1 and isinstance(nonempty[0], str):
        return BoqRowClass.SUMMARY_HEADER, ["SINGLE_TEXT_SECTION_ROW"], PlanningConfidence.MEDIUM
    return (
        BoqRowClass.UNKNOWN_REVIEW_REQUIRED,
        ["NO_CONTROLLED_CLASSIFICATION_RULE_MATCHED"],
        PlanningConfidence.INSUFFICIENT,
    )


def normalize_boq_source(
    session: Session,
    actor: ActorContext,
    project: Project,
    source: BoqSourceVersion,
    data: BoqNormalizeCreate,
    idempotency_key: str,
) -> BoqNormalizationRun:
    digest = hashlib.sha256(
        json.dumps(
            {"source_version_id": str(source.id), **data.model_dump(mode="json")},
            sort_keys=True,
            separators=(",", ":"),
        ).encode()
    ).hexdigest()
    existing = session.scalar(
        select(BoqNormalizationRun).where(
            BoqNormalizationRun.project_id == project.id,
            BoqNormalizationRun.idempotency_key == idempotency_key,
        )
    )
    if existing:
        if existing.request_hash != digest:
            raise HTTPException(status_code=409, detail="Idempotency key payload mismatch")
        return existing
    prior_run = session.scalar(
        select(BoqNormalizationRun).where(BoqNormalizationRun.source_version_id == source.id)
    )
    if prior_run:
        raise HTTPException(status_code=409, detail="BOQ source is already normalized")
    if source.extraction_status != BoqExtractionStatus.EXTRACTED:
        raise HTTPException(status_code=422, detail="BOQ source extraction is not ready")

    rows = list(
        session.scalars(
            select(BoqSourceRow)
            .where(BoqSourceRow.source_version_id == source.id)
            .order_by(BoqSourceRow.sequence_number)
        )
    )
    by_sheet: dict[str, list[BoqSourceRow]] = {}
    for row in rows:
        by_sheet.setdefault(row.sheet_name or "", []).append(row)

    header_rows: dict[str, int] = {}
    mappings: dict[str, dict[str, int]] = {}
    run_warnings: list[dict[str, Any]] = []
    for sheet_name, sheet_rows in by_sheet.items():
        explicit_mapping = data.column_mapping.get(sheet_name)
        if explicit_mapping:
            invalid = set(explicit_mapping) - set(CANONICAL_FIELDS)
            if invalid or any(index < 1 for index in explicit_mapping.values()):
                raise HTTPException(status_code=422, detail="Invalid explicit BOQ column mapping")
            mappings[sheet_name] = explicit_mapping
            header_rows[sheet_name] = data.header_rows.get(sheet_name, 0)
            continue
        candidates = [
            (
                len(_header_mapping(row.raw_values)),
                row.row_number or 0,
                _header_mapping(row.raw_values),
            )
            for row in sheet_rows
        ]
        score, row_number, mapping = max(candidates, default=(0, 0, {}), key=lambda value: value[0])
        if score < 2:
            mappings[sheet_name] = {}
            header_rows[sheet_name] = 0
            run_warnings.append(
                {"code": "HEADER_MAPPING_REVIEW_REQUIRED", "sheet_name": sheet_name}
            )
        else:
            mappings[sheet_name] = mapping
            header_rows[sheet_name] = data.header_rows.get(sheet_name, row_number)

    run = BoqNormalizationRun(
        id=uuid.uuid4(),
        organization_id=project.organization_id,
        project_id=project.id,
        source_version_id=source.id,
        normalizer_version=NORMALIZER_VERSION,
        header_rows=header_rows,
        column_mapping=mappings,
        warnings=run_warnings,
        total_rows=len(rows),
        schedule_relevant_rows=0,
        review_required_rows=0,
        idempotency_key=idempotency_key,
        request_hash=digest,
        created_by=actor.actor_id,
    )
    session.add(run)
    parent_sections: dict[str, str] = {}
    for row in rows:
        sheet_name = row.sheet_name or ""
        mapping = mappings[sheet_name]
        extraction_warnings = list(row.warnings)
        warnings = list(extraction_warnings)
        raw = {field: _value(row.raw_values, mapping, field) for field in CANONICAL_FIELDS}
        normalized: dict[str, Any] = {
            field: _text(raw[field])
            for field in CANONICAL_FIELDS
            if field not in {"quantity", "unit_price", "total_price"}
        }
        for field in ("quantity", "unit_price", "total_price"):
            normalized[field] = _decimal(raw[field], field, warnings)
        currency = normalized.get("currency")
        if currency:
            currency = currency.upper()
            if not re.fullmatch(r"[A-Z]{3}", currency):
                warnings.append({"code": "INVALID_CURRENCY", "value": currency})
                currency = None
            normalized["currency"] = currency
        is_header_row = row.row_number == header_rows[sheet_name]
        classification, basis, confidence = _classification(
            values=row.raw_values,
            normalized=normalized,
            is_header=is_header_row,
        )
        if is_header_row:
            warnings = extraction_warnings
        relevant = classification in SCHEDULE_RELEVANT_CLASSES
        review_state = (
            PlanningReviewState.REVIEW_REQUIRED
            if classification == BoqRowClass.UNKNOWN_REVIEW_REQUIRED or warnings
            else PlanningReviewState.NOT_REVIEWED
        )
        if review_state == PlanningReviewState.REVIEW_REQUIRED:
            run.review_required_rows += 1
        if relevant:
            run.schedule_relevant_rows += 1
        source_text = " | ".join(str(value) for value in row.raw_values if value not in (None, ""))
        mapped_indexes = set(mapping.values())
        unmapped = {
            str(index): value
            for index, value in enumerate(row.raw_values, start=1)
            if index not in mapped_indexes and value not in (None, "")
        }
        if (
            classification == BoqRowClass.SUMMARY_HEADER
            and row.row_number != header_rows[sheet_name]
        ):
            section = normalized.get("description") or normalized.get("item_name") or source_text
            if section:
                parent_sections[sheet_name] = section
        stable_digest = hashlib.sha256(
            f"{source.content_sha256}|{sheet_name}|{row.row_number}|{row.sequence_number}".encode()
        ).hexdigest()[:24]
        session.add(
            BoqLine(
                id=uuid.uuid4(),
                stable_line_id=f"BOQ-{stable_digest.upper()}",
                organization_id=project.organization_id,
                project_id=project.id,
                normalization_run_id=run.id,
                source_version_id=source.id,
                source_row_id=row.id,
                source_location={
                    "sheet_name": row.sheet_name,
                    "page_number": row.page_number,
                    "row_number": row.row_number,
                    "region": row.region,
                },
                parent_section=parent_sections.get(sheet_name),
                source_row_text=source_text or row.original_text,
                normalized_values={
                    key: (str(value) if isinstance(value, Decimal) else value)
                    for key, value in normalized.items()
                },
                unmapped_values=unmapped,
                classification=classification,
                classification_basis=basis,
                schedule_relevant=relevant,
                confidence=confidence,
                review_state=review_state,
                warnings=warnings,
                **normalized,
            )
        )
    record_audit(
        session,
        actor,
        organization_id=project.organization_id,
        project_id=project.id,
        action="BOQ_SOURCE_NORMALIZED",
        object_type="BOQ_NORMALIZATION_RUN",
        object_id=str(run.id),
        details={
            "source_version_id": str(source.id),
            "total_rows": run.total_rows,
            "schedule_relevant_rows": run.schedule_relevant_rows,
            "review_required_rows": run.review_required_rows,
            "normalizer_version": NORMALIZER_VERSION,
        },
    )
    return run


def _structure_digest(payload: dict[str, Any]) -> str:
    return hashlib.sha256(
        json.dumps(payload, sort_keys=True, separators=(",", ":")).encode()
    ).hexdigest()


def create_planning_structure(
    session: Session,
    actor: ActorContext,
    project: Project,
    normalization: BoqNormalizationRun,
    data: PlanningStructureCreate,
    idempotency_key: str,
) -> BoqPlanningStructureVersion:
    invalid = set(data.grouping_dimensions) - GROUPING_DIMENSIONS
    if invalid or not data.grouping_dimensions:
        raise HTTPException(status_code=422, detail="Unsupported WBS grouping dimension")
    request = {
        "normalization_run_id": str(normalization.id),
        **data.model_dump(mode="json"),
    }
    digest = _structure_digest(request)
    existing = session.scalar(
        select(BoqPlanningStructureVersion).where(
            BoqPlanningStructureVersion.project_id == project.id,
            BoqPlanningStructureVersion.idempotency_key == idempotency_key,
        )
    )
    if existing:
        if existing.request_hash != digest:
            raise HTTPException(status_code=409, detail="Idempotency key payload mismatch")
        return existing
    if session.scalar(
        select(BoqPlanningStructureVersion.id).where(
            BoqPlanningStructureVersion.normalization_run_id == normalization.id
        )
    ):
        raise HTTPException(status_code=409, detail="Planning structure already exists")

    lines = list(
        session.scalars(
            select(BoqLine)
            .where(BoqLine.normalization_run_id == normalization.id)
            .order_by(BoqLine.stable_line_id)
        )
    )
    root_id = str(uuid.uuid4())
    nodes = [
        {
            "id": root_id,
            "code": project.code,
            "name": project.name,
            "node_type": "PROJECT",
            "parent_id": None,
            "source_basis": "PROJECT_CONTEXT",
            "confidence": "HIGH",
        }
    ]
    node_by_key: dict[tuple[str, str], str] = {}
    package_by_key: dict[tuple[Any, ...], dict[str, Any]] = {}
    mappings: list[dict[str, Any]] = []
    unmapped: list[dict[str, Any]] = []
    assumptions: list[dict[str, Any]] = []
    warnings: list[dict[str, Any]] = []
    for line in lines:
        if not line.schedule_relevant:
            unmapped.append(
                {
                    "boq_line_id": str(line.id),
                    "stable_line_id": line.stable_line_id,
                    "classification": line.classification,
                    "reason": "ROW_CLASS_NOT_DIRECTLY_SCHEDULABLE",
                    "review_required": line.review_state == PlanningReviewState.REVIEW_REQUIRED,
                }
            )
            continue
        dimensions = [
            (dimension, getattr(line, dimension))
            for dimension in data.grouping_dimensions
            if getattr(line, dimension)
        ]
        if dimensions:
            node_type, node_value = dimensions[0]
            node_key = (node_type, str(node_value))
            node_id = node_by_key.get(node_key)
            if node_id is None:
                node_id = str(uuid.uuid4())
                node_by_key[node_key] = node_id
                nodes.append(
                    {
                        "id": node_id,
                        "code": str(node_value) if node_type == "source_wbs_code" else None,
                        "name": str(node_value),
                        "node_type": "SOURCE_WBS"
                        if node_type == "source_wbs_code"
                        else node_type.upper(),
                        "parent_id": root_id,
                        "source_basis": node_type.upper(),
                        "confidence": "HIGH" if node_type == "source_wbs_code" else "MEDIUM",
                    }
                )
        else:
            node_id = root_id
            assumptions.append(
                {
                    "boq_line_id": str(line.id),
                    "assumption": (
                        "Grouped under project because no configured WBS dimension was present"
                    ),
                    "confidence": "LOW",
                    "review_required": True,
                }
            )
        package_key = (*[value or "" for _, value in dimensions], line.classification)
        package = package_by_key.get(package_key)
        if package is None:
            package_id = str(uuid.uuid4())
            label = (
                " / ".join(str(value) for _, value in dimensions)
                or line.classification.replace("_", " ").title()
            )
            package = {
                "id": package_id,
                "code": f"WP-{len(package_by_key) + 1:03d}",
                "name": label,
                "wbs_node_id": node_id,
                "grouping_basis": [name for name, _ in dimensions] + ["classification"],
                "classification": line.classification,
                "review_state": "NOT_REVIEWED",
            }
            package_by_key[package_key] = package
        mappings.append(
            {
                "boq_line_id": str(line.id),
                "stable_line_id": line.stable_line_id,
                "work_package_id": package["id"],
                "mapping_basis": "DETERMINISTIC_GROUPING",
                "confidence": "MEDIUM" if dimensions else "LOW",
                "review_state": "REVIEW_REQUIRED" if not dimensions else "NOT_REVIEWED",
            }
        )
    structure = BoqPlanningStructureVersion(
        id=uuid.uuid4(),
        organization_id=project.organization_id,
        project_id=project.id,
        normalization_run_id=normalization.id,
        supersedes_version_id=None,
        version_number=1,
        status=PlanningStructureStatus.PROPOSED,
        action=PlanningStructureAction.GENERATE,
        wbs_nodes=nodes,
        work_packages=list(package_by_key.values()),
        line_mappings=mappings,
        unmapped_lines=unmapped,
        assumptions=assumptions,
        warnings=warnings,
        change_summary={"generated_packages": len(package_by_key), "mapped_lines": len(mappings)},
        reason="Initial deterministic WBS and work-package proposal",
        idempotency_key=idempotency_key,
        request_hash=digest,
        created_by=actor.actor_id,
    )
    session.add(structure)
    record_audit(
        session,
        actor,
        organization_id=project.organization_id,
        project_id=project.id,
        action="BOQ_PLANNING_STRUCTURE_PROPOSED",
        object_type="BOQ_PLANNING_STRUCTURE_VERSION",
        object_id=str(structure.id),
        object_version=1,
        details=structure.change_summary,
    )
    return structure


def revise_planning_structure(
    session: Session,
    actor: ActorContext,
    project: Project,
    current: BoqPlanningStructureVersion,
    data: PlanningStructureRevisionCreate,
    idempotency_key: str,
) -> BoqPlanningStructureVersion:
    request = {"current_version_id": str(current.id), **data.model_dump(mode="json")}
    digest = _structure_digest(request)
    existing = session.scalar(
        select(BoqPlanningStructureVersion).where(
            BoqPlanningStructureVersion.project_id == project.id,
            BoqPlanningStructureVersion.idempotency_key == idempotency_key,
        )
    )
    if existing:
        if existing.request_hash != digest:
            raise HTTPException(status_code=409, detail="Idempotency key payload mismatch")
        return existing
    latest = session.scalar(
        select(BoqPlanningStructureVersion)
        .where(BoqPlanningStructureVersion.normalization_run_id == current.normalization_run_id)
        .order_by(BoqPlanningStructureVersion.version_number.desc())
        .limit(1)
    )
    if latest is None or latest.id != current.id:
        raise HTTPException(status_code=409, detail="Planning structure version is stale")
    packages = deepcopy(current.work_packages)
    mappings = deepcopy(current.line_mappings)
    warnings = deepcopy(current.warnings)
    package_by_id = {item["id"]: item for item in packages}
    action = data.action
    status = PlanningStructureStatus.PROPOSED
    summary: dict[str, Any] = {"action": action}
    if action == PlanningStructureAction.SPLIT_PACKAGE:
        if len(data.package_ids) != 1 or len(data.new_package_names) != 2 or not data.line_ids:
            raise HTTPException(
                status_code=422, detail="Split requires one package, two names, and lines"
            )
        source_id = data.package_ids[0]
        source = package_by_id.get(source_id)
        source_lines = [item for item in mappings if item["work_package_id"] == source_id]
        selected = {str(value) for value in data.line_ids}
        if source is None or not selected < {item["boq_line_id"] for item in source_lines}:
            raise HTTPException(
                status_code=422, detail="Split lines must be a proper package subset"
            )
        packages.remove(source)
        new_ids = [str(uuid.uuid4()), str(uuid.uuid4())]
        for index, package_id in enumerate(new_ids):
            packages.append(
                {
                    **source,
                    "id": package_id,
                    "code": f"{source['code']}-{index + 1}",
                    "name": data.new_package_names[index],
                    "review_state": "REVISED",
                }
            )
        for item in source_lines:
            item["work_package_id"] = new_ids[0] if item["boq_line_id"] in selected else new_ids[1]
            item["review_state"] = "REVISED"
        summary["split_package_id"] = source_id
        summary["new_package_ids"] = new_ids
    elif action == PlanningStructureAction.MERGE_PACKAGES:
        if len(set(data.package_ids)) < 2 or len(data.new_package_names) != 1:
            raise HTTPException(status_code=422, detail="Merge requires packages and one new name")
        selected_packages = [package_by_id.get(value) for value in data.package_ids]
        if any(value is None for value in selected_packages):
            raise HTTPException(status_code=422, detail="Merge package not found")
        new_id = str(uuid.uuid4())
        base = selected_packages[0]
        assert base is not None
        packages = [item for item in packages if item["id"] not in set(data.package_ids)]
        packages.append(
            {
                **base,
                "id": new_id,
                "code": f"WP-M{current.version_number + 1:03d}",
                "name": data.new_package_names[0],
                "review_state": "REVISED",
            }
        )
        for item in mappings:
            if item["work_package_id"] in data.package_ids:
                item["work_package_id"] = new_id
                item["review_state"] = "REVISED"
        summary["merged_package_ids"] = data.package_ids
        summary["new_package_id"] = new_id
    elif action == PlanningStructureAction.REMAP_LINE:
        if not data.line_ids or data.target_package_id not in package_by_id:
            raise HTTPException(status_code=422, detail="Remap requires lines and target package")
        targets = {str(value) for value in data.line_ids}
        found = 0
        for item in mappings:
            if item["boq_line_id"] in targets:
                item["work_package_id"] = data.target_package_id
                item["mapping_basis"] = "HUMAN_REMAP"
                item["review_state"] = "REVISED"
                found += 1
        if found != len(targets):
            raise HTTPException(status_code=422, detail="Remap line is not currently mapped")
        summary["remapped_line_ids"] = sorted(targets)
    elif action == PlanningStructureAction.ACCEPT_PROPOSAL:
        status = PlanningStructureStatus.REVIEWED
        packages = [{**item, "review_state": "ACCEPTED"} for item in packages]
        mappings = [{**item, "review_state": "ACCEPTED"} for item in mappings]
    elif action == PlanningStructureAction.REJECT_PROPOSAL:
        status = PlanningStructureStatus.REJECTED
    else:
        raise HTTPException(status_code=422, detail="Unsupported planning structure revision")
    revised = BoqPlanningStructureVersion(
        id=uuid.uuid4(),
        organization_id=project.organization_id,
        project_id=project.id,
        normalization_run_id=current.normalization_run_id,
        supersedes_version_id=current.id,
        version_number=current.version_number + 1,
        status=status,
        action=action,
        wbs_nodes=deepcopy(current.wbs_nodes),
        work_packages=packages,
        line_mappings=mappings,
        unmapped_lines=deepcopy(current.unmapped_lines),
        assumptions=deepcopy(current.assumptions),
        warnings=warnings,
        change_summary=summary,
        reason=data.reason,
        idempotency_key=idempotency_key,
        request_hash=digest,
        created_by=actor.actor_id,
    )
    session.add(revised)
    record_audit(
        session,
        actor,
        organization_id=project.organization_id,
        project_id=project.id,
        action="BOQ_PLANNING_STRUCTURE_REVISED",
        object_type="BOQ_PLANNING_STRUCTURE_VERSION",
        object_id=str(revised.id),
        object_version=revised.version_number,
        details=summary,
    )
    return revised


ACTIVITY_TYPES_BY_CLASS = {
    BoqRowClass.DIRECT_EXECUTION_SCOPE: [ScheduleActivityArchetype.EXECUTION],
    BoqRowClass.PROCUREMENT_OR_SUPPLY: [
        ScheduleActivityArchetype.SUBMITTAL,
        ScheduleActivityArchetype.PROCUREMENT,
        ScheduleActivityArchetype.DELIVERY,
    ],
    BoqRowClass.TESTING_COMMISSIONING: [ScheduleActivityArchetype.TESTING_COMMISSIONING],
    BoqRowClass.APPROVAL_INSPECTION: [ScheduleActivityArchetype.INSPECTION_RELEASE],
    BoqRowClass.PRELIMINARIES_GENERAL: [ScheduleActivityArchetype.MOBILIZATION],
}


def generate_schedule_draft(
    session: Session,
    actor: ActorContext,
    project: Project,
    structure: BoqPlanningStructureVersion,
    data: ScheduleDraftGenerateCreate,
    idempotency_key: str,
) -> ScheduleDraftGeneration:
    request = {"planning_structure_id": str(structure.id), **data.model_dump(mode="json")}
    digest = _structure_digest(request)
    existing = session.scalar(
        select(ScheduleDraftGeneration).where(
            ScheduleDraftGeneration.project_id == project.id,
            ScheduleDraftGeneration.idempotency_key == idempotency_key,
        )
    )
    if existing:
        if existing.request_hash != digest:
            raise HTTPException(status_code=409, detail="Idempotency key payload mismatch")
        return existing
    latest = session.scalar(
        select(BoqPlanningStructureVersion)
        .where(BoqPlanningStructureVersion.normalization_run_id == structure.normalization_run_id)
        .order_by(BoqPlanningStructureVersion.version_number.desc())
        .limit(1)
    )
    if latest is None or latest.id != structure.id:
        raise HTTPException(status_code=409, detail="Planning structure version is stale")
    if structure.status != PlanningStructureStatus.REVIEWED:
        raise HTTPException(status_code=422, detail="Planner-reviewed structure is required")
    if session.scalar(
        select(ScheduleDraftGeneration.id).where(
            ScheduleDraftGeneration.planning_structure_id == structure.id
        )
    ):
        raise HTTPException(status_code=409, detail="Schedule draft already generated")

    package_ids = {item["id"] for item in structure.work_packages}
    productivity_by_package = {item.work_package_id: item for item in data.productivity_inputs}
    if len(productivity_by_package) != len(data.productivity_inputs):
        raise HTTPException(status_code=422, detail="Duplicate productivity input")
    duration_by_key = {
        (item.work_package_id, item.activity_type): item for item in data.duration_inputs
    }
    if len(duration_by_key) != len(data.duration_inputs):
        raise HTTPException(status_code=422, detail="Duplicate duration input")
    input_packages = set(productivity_by_package) | {
        item.work_package_id for item in data.duration_inputs
    }
    if not input_packages <= package_ids:
        raise HTTPException(status_code=422, detail="Duration/productivity package not found")

    normalization = session.get(BoqNormalizationRun, structure.normalization_run_id)
    if normalization is None:
        raise HTTPException(status_code=422, detail="BOQ normalization unavailable")
    lines = {
        str(line.id): line
        for line in session.scalars(
            select(BoqLine).where(BoqLine.normalization_run_id == normalization.id)
        )
    }
    mappings_by_package: dict[str, list[dict[str, Any]]] = {}
    for mapping in structure.line_mappings:
        mappings_by_package.setdefault(mapping["work_package_id"], []).append(mapping)

    generation = ScheduleDraftGeneration(
        id=uuid.uuid4(),
        organization_id=project.organization_id,
        project_id=project.id,
        planning_structure_id=structure.id,
        generator_version=ACTIVITY_GENERATOR_VERSION,
        draft_state=ScheduleDraftState.DRAFT_GENERATED,
        productivity_inputs=data.model_dump(mode="json")["productivity_inputs"],
        duration_inputs=data.model_dump(mode="json")["duration_inputs"],
        warnings=[],
        activity_count=0,
        unresolved_duration_count=0,
        idempotency_key=idempotency_key,
        request_hash=digest,
        created_by=actor.actor_id,
    )
    session.add(generation)
    assumptions: list[PlanningAssumption] = []
    sequence = 0
    for package in structure.work_packages:
        package_mappings = mappings_by_package.get(package["id"], [])
        if not package_mappings:
            generation.warnings.append(
                {"code": "EMPTY_WORK_PACKAGE_SKIPPED", "work_package_id": package["id"]}
            )
            continue
        package_lines = [lines[item["boq_line_id"]] for item in package_mappings]
        activity_types = ACTIVITY_TYPES_BY_CLASS.get(BoqRowClass(package["classification"]), [])
        if not activity_types:
            continue
        quantities = [line.quantity for line in package_lines if line.quantity is not None]
        units = {line.unit for line in package_lines if line.quantity is not None and line.unit}
        quantity = sum(quantities, Decimal("0")) if quantities and len(units) == 1 else None
        unit = next(iter(units)) if quantity is not None else None
        if quantities and (len(units) != 1 or len(quantities) != len(package_lines)):
            generation.warnings.append(
                {
                    "code": "PACKAGE_QUANTITY_NOT_AGGREGATABLE",
                    "work_package_id": package["id"],
                }
            )
        for activity_type in activity_types:
            sequence += 1
            activity_id = uuid.uuid4()
            warnings: list[dict[str, Any]] = []
            explicit = duration_by_key.get((package["id"], activity_type))
            productivity_input = productivity_by_package.get(package["id"])
            duration: Decimal | None = None
            unrounded: Decimal | None = None
            productivity: dict[str, Any] | None = None
            confidence = PlanningConfidence.INSUFFICIENT
            if explicit:
                duration = explicit.duration_working_days
                unrounded = duration
                duration_status = DurationStatus.EXPLICIT
                duration_basis = explicit.basis
                confidence = explicit.confidence
            elif (
                activity_type == ScheduleActivityArchetype.EXECUTION
                and productivity_input
                and quantity is not None
                and productivity_input.quantity_unit.casefold() == (unit or "").casefold()
            ):
                unrounded = quantity / productivity_input.rate_per_working_day
                duration = unrounded.quantize(Decimal("1"), rounding=ROUND_CEILING)
                duration_status = DurationStatus.CALCULATED
                duration_basis = (
                    DurationBasis.PROJECT_PRODUCTIVITY
                    if productivity_input.source_type
                    in {
                        ProductivitySourceType.PROJECT_HISTORICAL_ACTUAL,
                        ProductivitySourceType.PLANNER_INPUT,
                        ProductivitySourceType.SUBCONTRACTOR_PLAN,
                    }
                    else DurationBasis.CONTROLLED_PRODUCTIVITY_LIBRARY
                )
                confidence = productivity_input.confidence
                productivity = productivity_input.model_dump(mode="json")
            else:
                duration_status = DurationStatus.VERIFICATION_REQUIRED
                duration_basis = DurationBasis.UNRESOLVED
                generation.unresolved_duration_count += 1
                warnings.append({"code": "DURATION_BASIS_REQUIRED"})
            activity = ProposedScheduleActivity(
                id=activity_id,
                organization_id=project.organization_id,
                project_id=project.id,
                generation_id=generation.id,
                activity_code=f"ACT-{sequence:04d}",
                wbs_node_id=package["wbs_node_id"],
                work_package_id=package["id"],
                boq_line_refs=[item["boq_line_id"] for item in package_mappings],
                activity_name=(
                    f"{activity_type.value.replace('_', ' ').title()} — {package['name']}"
                ),
                activity_type=activity_type,
                trade=next((line.trade for line in package_lines if line.trade), None),
                location=next((line.location for line in package_lines if line.location), None),
                description=f"Proposed {activity_type.value} activity for {package['name']}",
                quantity=quantity,
                unit=unit,
                duration_working_days=duration,
                duration_unrounded=unrounded,
                duration_status=duration_status,
                duration_basis=duration_basis,
                productivity=productivity,
                calendar_id=None,
                responsible_role=None,
                confidence=confidence,
                review_state=PlanningReviewState.REVIEW_REQUIRED,
                warnings=warnings,
            )
            session.add(activity)
            if duration_status == DurationStatus.VERIFICATION_REQUIRED:
                assumptions.append(
                    PlanningAssumption(
                        id=uuid.uuid4(),
                        organization_id=project.organization_id,
                        project_id=project.id,
                        generation_id=generation.id,
                        proposition=f"Duration is unresolved for {activity.activity_code}",
                        reason_needed=(
                            "No defensible explicit duration or matching productivity basis"
                        ),
                        affected_activity_ids=[str(activity.id)],
                        source_basis="UNRESOLVED",
                        confidence=PlanningConfidence.INSUFFICIENT,
                        consequence_if_wrong="Schedule dates and critical path may be unreliable",
                        validation_owner=data.validation_owner,
                        status="OPEN",
                        resolution_note=None,
                    )
                )
            elif productivity_input and productivity_input.source_type in {
                ProductivitySourceType.COMPANY_HISTORICAL_BENCHMARK,
                ProductivitySourceType.CONTROLLED_REFERENCE_LIBRARY,
                ProductivitySourceType.EXPLICIT_ASSUMPTION,
            }:
                assumptions.append(
                    PlanningAssumption(
                        id=uuid.uuid4(),
                        organization_id=project.organization_id,
                        project_id=project.id,
                        generation_id=generation.id,
                        proposition=(
                            f"Productivity {productivity_input.rate_per_working_day} "
                            f"{productivity_input.quantity_unit}/working day applies"
                        ),
                        reason_needed="Non-project productivity is required to propose duration",
                        affected_activity_ids=[str(activity.id)],
                        source_basis=productivity_input.source_reference,
                        confidence=productivity_input.confidence,
                        consequence_if_wrong="Activity duration and downstream dates may change",
                        validation_owner=data.validation_owner,
                        status="OPEN",
                        resolution_note=None,
                    )
                )
    generation.activity_count = sequence
    if generation.unresolved_duration_count:
        generation.draft_state = ScheduleDraftState.VERIFICATION_REQUIRED
    session.add_all(assumptions)
    record_audit(
        session,
        actor,
        organization_id=project.organization_id,
        project_id=project.id,
        action="SCHEDULE_DRAFT_GENERATED",
        object_type="SCHEDULE_DRAFT_GENERATION",
        object_id=str(generation.id),
        details={
            "planning_structure_id": str(structure.id),
            "activity_count": generation.activity_count,
            "unresolved_duration_count": generation.unresolved_duration_count,
            "draft_state": generation.draft_state,
        },
    )
    return generation


def create_schedule_logic(
    session: Session,
    actor: ActorContext,
    project: Project,
    generation: ScheduleDraftGeneration,
    data: ScheduleLogicCreate,
    idempotency_key: str,
) -> ScheduleLogicProposal:
    request = {"generation_id": str(generation.id), **data.model_dump(mode="json")}
    digest = _structure_digest(request)
    existing = session.scalar(
        select(ScheduleLogicProposal).where(
            ScheduleLogicProposal.project_id == project.id,
            ScheduleLogicProposal.idempotency_key == idempotency_key,
        )
    )
    if existing:
        if existing.request_hash != digest:
            raise HTTPException(status_code=409, detail="Idempotency key payload mismatch")
        return existing
    if session.scalar(
        select(ScheduleLogicProposal.id).where(ScheduleLogicProposal.generation_id == generation.id)
    ):
        raise HTTPException(status_code=409, detail="Schedule logic already proposed")
    activities = list(
        session.scalars(
            select(ProposedScheduleActivity)
            .where(ProposedScheduleActivity.generation_id == generation.id)
            .order_by(ProposedScheduleActivity.activity_code)
        )
    )
    activity_by_id = {str(item.id): item for item in activities}
    dependencies: list[dict[str, Any]] = []
    warnings: list[dict[str, Any]] = []
    seen: set[tuple[str, str, str, str]] = set()

    def add_dependency(payload: dict[str, Any]) -> None:
        predecessor = payload["predecessor_activity_id"]
        successor = payload["successor_activity_id"]
        if predecessor == successor:
            raise HTTPException(status_code=422, detail="Dependency cannot reference itself")
        if predecessor not in activity_by_id or successor not in activity_by_id:
            raise HTTPException(status_code=422, detail="Dependency activity not found")
        key = (
            predecessor,
            successor,
            payload["relation_type"],
            str(payload["lag_working_days"]),
        )
        if key in seen:
            raise HTTPException(status_code=422, detail="Duplicate dependency")
        seen.add(key)
        dependencies.append(payload)

    by_package: dict[str, dict[str, ProposedScheduleActivity]] = {}
    for activity in activities:
        by_package.setdefault(activity.work_package_id, {})[activity.activity_type] = activity
    for package_activities in by_package.values():
        chain = [
            package_activities.get(ScheduleActivityArchetype.SUBMITTAL),
            package_activities.get(ScheduleActivityArchetype.PROCUREMENT),
            package_activities.get(ScheduleActivityArchetype.DELIVERY),
        ]
        for predecessor, successor in zip(chain, chain[1:], strict=False):
            if predecessor and successor:
                add_dependency(
                    {
                        "predecessor_activity_id": str(predecessor.id),
                        "successor_activity_id": str(successor.id),
                        "relation_type": "FINISH_TO_START",
                        "lag_working_days": "0",
                        "basis": DependencyBasis.PROCUREMENT_PREREQUISITE,
                        "source_reference": "VAI controlled procurement chain",
                        "confidence": "HIGH",
                        "review_state": "NOT_REVIEWED",
                        "generated": True,
                    }
                )
    for dependency in data.dependencies:
        add_dependency({**dependency.model_dump(mode="json"), "generated": False})
    templates = [item.model_dump(mode="json") for item in data.sequence_templates]
    for template in data.sequence_templates:
        if template.basis != DependencyBasis.VALIDATED_SEQUENCE_TEMPLATE:
            raise HTTPException(
                status_code=422,
                detail="Sequence template dependency requires validated-template basis",
            )
        add_dependency({**template.model_dump(mode="json"), "generated": True})

    if data.calendar:
        calendar = data.calendar.model_dump(mode="json")
        assumptions: list[dict[str, Any]] = []
        review_state = data.calendar.review_state
    else:
        calendar = {
            "calendar_id": "PROPOSED-DEFAULT",
            "name": "Proposed five-day calendar",
            "working_weekdays": [0, 1, 2, 3, 4],
            "working_hours_per_day": "8",
            "holidays": [],
            "shift_pattern": None,
            "review_state": "REVIEW_REQUIRED",
        }
        assumptions = [
            {
                "proposition": "Monday-Friday, eight working hours per day",
                "reason_needed": "No project working calendar was supplied",
                "source_basis": "VISIBLE_DEFAULT_ASSUMPTION",
                "confidence": "LOW",
                "consequence_if_wrong": "All calculated working dates may change",
                "validation_owner": data.validation_owner,
                "status": "OPEN",
            }
        ]
        review_state = PlanningReviewState.REVIEW_REQUIRED
    activity_ids = set(activity_by_id)
    milestones = []
    for milestone in data.milestones:
        feeds = [str(value) for value in milestone.feeds_from_activity_ids]
        if not set(feeds) <= activity_ids:
            raise HTTPException(status_code=422, detail="Milestone feeder activity not found")
        milestones.append(
            {
                **milestone.model_dump(mode="json"),
                "feeds_from_activity_ids": feeds,
                "status": "PROPOSED",
                "authorized": False,
            }
        )
    constraints = []
    for constraint in data.constraints:
        if str(constraint.activity_id) not in activity_ids:
            raise HTTPException(status_code=422, detail="Constraint activity not found")
        constraints.append(
            {**constraint.model_dump(mode="json"), "status": "PROPOSED", "authorized": False}
        )
    proposal = ScheduleLogicProposal(
        id=uuid.uuid4(),
        organization_id=project.organization_id,
        project_id=project.id,
        generation_id=generation.id,
        logic_version=SCHEDULE_LOGIC_VERSION,
        dependencies=dependencies,
        milestones=milestones,
        constraints=constraints,
        calendar=calendar,
        sequence_templates=templates,
        assumptions=assumptions,
        warnings=warnings,
        review_state=review_state,
        idempotency_key=idempotency_key,
        request_hash=digest,
        created_by=actor.actor_id,
    )
    session.add(proposal)
    record_audit(
        session,
        actor,
        organization_id=project.organization_id,
        project_id=project.id,
        action="SCHEDULE_LOGIC_PROPOSED",
        object_type="SCHEDULE_LOGIC_PROPOSAL",
        object_id=str(proposal.id),
        details={
            "generation_id": str(generation.id),
            "dependency_count": len(dependencies),
            "milestone_count": len(milestones),
            "calendar_review_state": review_state,
        },
    )
    return proposal
