"""Project-specific editable progress templates; no shared company or sample data."""

import hashlib
import io
import re
from datetime import date

from app.archive_limits import load_workbook
from app.services.ms_project_export import HEADERS
from app.services.progress_columns import TASK_PROGRESS_HEADERS
from app.services.progress_workbook import _date, _duration
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

TEMPLATE_VERSION = "VAI-PROGRESS-TEMPLATE-2"


def build_progress_template(
    content: bytes, *, project_id: str, as_of: date, sample: bool = False
) -> bytes:
    try:
        workbook = load_workbook(io.BytesIO(content), data_only=False, keep_links=False)
    except Exception as exc:
        raise ValueError("Upload a readable VAI XLSX schedule workbook") from exc
    try:
        if "Tasks" not in workbook:
            raise ValueError("Schedule must contain a Tasks sheet")
        tasks = workbook["Tasks"]
        if tuple(c.value for c in tasks[1][:6]) != HEADERS:
            raise ValueError("Tasks must use the VAI six-column schedule headers")
        rows, seen = [], set()
        for cells in tasks.iter_rows(min_row=2):
            values = [cell.value for cell in cells[:6]]
            if values[0] is None:
                continue
            if any(cell.data_type == "f" for cell in cells[:6]):
                raise ValueError("Schedule task inputs must be values, not formulas")
            try:
                task_id = int(values[0])
            except (ValueError, TypeError) as exc:
                raise ValueError("Task IDs must be unique positive integers") from exc
            if task_id < 1 or str(task_id) != str(values[0]) or task_id in seen:
                raise ValueError("Task IDs must be unique positive integers")
            seen.add(task_id)
            start, finish = _date(values[3], "Planned Start"), _date(values[4], "Planned Finish")
            if finish < start:
                raise ValueError("Planned finish must not precede planned start")
            if not values[1]:
                raise ValueError("Every schedule task needs an activity name")
            rows.append((values, start, finish, _duration(values[2])))
        if not rows:
            raise ValueError("Schedule must contain at least one task")

        for title in ("Progress Report", "Schedule Comparison", "Progress Template Metadata"):
            if title in workbook:
                del workbook[title]
        progress = workbook.create_sheet("Progress Report", 0)
        comparison = workbook.create_sheet("Schedule Comparison", 1)
        metadata = workbook.create_sheet("Progress Template Metadata")
        mode = "SYNTHETIC_SAMPLE" if sample else "BLANK_TEMPLATE"
        progress.append(
            [
                "FICTIONAL SAMPLE" if sample else "PROGRESS INPUT",
                "Edit yellow cells F–J. Preserve task IDs, task names and dates.",
            ]
        )
        progress.append(
            [
                "Evidence boundary",
                "Reported inputs require separate verification. "
                "Schedule dates remain proposals unless separately authorized.",
            ]
        )
        progress.append([])
        progress.append(TASK_PROGRESS_HEADERS)
        comparison.append(
            [
                "Progress against source schedule",
                "Calendar-day approximation; not a critical-path deadline forecast.",
            ]
        )
        comparison.append(["Reporting cut-off", as_of])
        comparison.append(
            [
                "FICTIONAL SAMPLE" if sample else "Completion input",
                "Use 0–100 in Progress Report. Blank means not reported. "
                "Positive variance = ahead; negative = behind. "
                "Status tolerance ±5 percentage points.",
            ]
        )
        comparison.append(
            [
                "Task ID",
                "Task Name",
                "Planned Start",
                "Planned Finish",
                "Duration days",
                "Planned %",
                "Reported %",
                "Variance pp",
                "Approx. days +/−",
                "Status",
            ]
        )
        for i, (values, start, finish, duration) in enumerate(rows, 5):
            task_id = values[0]
            actual = None
            if sample:
                planned = (
                    0
                    if as_of < start
                    else 100
                    if as_of >= finish
                    else ((as_of - start).days + 1) / max(1, (finish - start).days + 1) * 100
                )
                actual = (
                    max(0, min(100, round(planned + (-25, 0, 15)[(i - 5) % 3]))) if planned else 0
                )
            progress_values = [
                task_id,
                values[1],
                values[2],
                values[3],
                values[4],
                actual,
                as_of if sample else None,
                "FICTIONAL EXAMPLE — replace with site observations" if sample else None,
                None,
                None,
                "SOURCE SCHEDULE — authorization not asserted",
            ]
            progress.append(
                [
                    "'" + value
                    if isinstance(value, str) and value.startswith(("=", "+", "@"))
                    else value
                    for value in progress_values
                ]
            )
            comparison.append(
                [
                    task_id,
                    f"='Progress Report'!B{i}",
                    start,
                    finish,
                    duration,
                    f"=IF($B$2<C{i},0,IF($B$2>=D{i},100,MIN(100,($B$2-C{i}+1)/(D{i}-C{i}+1)*100)))",
                    f"=IF(ISNUMBER('Progress Report'!F{i}),'Progress Report'!F{i},\"\")",
                    f'=IF(ISNUMBER(G{i}),G{i}-F{i},"")',
                    f'=IF(ISNUMBER(H{i}),SIGN(H{i})*ROUND(ABS(H{i})/100*E{i},1),"")',
                    f'=IF(ISNUMBER(G{i}),IF(H{i}>5,"EARLY",'
                    f'IF(H{i}<-5,"LATE","ON TIME")),"NOT REPORTED")',
                ]
            )
        metadata.append(["Field", "Value"])
        for key, value in {
            "template_version": TEMPLATE_VERSION,
            "mode": mode,
            "project_id": project_id,
            "source_sha256": hashlib.sha256(content).hexdigest(),
            "task_count": len(rows),
            "reporting_cutoff": as_of.isoformat(),
            "comparison_basis": (
                "Inclusive calendar-day linear approximation; proposed duration equivalent; "
                "no critical-path forecast"
            ),
            "semantic_boundary": (
                "No verification, authorization or project-wide weighting is created"
            ),
        }.items():
            metadata.append([key, value])
        for sheet in (progress, comparison):
            sheet.freeze_panes = "C5"
            sheet.sheet_view.showGridLines = False
            sheet.auto_filter.ref = f"A4:{'K' if sheet is progress else 'J'}{sheet.max_row}"
            for cell in sheet[4]:
                cell.fill = PatternFill("solid", fgColor="781633")
                cell.font = Font(bold=True, color="FFFFFF")
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            sheet.row_dimensions[4].height = 34
            for row in sheet.iter_rows(min_row=5):
                for cell in row:
                    if (
                        isinstance(cell.value, str)
                        and cell.data_type != "f"
                        and re.match(r"^[=+@]", cell.value)
                    ):
                        cell.value = "'" + cell.value
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                sheet.row_dimensions[row[0].row].height = 45
            for col in "ABCDEFGHIJK":
                sheet.column_dimensions[col].width = 20
            sheet.column_dimensions["B"].width = 64
            sheet.column_dimensions["A"].width = 12
        progress.column_dimensions["H"].width = 50
        progress.column_dimensions["K"].width = 40
        for row in (1, 2):
            note = f"{progress.cell(row, 1).value}: {progress.cell(row, 2).value}"
            progress.merge_cells(start_row=row, start_column=1, end_row=row, end_column=11)
            progress.cell(row, 1).value = note
            progress.cell(row, 1).alignment = Alignment(wrap_text=True, vertical="center")
            progress.row_dimensions[row].height = 30
        comparison.merge_cells("A1:J1")
        comparison["A1"] = (
            "Progress against source schedule — calendar-day approximation, "
            "not a critical-path deadline forecast"
        )
        comparison.row_dimensions[1].height = 30
        comparison.column_dimensions["A"].width = 24
        comparison.merge_cells("B3:J3")
        comparison["B3"].alignment = Alignment(wrap_text=True, vertical="center")
        comparison.row_dimensions[3].height = 30
        for row in progress.iter_rows(min_row=5, min_col=6, max_col=10):
            for cell in row:
                cell.fill = PatternFill("solid", fgColor="FFF2CC")
        validation = DataValidation(
            type="decimal", operator="between", formula1=0, formula2=100, allow_blank=True
        )
        validation.showErrorMessage = True
        validation.error = "Enter a percentage between 0 and 100."
        progress.add_data_validation(validation)
        validation.add(f"F5:F{progress.max_row}")
        comparison["B2"].fill = PatternFill("solid", fgColor="FFF2CC")
        comparison["B2"].number_format = "yyyy-mm-dd"
        for row in comparison.iter_rows(min_row=5):
            for cell in row[2:4]:
                cell.number_format = "yyyy-mm-dd"
            for cell in row[5:9]:
                cell.number_format = "0.0"
        for row in progress.iter_rows(min_row=5):
            row[6].number_format = "yyyy-mm-dd"
        metadata.column_dimensions["A"].width = 26
        metadata.column_dimensions["B"].width = 100
        output = io.BytesIO()
        workbook.save(output)
        return output.getvalue()
    finally:
        workbook.close()
