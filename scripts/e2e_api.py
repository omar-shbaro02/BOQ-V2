"""Disposable browser-test API. Never loads developer credentials or a persistent database."""

import os
import sys
import tempfile
import uuid
from pathlib import Path

import uvicorn

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "services" / "api"))


def main() -> None:
    scratch = ROOT / ".test-tmp"
    scratch.mkdir(exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="browser-", dir=scratch) as directory:
        os.environ.update(
            VAI_ENVIRONMENT="test",
            VAI_DATABASE_URL=f"sqlite+pysqlite:///{Path(directory) / 'test.db'}",
            VAI_EVIDENCE_STORAGE_PATH=str(Path(directory) / "evidence"),
            VAI_CORS_ORIGINS='["http://127.0.0.1:3011"]',
            VAI_OPENAI_API_KEY="",
            VAI_OPENAI_AGENTS_ENABLED="false",
        )
        from datetime import UTC, datetime

        from app.database import Base, engine
        from app.main import app
        from app.models import (
            AuthorizedContextVersion,
            ControlledObject,
            DecisionCaseShell,
            Membership,
            Organization,
            Project,
        )
        from app.services.openai_schedule import build_openai_schedule_workbook
        from sqlalchemy.orm import Session

        Base.metadata.create_all(engine)
        org_id, project_id = uuid.UUID(int=101), uuid.UUID(int=102)
        with Session(engine) as session:
            session.add(Organization(id=org_id, name="Browser fixture", slug="browser-fixture"))
            session.flush()
            session.add(
                ControlledObject(
                    id=uuid.UUID(int=103),
                    organization_id=org_id,
                    project_id=project_id,
                    object_type="ACTIVITY",
                    code="A1",
                    name="Excavation",
                )
            )
            session.add(
                AuthorizedContextVersion(
                    id=uuid.UUID(int=104),
                    organization_id=org_id,
                    project_id=project_id,
                    context_type="SCHEDULE",
                    version_number=1,
                    semantic_state="CURRENT_AUTHORIZED",
                    effective_from=datetime(2026, 9, 1, tzinfo=UTC),
                    activated_at=datetime(2026, 9, 1, tzinfo=UTC),
                    created_by="browser@example.test",
                    payload={
                        "data_date": "2026-10-05",
                        "activities": [
                            {
                                "code": "A1",
                                "name": "Excavation",
                                "planned_start": "2026-10-01",
                                "planned_finish": "2026-10-10",
                                "controlled_object_code": "A1",
                                "calendar_id": "SITE",
                                "progress_plan": [],
                            }
                        ],
                        "calendars": [
                            {
                                "calendar_id": "SITE",
                                "working_weekdays": [0, 1, 2, 3, 4],
                                "holidays": [],
                            }
                        ],
                        "dependencies": [],
                    },
                )
            )
            session.add(
                Project(
                    id=project_id,
                    organization_id=org_id,
                    code="E2E",
                    name="Browser project",
                    timezone="Asia/Beirut",
                    currency="USD",
                    delivery_model="DESIGN_BID_BUILD",
                    reporting_cadence="WEEKLY",
                )
            )
            session.flush()
            session.add(
                Membership(
                    organization_id=org_id,
                    project_id=project_id,
                    actor_id="browser@example.test",
                    role="PROJECT_ADMIN",
                )
            )
            session.add(
                DecisionCaseShell(
                    organization_id=org_id,
                    project_id=project_id,
                    case_number="DC-E2E-001",
                    title="Evidence gap requiring human review",
                    case_type="PROGRESS",
                    lifecycle="OPEN",
                    owner_actor_id="browser@example.test",
                    created_by="browser@example.test",
                    governance_state="HUMAN_REVIEW_REQUIRED",
                )
            )
            session.commit()
        import io

        from openpyxl import Workbook, load_workbook

        content = build_openai_schedule_workbook(
            [
                {
                    "id": 1,
                    "name": "Excavation",
                    "duration_days": 7,
                    "start": "2026-10-01",
                    "finish": "2026-10-10",
                    "dependency_ids": [],
                    "boq_references": [],
                }
            ]
        )
        workbook = load_workbook(io.BytesIO(content))
        workbook["Progress Report"]["G5"] = 25
        workbook["Progress Report"]["H5"] = "2026-10-05"
        workbook.save(scratch / "browser-progress.xlsx")
        workbook.close()
        boq = Workbook()
        sheet = boq.active
        sheet.title = "Concrete works"
        sheet.append(["Item", "Description", "Quantity", "Unit"])
        sheet.append(["C-01", "Concrete foundations", 120, "m3"])
        sheet.append(["C-02", "Concrete columns", 80, "m3"])
        sheet.append([None, "Subtotal", 200, "m3"])
        boq.save(scratch / "browser-boq.xlsx")
        sheet.append(["P-01", "Supply pumps", 2, "No"])
        boq.save(scratch / "browser-boq-procurement.xlsx")
        sheet.delete_rows(5)
        sheet.append(["C-03", "Concrete beams", 20, "m3"])
        boq.save(scratch / "browser-boq-v2.xlsx")
        boq.close()
        # Share the reviewed two-page PDF fixture with the parser's real-engine checks.
        sys.path.append(str(ROOT / "services" / "api" / "tests"))
        from test_boq_pdf import text_pdf

        (scratch / "browser-boq.pdf").write_bytes(text_pdf())
        uvicorn.run(app, host="127.0.0.1", port=8011)
        engine.dispose()


if __name__ == "__main__":
    main()
